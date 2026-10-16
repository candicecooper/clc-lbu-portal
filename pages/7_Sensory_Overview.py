import streamlit as st
from supabase import create_client
import json
import io
from datetime import date

# ── Page config ────────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Sensory Profiles",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="collapsed",
)

TEAL  = "#1A7A7A"
AMBER = "#C8760A"

# Sensory Overview entry tool inside the portal (pages/ folder)
SENSORY_OVERVIEW_PAGE = "pages/7_Sensory_Overview.py"
DOC_TYPE = "Sensory Overview"

# ── Document Archive / Student Register integration ───────────────────────────
# archive_utils self-locates the students project (SUPABASE_URL_STUDENTS first,
# verified against document_archive), so the register and archive are read from
# the right database even though sensory_overviews lives in the default project.
try:
    from archive_utils import archive_document, fetch_archive, fetch_register, initials_from_name
    ARCHIVE_OK = True
    ARCHIVE_ERR = ""
except Exception as _imp_err:  # noqa: BLE001
    ARCHIVE_OK = False
    ARCHIVE_ERR = str(_imp_err)

    def archive_document(**_k):
        return None

    def fetch_archive(**_k):
        return []

    def fetch_register(_status=None):
        return []

    def initials_from_name(name):
        parts = [p for p in str(name or "").replace("_", " ").split() if p]
        return "".join(p[0].upper() for p in parts[:2]) or None

st.markdown(f"""
<style>
    .stApp {{ background-color: #f5f7f7; }}
    h1, h2, h3 {{ color: {TEAL}; }}

    .student-card {{
        background: white;
        border-radius: 10px;
        border-left: 6px solid {TEAL};
        padding: 14px 18px;
        margin-bottom: 10px;
        box-shadow: 0 1px 4px rgba(0,0,0,0.07);
    }}
    .student-card .sname  {{ font-size: 1.1rem; font-weight: 700; color: {TEAL}; }}
    .student-card .smeta  {{ font-size: 0.8rem; color: #666; margin-top: 2px; }}

    .pill {{
        display: inline-block;
        padding: 3px 10px;
        border-radius: 12px;
        font-size: 0.75rem;
        font-weight: 700;
        margin: 2px 2px;
    }}
    .pill.high {{ background: #fde8e8; color: #c62828; }}
    .pill.med  {{ background: #fff3e0; color: #e65100; }}
    .pill.low  {{ background: #e8f5e9; color: #2e7d32; }}
    .pill.none {{ background: #eeeeee; color: #555; }}

    .pri-card {{
        padding: 10px 14px;
        border-radius: 8px;
        margin: 4px 0;
        background: white;
        border-left: 5px solid #ccc;
        box-shadow: 0 1px 3px rgba(0,0,0,0.06);
    }}
    .pri-card.high {{ border-left-color: #c62828; background: #fff8f8; }}
    .pri-card.med  {{ border-left-color: #e65100; background: #fffaf5; }}
    .pri-card.low  {{ border-left-color: #2e7d32; background: #f6fff6; }}
    .pri-card .plabel {{ font-size: 0.72rem; font-weight: 700; color: #555; }}
    .pri-card .pval   {{ font-size: 1rem;   font-weight: 800; }}
    .pri-card.high .pval {{ color: #c62828; }}
    .pri-card.med  .pval {{ color: #e65100; }}
    .pri-card.low  .pval {{ color: #2e7d32; }}
    .pri-card .psub  {{ font-size: 0.7rem; color: #888; }}

    .ov-meta {{
        background: #e8f4f4;
        border-radius: 6px;
        padding: 8px 14px;
        font-size: 0.82rem;
        color: #2a5a5a;
        margin-bottom: 10px;
    }}

    .arc-box {{
        background: #fff8e1;
        border: 1px solid {AMBER};
        border-radius: 8px;
        padding: 10px 14px;
        font-size: 0.85rem;
        color: #7a5800;
        margin: 6px 0;
    }}
    .arc-done {{
        background: #e8f5e9;
        border: 1px solid #2e7d32;
        border-radius: 8px;
        padding: 8px 14px;
        font-size: 0.85rem;
        color: #1b5e20;
        margin: 6px 0;
    }}

    .stTabs [data-baseweb="tab-list"] {{ gap: 4px; }}
    .stTabs [data-baseweb="tab"] {{
        background: #e0ecec; border-radius: 6px 6px 0 0;
        color: {TEAL}; font-weight: 600;
    }}
    .stTabs [aria-selected="true"] {{ background: {TEAL} !important; color: white !important; }}

    .prog-badge {{
        display: inline-block;
        background: {TEAL};
        color: white;
        padding: 2px 9px;
        border-radius: 10px;
        font-size: 0.75rem;
        font-weight: 700;
    }}
    .prog-badge.jp {{ background: #1565c0; }}
    .prog-badge.py {{ background: #6a1b9a; }}
    .prog-badge.sy {{ background: {TEAL};  }}
    .prog-badge.arch {{ background: #757575; }}
</style>
""", unsafe_allow_html=True)

# ── Supabase (sensory_overviews lives in the default project) ─────────────────
@st.cache_resource
def get_supabase():
    return create_client(st.secrets["SUPABASE_URL"], st.secrets["SUPABASE_KEY"])

supabase = get_supabase()

# ── Constants ──────────────────────────────────────────────────────────────────
SECTION_ICONS = {
    "Body Awareness (Interoception)": "🫀",
    "Visual":                         "👁️",
    "Sense of Body in Space":         "🗺️",
    "Auditory":                       "👂",
    "Auditory Processing":            "🔊",
    "Smell":                          "👃",
    "Touch and Textures":             "🤚",
    "Kinaesthetic":                   "🏃",
}

PROGRAM_LABELS = {"JP": "Junior Primary", "PY": "Primary Years", "SY": "Senior Years"}

# ── Helpers ────────────────────────────────────────────────────────────────────
def priority_info(pct):
    if pct >= 60:
        return "HIGH PRIORITY",   "high", "🔴"
    elif pct >= 30:
        return "MEDIUM PRIORITY", "med",  "🟡"
    return "LOW PRIORITY",        "low",  "🟢"


def pri_card(section_name, info, icon=""):
    return f"""
    <div class="pri-card {info['css']}">
        <div class="plabel">{icon} {section_name}</div>
        <div class="pval">{info['emoji']} {info['priority']}</div>
        <div class="psub">{info['checked']}/{info['total']} ({info['pct']}%)</div>
    </div>"""


def prog_badge(program, archived=False):
    if archived:
        return '<span class="prog-badge arch">Archived</span>'
    css = program.lower() if program in ("JP", "PY", "SY") else ""
    return f'<span class="prog-badge {css}">{program or "—"}</span>'


def student_label(reg):
    """Initials · Program · Yr — the standard display used across the portal."""
    bits = [reg["initials"]]
    if reg.get("program"):
        bits.append(reg["program"])
    if reg.get("year_level") is not None and str(reg.get("year_level")).strip() != "":
        bits.append(f"Yr {reg['year_level']}")
    return " · ".join(bits)


def go_to_sensory_overview(initials=None):
    """Jump to the Sensory Overview page inside the portal, passing the student."""
    if initials:
        st.session_state["sensory_prefill_initials"] = initials
    st.switch_page(SENSORY_OVERVIEW_PAGE)


# ── Data loaders ───────────────────────────────────────────────────────────────
def load_all_overviews():
    try:
        res = supabase.table("sensory_overviews") \
            .select("id,student_id,student_name,overview_date,review_date,completed_by,priority_results,responses") \
            .order("overview_date", desc=True).execute()
        return res.data or []
    except Exception as e:
        st.error(f"Could not load overviews: {e}")
        return []


@st.cache_data(ttl=60, show_spinner=False)
def load_register():
    """All register students (active + archived). Returns (rows, error)."""
    if not ARCHIVE_OK:
        return [], ARCHIVE_ERR
    try:
        rows = fetch_register(None) or []
        clean = []
        for r in rows:
            ini = (r.get("initials") or "").strip().upper()
            if not ini:
                continue
            clean.append({**r, "initials": ini})
        return clean, None
    except Exception as e:  # noqa: BLE001
        return [], str(e)


def match_to_register(overview, register):
    """
    Link an overview to a register student. Tries, in order:
    register id, exact full name, exact initials, then derived initials
    (only when that is unambiguous).
    """
    sid  = str(overview.get("student_id") or "").strip()
    name = str(overview.get("student_name") or "").strip()

    if sid:
        for r in register:
            if str(r.get("id")) == sid:
                return r
    if name:
        lname = name.lower()
        for r in register:
            if (r.get("full_name") or "").strip().lower() == lname:
                return r
        for r in register:
            if r["initials"] == name.upper():
                return r
        derived = initials_from_name(name)
        if derived:
            hits = [r for r in register if r["initials"] == derived.upper()]
            active_hits = [r for r in hits if r.get("status") == "active"]
            if len(active_hits) == 1:
                return active_hits[0]
            if len(hits) == 1:
                return hits[0]
    return None


def build_student_list(all_overviews, register):
    """
    Group overviews by register student. Overviews that can't be matched are
    grouped by their typed name and flagged as 'not on register'.
    """
    groups = {}
    for ov in all_overviews:
        reg = match_to_register(ov, register)
        if reg:
            key = f"reg::{reg['initials']}::{reg.get('id')}"
            if key not in groups:
                groups[key] = {
                    "key": key,
                    "register": reg,
                    "initials": reg["initials"],
                    "label": student_label(reg),
                    "program": reg.get("program") or "",
                    "status": "archived" if reg.get("status") == "archived" else "active",
                    "overviews": [],
                }
        else:
            raw = (ov.get("student_name") or "Unknown").strip()
            key = f"raw::{raw.lower()}"
            if key not in groups:
                groups[key] = {
                    "key": key,
                    "register": None,
                    "initials": None,
                    "label": raw,
                    "program": "",
                    "status": "unlinked",
                    "overviews": [],
                }
        groups[key]["overviews"].append(ov)

    result = []
    for g in groups.values():
        g["overviews"].sort(key=lambda o: o.get("overview_date") or "", reverse=True)
        g["latest_overview"] = g["overviews"][0]
        result.append(g)
    return sorted(result, key=lambda x: x["label"])


def parse_priorities(raw):
    if not raw:
        return {}
    try:
        if isinstance(raw, str):
            raw = json.loads(raw)
        out = {}
        for sec, info in raw.items():
            _, css, emoji = priority_info(info.get("pct", 0))
            out[sec] = {**info, "css": css, "emoji": emoji}
        return out
    except Exception:
        return {}


def parse_responses(raw):
    if not raw:
        return {}
    try:
        if isinstance(raw, str):
            return json.loads(raw)
        return raw
    except Exception:
        return {}


def review_status(review_date_str):
    if not review_date_str:
        return "No review date set", "#888"
    try:
        rv = date.fromisoformat(review_date_str)
        delta = (rv - date.today()).days
        if delta < 0:
            return f"⚠️ Review overdue ({abs(delta)} days)", "#c62828"
        elif delta <= 14:
            return f"🟡 Review due soon ({delta} days)", "#e65100"
        else:
            return f"🟢 Review due {rv.strftime('%d %b %Y')}", "#2e7d32"
    except Exception:
        return review_date_str, "#555"


# ── Word report for archiving ─────────────────────────────────────────────────
def build_overview_docx(overview, initials, label):
    """Builds a Word copy of one saved overview (initials only — no full name)."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)

    h = doc.add_heading("Sensory Overview", level=0)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x7A, 0x7A)

    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Light Grid Accent 1"
    for k, v in [
        ("Student", label or initials),
        ("Overview date", overview.get("overview_date") or "—"),
        ("Review date", overview.get("review_date") or "—"),
        ("Completed by", overview.get("completed_by") or "—"),
    ]:
        cells = meta.add_row().cells
        cells[0].text = k
        cells[1].text = str(v)

    pr = parse_priorities(overview.get("priority_results"))
    doc.add_heading("Priority summary", level=1)
    if pr:
        order = {"HIGH PRIORITY": 0, "MEDIUM PRIORITY": 1, "LOW PRIORITY": 2}
        t = doc.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Area", "Priority", "Items checked"
        for sec, info in sorted(pr.items(), key=lambda kv: order.get(kv[1].get("priority"), 9)):
            c = t.add_row().cells
            c[0].text = sec
            c[1].text = (info.get("priority") or "").title()
            c[2].text = f"{info.get('checked', 0)}/{info.get('total', 0)} ({info.get('pct', 0)}%)"
    else:
        doc.add_paragraph("No priority data recorded.")

    responses = parse_responses(overview.get("responses"))
    doc.add_heading("Checked items", level=1)
    any_checked = False
    for sec_name, items_dict in (responses or {}).items():
        checked = [it for it, val in items_dict.items() if val]
        if not checked:
            continue
        any_checked = True
        doc.add_heading(sec_name, level=2)
        for it in checked:
            doc.add_paragraph(it, style="List Bullet")
    if not any_checked:
        doc.add_paragraph("No items were checked in this overview.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def archive_title(initials, ov_date):
    # Same title format the Sensory Overview tool uses, so both tools'
    # copies are recognised as the same document.
    return f"Sensory Overview — {initials} — {ov_date}"


@st.cache_data(ttl=30, show_spinner=False)
def archived_titles(initials):
    if not ARCHIVE_OK or not initials:
        return []
    try:
        rows = fetch_archive(student_initials=initials, doc_type=DOC_TYPE,
                             include_confidential=True) or []
        return [r.get("title") or "" for r in rows]
    except Exception:  # noqa: BLE001
        return []


def render_archive_control(student, overview):
    """Save-to-archive block for one overview, with replace-vs-keep-both choice."""
    ov_id   = overview.get("id")
    ov_date = overview.get("overview_date") or "undated"
    ini     = student["initials"]

    if not ARCHIVE_OK:
        st.caption("🗄️ Archive unavailable — archive_utils could not be loaded.")
        return
    if not ini:
        st.markdown(
            '<div class="arc-box">📁 This overview isn\'t linked to a registered student, '
            "so it can't be filed. Add the student in the <b>Student Register</b> "
            "(or check the name matches), then it can be saved here.</div>",
            unsafe_allow_html=True,
        )
        return

    title = archive_title(ini, ov_date)
    existing = archived_titles(ini)
    if title in existing:
        st.markdown(f'<div class="arc-done">✅ This overview is in {ini}\'s archive folder.</div>',
                    unsafe_allow_html=True)
        return

    st.markdown(f'<div class="arc-box">📁 <b>Not yet in {ini}\'s archive.</b> '
                "Save a copy to their Document Archive folder?</div>",
                unsafe_allow_html=True)

    supersede = False
    if existing:
        mode = st.radio(
            "A Sensory Overview is already in this student's folder:",
            ["Replace previous version", "Keep both"],
            horizontal=True, key=f"arcmode_{ov_id}",
        )
        supersede = (mode == "Replace previous version")

    if st.button(f"🗄️ Save to {ini}'s archive", type="primary", key=f"arcsave_{ov_id}"):
        try:
            file_bytes = build_overview_docx(overview, ini, student["label"])
            saved = archive_document(
                source_tool="Sensory Profiles",
                doc_type=DOC_TYPE,
                title=title,
                file_bytes=file_bytes,
                filename=f"sensory_{ini}_{ov_date}.docx",
                student_initials=ini,
                created_by=st.session_state.get("user_email", "unknown"),
                supersede=supersede,
            )
            if saved:
                archived_titles.clear()
                st.success("Saved to the Document Archive"
                           + (" (previous version moved to restorable Deleted)." if supersede else "."))
                st.rerun()
            else:
                st.error("Could not save to the archive — nothing was filed.")
        except Exception as e:  # noqa: BLE001
            st.error(f"Could not save to the archive: {e}")


# ── Page header ────────────────────────────────────────────────────────────────
st.markdown("<h1>🧠 Sensory Profiles</h1>", unsafe_allow_html=True)
st.caption("View sensory overviews for registered students and file them to the Document Archive.")

register, reg_err = load_register()
if reg_err:
    st.warning(f"Student Register couldn't be loaded, so students are shown by typed name only. ({reg_err})")

all_overviews = load_all_overviews()

if not all_overviews:
    st.info("No overviews have been saved yet. Complete a Sensory Overview first.")
    if st.button("➕ Go to Sensory Overview"):
        go_to_sensory_overview()
    st.stop()

students = build_student_list(all_overviews, register)

# ── Top controls ───────────────────────────────────────────────────────────────
c1, c2, c3, c4 = st.columns([1.3, 1.1, 1.5, 2])
with c1:
    status_filter = st.selectbox(
        "Students",
        ["Current", "Archived", "Not on register", "All"],
    )
with c2:
    program_filter = st.selectbox("Program", ["All", "JP", "PY", "SY"])
with c3:
    overview_filter = st.selectbox(
        "Review status",
        ["Any", "Review overdue", "Review due soon"],
    )
with c4:
    search = st.text_input("Search initials", placeholder="Type to filter…")

if status_filter == "Current":
    students = [s for s in students if s["status"] == "active"]
elif status_filter == "Archived":
    students = [s for s in students if s["status"] == "archived"]
elif status_filter == "Not on register":
    students = [s for s in students if s["status"] == "unlinked"]

if program_filter != "All":
    students = [s for s in students if s["program"] == program_filter]

if search:
    q = search.strip().lower()
    students = [s for s in students if q in s["label"].lower()]


def _review_delta(s):
    rv = s["latest_overview"].get("review_date")
    if not rv:
        return None
    try:
        return (date.fromisoformat(rv) - date.today()).days
    except Exception:
        return None


if overview_filter == "Review overdue":
    students = [s for s in students if (d := _review_delta(s)) is not None and d < 0]
elif overview_filter == "Review due soon":
    students = [s for s in students if (d := _review_delta(s)) is not None and 0 <= d <= 14]

unlinked_total = sum(1 for s in build_student_list(all_overviews, register) if s["status"] == "unlinked")
if unlinked_total and status_filter != "Not on register":
    st.caption(f"🗂️ {unlinked_total} student name{'s' if unlinked_total != 1 else ''} on saved overviews "
               "don't match the Student Register — choose **Not on register** to review them.")

st.markdown(f"**{len(students)} student{'s' if len(students) != 1 else ''}** matching filters")
st.markdown("---")

if not students:
    st.info("No students match the current filters.")
    st.stop()

# ── Student list ───────────────────────────────────────────────────────────────
for student in students:
    label    = student["label"]
    ini      = student["initials"]
    overview = student["latest_overview"]
    all_ovs  = student["overviews"]

    pills_html = ""
    pr = parse_priorities(overview.get("priority_results"))
    for sec, info in pr.items():
        icon = SECTION_ICONS.get(sec, "")
        pills_html += f'<span class="pill {info["css"]}">{icon} {sec.split("(")[0].strip()}</span> '

    has_high = any(i.get("priority") == "HIGH PRIORITY" for i in pr.values())
    ov_date  = overview.get("overview_date", "—")
    by       = overview.get("completed_by", "—")
    tag = "📦 " if student["status"] == "archived" else ("🗂️ " if student["status"] == "unlinked" else "")

    with st.expander(
        f"{'🔴 ' if has_high else ''}{tag}{label}  ·  Last overview: {ov_date}  ·  {by}",
        expanded=False,
    ):
        mc1, mc2 = st.columns([3, 1])
        with mc1:
            badge = prog_badge(student["program"], archived=(student["status"] == "archived"))
            sub = ("Not on Student Register" if student["status"] == "unlinked"
                   else PROGRAM_LABELS.get(student["program"], ""))
            st.markdown(
                f'<div class="student-card">'
                f'<div class="sname">{label} &nbsp;{badge}</div>'
                f'<div class="smeta">{sub}</div>'
                f'<div style="margin-top:6px;">{pills_html}</div>'
                f'</div>',
                unsafe_allow_html=True,
            )
        with mc2:
            st.markdown("&nbsp;")
            if st.button("➕ New Overview", key=f"new_{student['key']}", use_container_width=True):
                go_to_sensory_overview(ini)

        rev_label, rev_colour = review_status(overview.get("review_date"))
        st.markdown(
            f'<div class="ov-meta">'
            f'📅 Overview: <strong>{ov_date}</strong> &nbsp;·&nbsp; '
            f'Completed by: <strong>{by}</strong> &nbsp;·&nbsp; '
            f'<span style="color:{rev_colour};font-weight:700;">{rev_label}</span>'
            f'</div>',
            unsafe_allow_html=True,
        )

        # Archive status for the latest overview
        render_archive_control(student, overview)

        tab_summary, tab_detail, tab_history = st.tabs(
            ["📊 Priority Summary", "✅ Checked Items", "🕓 History"]
        )

        with tab_summary:
            if not pr:
                st.warning("No priority data available.")
            else:
                for heading, level in [("🔴 High Priority", "HIGH PRIORITY"),
                                       ("🟡 Medium Priority", "MEDIUM PRIORITY"),
                                       ("🟢 Low Priority", "LOW PRIORITY")]:
                    secs = [(s, i) for s, i in pr.items() if i.get("priority") == level]
                    if not secs:
                        continue
                    st.markdown(f"##### {heading}")
                    cols = st.columns(min(len(secs), 4))
                    for idx, (sec, info) in enumerate(secs):
                        with cols[idx % 4]:
                            st.markdown(pri_card(sec, info, SECTION_ICONS.get(sec, "")), unsafe_allow_html=True)

        with tab_detail:
            responses = parse_responses(overview.get("responses"))
            if not responses:
                st.info("No response detail available.")
            else:
                any_checked = False
                for sec_name, items_dict in responses.items():
                    icon = SECTION_ICONS.get(sec_name, "")
                    checked_items = [it for it, val in items_dict.items() if val]
                    if not checked_items:
                        continue
                    any_checked = True
                    _, css, emoji = priority_info(
                        round(len(checked_items) / max(len(items_dict), 1) * 100, 1)
                    )
                    st.markdown(f"**{icon} {sec_name}** {emoji}")
                    ic1, ic2 = st.columns(2)
                    for i, item in enumerate(checked_items):
                        with [ic1, ic2][i % 2]:
                            st.markdown(f"✅ {item}")
                    st.markdown("---")
                if not any_checked:
                    st.info("No items were checked in this overview.")

        with tab_history:
            if len(all_ovs) <= 1:
                st.info("Only one overview on file — history will appear here as more are added.")
            else:
                for ov in all_ovs[1:]:
                    o_date = ov.get("overview_date", "—")
                    o_by   = ov.get("completed_by", "—")
                    o_rev  = ov.get("review_date", "—")
                    with st.expander(f"📄 {o_date}  ·  {o_by}", expanded=False):
                        st.caption(f"Review date: {o_rev}")
                        hist_pr = parse_priorities(ov.get("priority_results"))
                        if hist_pr:
                            hc = st.columns(4)
                            for idx, (sec, info) in enumerate(hist_pr.items()):
                                with hc[idx % 4]:
                                    st.markdown(
                                        pri_card(sec, info, SECTION_ICONS.get(sec, "")),
                                        unsafe_allow_html=True,
                                    )
                        render_archive_control(student, ov)

        st.markdown("---")
        fa1, fa2 = st.columns(2)
        with fa1:
            if st.button("🔄 Start updated overview", key=f"upd_{student['key']}", use_container_width=True):
                go_to_sensory_overview(ini)
        with fa2:
            st.page_link(SENSORY_OVERVIEW_PAGE, label="📋 Go to Sensory Overview", icon=None)
