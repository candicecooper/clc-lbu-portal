"""
72_Student_Snapshot.py — CLC-LBU Portal
=======================================

One screen per student for case meetings: SMARTAR goals, behaviour data
(Behaviour Support Tool + Scatter Plot Tracker) and the Transition
Timetable, plus the student's filed documents — without opening each tool
and finding the student again.

READ-ONLY. This page never writes to any table, so it can't change or
break data in the tools it reads from.

Where the data comes from
  • Student list ........ Student Register (archive_utils.fetch_register)
  • Goals ............... smartar_goals / smartar_tracking   (SMARTAR Goal Builder, SGoal Quick Track)
  • Incidents ........... incidents / critical_incidents     (Behaviour Support Tool, Quick Incident Log)
  • Absences ............ student_attendance                  (Behaviour Support Tool)
  • Scatter plot ........ scatter_students / scatter_weekly_data (Scatter Plot Tracker)
  • Transition .......... transition_timetables               (Transition Timetables, via transition_utils)
  • Documents ........... document_archive                    (Document Archive)

Privacy: the page shows INITIALS ONLY. Behaviour Support rows hold full
names, but they are only used to match the student and are never shown.
The Word summary is initials-only too.

Access: staff password (staff_auth). To restrict it to leadership, swap
the two auth lines below for confidential_auth.require_confidential_auth.
"""

from __future__ import annotations

import io
import json
import os
import re
import time as _time
from collections import Counter
from datetime import date, datetime, timedelta
from urllib.parse import quote
from zoneinfo import ZoneInfo

import pandas as pd
import plotly.graph_objects as go
import streamlit as st

st.set_page_config(page_title="Student Snapshot", page_icon="🧭", layout="wide")

try:
    from theme import apply_clc_theme
    apply_clc_theme()
except Exception:
    pass

from staff_auth import require_staff_auth  # noqa: E402

require_staff_auth("Student Snapshot", "🧭")

from archive_utils import (  # noqa: E402
    _get_client as _archive_client,
    download_bytes,
    fetch_archive,
    fetch_register,
    initials_from_name,
)

try:
    import transition_utils as tu
except Exception:  # transition tool not deployed
    tu = None

# ──────────────────────────────────────────────────────────────────────────
# Constants
# ──────────────────────────────────────────────────────────────────────────
TEAL, TEAL_LIGHT, AMBER, AMBER_LIGHT = "#1A7A7A", "#e8f2f2", "#C8760A", "#fdf3e3"
RED, GREEN, GREY = "#c0392b", "#1a7a3c", "#6b7c7c"
ADELAIDE_TZ = ZoneInfo("Australia/Adelaide")
PORTAL = "https://clc-lbu-portal.onrender.com"

LEVEL_ORDER = ["EXCEEDING", "ABOVE", "EXPECTED", "BELOW", "NOT YET"]
LEVEL_EMOJI = {"EXCEEDING": "🌟", "ABOVE": "✅", "EXPECTED": "🎯", "BELOW": "⚠️", "NOT YET": "❌"}
LEVEL_SCORE = {"EXCEEDING": 5, "ABOVE": 4, "EXPECTED": 3, "BELOW": 2, "NOT YET": 1}
ON_TRACK = {"EXCEEDING", "ABOVE", "EXPECTED"}
GOAL_COLOURS = [TEAL, AMBER, "#5b6ee1", "#8e44ad", "#2c3e50", "#16a085"]

DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
SESSIONS = ["Morning", "Middle", "Afternoon"]

PERIODS = ["This term", "Last 4 weeks", "Last 10 weeks", "Whole placement"]

TODAY = datetime.now(ADELAIDE_TZ).date()

# ──────────────────────────────────────────────────────────────────────────
# Styling
# ──────────────────────────────────────────────────────────────────────────
st.markdown(f"""
<style>
  .block-container {{ padding-top: 1.4rem; }}
  .snap-head {{
      background: linear-gradient(135deg, {TEAL} 0%, #145f5f 100%);
      color: white; border-radius: 12px; padding: 1rem 1.3rem; margin: .4rem 0 1rem;
  }}
  .snap-head h2 {{ color: white !important; margin: 0; font-size: 1.7rem; letter-spacing: .02em; }}
  .snap-head .sub {{ opacity: .88; font-size: .9rem; margin-top: .2rem; }}
  .snap-head .chip {{
      display:inline-block; background: rgba(255,255,255,.16); border-radius: 999px;
      padding: .12rem .65rem; margin: .45rem .35rem 0 0; font-size: .8rem;
  }}
  .tile {{
      background: white; border: 1px solid #d6e4e4; border-radius: 12px;
      padding: .85rem 1rem; height: 100%;
  }}
  .tile .lbl {{ color: {GREY}; font-size: .75rem; font-weight: 600;
               text-transform: uppercase; letter-spacing: .06em; }}
  .tile .big {{ font-size: 1.6rem; font-weight: 700; color: #1f2d2d; line-height: 1.2; margin-top: .15rem; }}
  .tile .note {{ font-size: .82rem; color: {GREY}; margin-top: .15rem; }}
  .tile.warn {{ border-left: 5px solid {AMBER}; }}
  .tile.bad  {{ border-left: 5px solid {RED}; }}
  .tile.good {{ border-left: 5px solid {GREEN}; }}
  .tile.info {{ border-left: 5px solid {TEAL}; }}
  .goal-card {{
      background: white; border: 1px solid #d6e4e4; border-left: 5px solid {TEAL};
      border-radius: 10px; padding: .75rem 1rem; margin: .6rem 0 .3rem;
  }}
  .goal-card .gtext {{ font-size: .95rem; color: #1f2d2d; }}
  .goal-card .meta {{ font-size: .8rem; color: {GREY}; margin-top: .3rem; }}
  .pill {{ display:inline-block; border-radius: 999px; padding: .05rem .55rem;
           font-size: .75rem; font-weight: 600; color: white; margin-right: .3rem; }}
  .ci-card {{ background: #fff7f5; border-left: 4px solid {RED}; border-radius: 8px;
              padding: .55rem .8rem; margin: .35rem 0; font-size: .87rem; }}
  .day-cell {{ border-radius: 8px; padding: .5rem .4rem; text-align: center;
               font-size: .82rem; min-height: 64px; border: 1px solid #d6e4e4; }}
  .day-cell .dn {{ font-weight: 700; font-size: .78rem; text-transform: uppercase; }}
  .muted {{ color: {GREY}; font-size: .85rem; }}
</style>
""", unsafe_allow_html=True)


# ──────────────────────────────────────────────────────────────────────────
# Supabase clients
# ──────────────────────────────────────────────────────────────────────────
def _secret(*names):
    for n in names:
        v = os.environ.get(n)
        if v:
            return v
        try:
            if n in st.secrets:
                return st.secrets[n]
        except Exception:
            pass
    return None


def _make_main_client():
    """The project the SMARTAR, Behaviour Support and Scatter Plot tools use
    (secrets [supabase] url/key). Falls back to the archive's project."""
    from supabase import create_client
    url = key = None
    try:
        url = st.secrets.get("supabase", {}).get("url")
        key = st.secrets.get("supabase", {}).get("key")
    except Exception:
        pass
    url = url or _secret("BEHAVIOUR_SUPABASE_URL")
    key = key or _secret("BEHAVIOUR_SUPABASE_KEY")
    if url and key:
        return create_client(url, key)
    return _archive_client()


def _make_students_client():
    """The students project (Register, Transition Timetables)."""
    if tu is not None:
        try:
            c = tu.get_students_client()
            if c is not None:
                return c
        except Exception:
            pass
    return _archive_client()


_CLIENTS: dict = {}


def _client(which: str):
    if which not in _CLIENTS:
        _CLIENTS[which] = _make_main_client() if which == "main" else _make_students_client()
    return _CLIENTS[which]


def _q(which: str, fn, retries: int = 3):
    """Run a query with retry + a fresh client (Render free tier drops idle connections)."""
    last = None
    for attempt in range(retries):
        try:
            return fn(_client(which))
        except Exception as exc:  # noqa: BLE001
            last = exc
            _CLIENTS.pop(which, None)
            _time.sleep(0.4 * (attempt + 1))
    raise last


# ──────────────────────────────────────────────────────────────────────────
# Matching a register student (initials) to each tool's own records
# ──────────────────────────────────────────────────────────────────────────
def _name_matches(name, ini: str) -> bool:
    """True if a tool's stored student name belongs to these initials.
    Covers initials-only rows (register sync), prefix names like 'JT (PY)',
    and legacy full-name rows (via the register's own initials rule)."""
    n = str(name or "").strip()
    if not n:
        return False
    up = n.upper()
    if up == ini:
        return True
    if up.startswith(ini) and not up[len(ini):len(ini) + 1].isalpha():
        return True
    return initials_from_name(n) == ini


# ──────────────────────────────────────────────────────────────────────────
# Loaders (cached briefly so moving between tabs is instant)
# ──────────────────────────────────────────────────────────────────────────
@st.cache_data(ttl=60, show_spinner=False)
def load_register(include_archived: bool) -> list[dict]:
    rows = fetch_register(None if include_archived else "active")
    return [r for r in rows if r.get("initials")]


@st.cache_data(ttl=60, show_spinner=False)
def load_smartar(ini: str) -> dict:
    out = {"names": [], "goals": [], "tracking": [], "error": None}
    try:
        students = _q("main", lambda c: c.table("smartar_students")
                      .select("student_name").execute().data or [])
        names = sorted({s["student_name"] for s in students if _name_matches(s.get("student_name"), ini)})
        if ini not in names:
            names.append(ini)  # goals can exist without a student row
        out["names"] = names
        goals = _q("main", lambda c: c.table("smartar_goals").select("*")
                   .in_("student_name", names).order("created_at").execute().data or [])
        out["goals"] = goals
        ids = [g["id"] for g in goals]
        if ids:
            out["tracking"] = _q("main", lambda c: c.table("smartar_tracking")
                                 .select("goal_id, session_date, session_time, performance_level")
                                 .in_("goal_id", ids).order("session_date").execute().data or [])
    except Exception as exc:  # noqa: BLE001
        out["error"] = str(exc)
    return out


@st.cache_data(ttl=60, show_spinner=False)
def load_behaviour(ini: str) -> dict:
    out = {"student_ids": [], "incidents": [], "critical": [], "absent_dates": [],
           "multiple": False, "error": None}
    try:
        rows = _q("main", lambda c: c.table("students")
                  .select("id, name, first_name, last_name, archived").execute().data or [])
        matched = [r for r in rows
                   if (str(r.get("first_name") or "").strip().upper() == ini)
                   or _name_matches(r.get("name"), ini)
                   or _name_matches(f"{r.get('first_name') or ''} {r.get('last_name') or ''}", ini)]
        current = [r for r in matched if not r.get("archived")]
        use = current or matched
        ids = [str(r["id"]) for r in use]
        out["student_ids"] = ids
        out["multiple"] = len(ids) > 1
        if not ids:
            return out
        out["incidents"] = _q("main", lambda c: c.table("incidents")
                              .select("id, student_id, date, time, day_of_week, session, location, "
                                      "behaviour_type, severity, hypothesis_function, hypothesis_item, "
                                      "is_critical, duration_minutes")
                              .in_("student_id", ids).execute().data or [])
        out["critical"] = _q("main", lambda c: c.table("critical_incidents")
                             .select("id, student_id, severity, primary_location, primary_time, "
                                     "primary_behaviour, primary_hypothesis_function, "
                                     "primary_hypothesis_item, status, created_at")
                             .in_("student_id", ids).execute().data or [])
        try:
            att = _q("main", lambda c: c.table("student_attendance")
                     .select("student_id, absent_dates").in_("student_id", ids).execute().data or [])
            dates = set()
            for a in att:
                ad = a.get("absent_dates") or []
                if isinstance(ad, str):
                    ad = json.loads(ad or "[]")
                for d in ad:
                    dates.add(str(d)[:10])
            out["absent_dates"] = sorted(dates)
        except Exception:
            pass  # attendance table optional
    except Exception as exc:  # noqa: BLE001
        out["error"] = str(exc)
    return out


def _loads(v, default):
    if v is None:
        return default
    if isinstance(v, (dict, list)):
        return v
    try:
        return json.loads(v)
    except Exception:
        return default


@st.cache_data(ttl=60, show_spinner=False)
def load_scatter(ini: str) -> dict:
    out = {"key": None, "definitions": {}, "weeks": [], "error": None}
    try:
        rows = _q("main", lambda c: c.table("scatter_students")
                  .select("student_name, behaviour_definitions").execute().data or [])
        match = next((r for r in rows if _name_matches(r.get("student_name"), ini)), None)
        key = match["student_name"] if match else ini
        out["key"] = key
        out["definitions"] = _loads(match.get("behaviour_definitions"), {}) if match else {}
        out["weeks"] = _q("main", lambda c: c.table("scatter_weekly_data")
                          .select("week_start, behaviours, schedule, grid_data")
                          .eq("student_name", key).order("week_start").execute().data or [])
    except Exception as exc:  # noqa: BLE001
        out["error"] = str(exc)
    return out


@st.cache_data(ttl=60, show_spinner=False)
def load_transition(ini: str) -> dict:
    out = {"timetables": [], "error": None}
    if tu is None:
        out["error"] = "Transition Timetables module not found."
        return out
    try:
        rows = _q("students", lambda c: c.table(tu.TABLE_TIMETABLES).select("*")
                  .eq("student_initials", ini).order("updated_at", desc=True).execute().data or [])
        out["timetables"] = rows
    except Exception as exc:  # noqa: BLE001
        out["error"] = str(exc)
    return out


@st.cache_data(ttl=60, show_spinner=False)
def load_documents(ini: str) -> list[dict]:
    return fetch_archive(student_initials=ini, limit=50)


# ──────────────────────────────────────────────────────────────────────────
# Date helpers
# ──────────────────────────────────────────────────────────────────────────
def _d(v) -> date | None:
    if not v:
        return None
    if isinstance(v, date) and not isinstance(v, datetime):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except Exception:
        return None


def _monday(d: date) -> date:
    return d - timedelta(days=d.weekday())


def _fmt(d: date | None, with_year: bool = False) -> str:
    if not d:
        return "—"
    return d.strftime("%-d %b %Y" if with_year else "%-d %b")


def current_term_range(today: date) -> tuple[date, date, str]:
    """Start/end of the current term (or the most recent one in holidays)."""
    if tu is not None:
        terms = sorted(((y, t, s, e) for y, ts in tu.TERM_DATES.items()
                        for t, (s, e) in ts.items()), key=lambda x: x[2])
        chosen = None
        for y, t, s, e in terms:
            if s - timedelta(days=s.weekday()) <= today:
                chosen = (y, t, s, e)
        if chosen:
            y, t, s, e = chosen
            return s, min(e, today), f"Term {t} {y}"
    return today - timedelta(days=70), today, "Last 10 weeks"


def period_range(choice: str, placement_start: date | None) -> tuple[date, date, str]:
    if choice == "This term":
        return current_term_range(TODAY)
    if choice == "Last 4 weeks":
        return TODAY - timedelta(days=27), TODAY, "last 4 weeks"
    if choice == "Last 10 weeks":
        return TODAY - timedelta(days=69), TODAY, "last 10 weeks"
    start = placement_start or (TODAY - timedelta(days=365))
    return start, TODAY, "whole placement"


def in_range(d: date | None, start: date, end: date) -> bool:
    return d is not None and start <= d <= end


def school_days_between(start: date, end: date) -> int:
    n, d = 0, start
    while d <= end:
        if d.weekday() < 5:
            n += 1
        d += timedelta(days=1)
    return n


# ──────────────────────────────────────────────────────────────────────────
# Analysis
# ──────────────────────────────────────────────────────────────────────────
def goal_stats(goal: dict, tracking: list[dict], start: date, end: date) -> dict:
    rows = [r for r in tracking
            if r.get("goal_id") == goal["id"] and in_range(_d(r.get("session_date")), start, end)]
    rows.sort(key=lambda r: (str(r.get("session_date")), SESSIONS.index(r["session_time"])
                             if r.get("session_time") in SESSIONS else 9))
    counts = Counter(r.get("performance_level") for r in rows)
    total = sum(counts[l] for l in LEVEL_ORDER)
    on = sum(counts[l] for l in ON_TRACK)
    trend = "insufficient"
    if len(rows) >= 4:  # same rule as the Goal Builder's Progress Report
        mid = len(rows) // 2
        a = sum(LEVEL_SCORE.get(r.get("performance_level"), 0) for r in rows[:mid]) / mid
        b = sum(LEVEL_SCORE.get(r.get("performance_level"), 0) for r in rows[mid:]) / (len(rows) - mid)
        trend = "improving" if b - a > 0.5 else "declining" if b - a < -0.5 else "stable"
    weekly: dict[date, list[int]] = {}
    for r in rows:
        wk = _monday(_d(r["session_date"]))
        weekly.setdefault(wk, [0, 0])
        weekly[wk][1] += 1
        if r.get("performance_level") in ON_TRACK:
            weekly[wk][0] += 1
    return {
        "rows": rows, "counts": counts, "total": total,
        "pct_on": round(on / total * 100) if total else None,
        "trend": trend,
        "last": _d(rows[-1]["session_date"]) if rows else None,
        "weekly": {k: round(v[0] / v[1] * 100) for k, v in sorted(weekly.items())},
    }


TREND_TXT = {
    "improving": ("📈 Improving", GREEN),
    "declining": ("📉 Declining", RED),
    "stable": ("➡️ Stable", TEAL),
    "insufficient": ("· Not enough data", GREY),
}


def ci_date(ci: dict) -> date | None:
    return _d(ci.get("primary_time")) or _d(ci.get("created_at"))


def top(counter: Counter, n: int = 3) -> list[tuple[str, int]]:
    return [(k, v) for k, v in counter.most_common() if k][:n]


def behaviour_stats(beh: dict, start: date, end: date) -> dict:
    inc = [i for i in beh["incidents"] if in_range(_d(i.get("date")), start, end)]
    span = (end - start).days + 1
    p_start, p_end = start - timedelta(days=span), start - timedelta(days=1)
    prev = [i for i in beh["incidents"] if in_range(_d(i.get("date")), p_start, p_end)]
    crit = [c for c in beh["critical"] if in_range(ci_date(c), start, end)]
    crit_n = max(len(crit), sum(1 for i in inc if i.get("is_critical")))
    sev = [float(i["severity"]) for i in inc if str(i.get("severity") or "").replace(".", "", 1).isdigit()]
    absent = [d for d in beh["absent_dates"] if in_range(_d(d), start, end)]
    func = Counter()
    for i in inc:
        f, it = (i.get("hypothesis_function") or "").strip(), (i.get("hypothesis_item") or "").strip()
        if f or it:
            func[f"{f} {it.lower()}".strip()] += 1
    day = Counter(i.get("day_of_week") or (_d(i.get("date")).strftime("%A") if _d(i.get("date")) else "")
                  for i in inc)
    return {
        "inc": inc, "prev_n": len(prev), "crit": crit, "crit_n": crit_n,
        "avg_sev": round(sum(sev) / len(sev), 1) if sev else None,
        "absent": absent,
        "behaviours": Counter(i.get("behaviour_type") for i in inc),
        "locations": Counter(i.get("location") for i in inc),
        "sessions": Counter(i.get("session") for i in inc),
        "days": day,
        "functions": func,
    }


def scatter_stats(sc: dict, start: date, end: date) -> dict:
    """Totals per (day-of-week, time slot) across the weeks in range."""
    grid: dict[tuple[str, str], int] = {}
    slots: list[str] = []
    behaviours: list[str] = []
    weeks_used = 0
    for w in sc["weeks"]:
        ws = _d(w.get("week_start"))
        if not ws or not (start - timedelta(days=6) <= ws <= end):
            continue
        weeks_used += 1
        sched = _loads(w.get("schedule"), [])
        behs = _loads(w.get("behaviours"), [])
        data = _loads(w.get("grid_data"), {})
        for b in behs:
            if b not in behaviours:
                behaviours.append(b)
        for slot in sched:
            t = slot.get("time")
            if t and t not in slots:
                slots.append(t)
        for i in range(5):
            d = ws + timedelta(days=i)
            if not in_range(d, start, end):
                continue
            for slot in sched:
                t = slot.get("time")
                for b_idx in range(max(1, len(behs))):
                    v = data.get(f"{d}_{t}_{b_idx}", 0) or 0
                    if v:
                        k = (DAY_NAMES[i], t)
                        grid[k] = grid.get(k, 0) + int(v)
    return {"grid": grid, "slots": slots, "behaviours": behaviours,
            "weeks": weeks_used, "total": sum(grid.values())}


def pick_timetable(tts: list[dict]) -> dict | None:
    """Most useful timetable: live (published) first, then draft, then complete."""
    if not tts:
        return None
    for want in ("published", "active", "draft", "complete"):
        for t in tts:
            if t.get("status") == want:
                return t
    return tts[0]


def tt_view(tt: dict) -> dict:
    if tu is not None and tu.is_live(tt):
        return tu.live_view(tt) or tt
    return tt


def tt_weeks(view: dict) -> list[dict]:
    weeks = _loads(view.get("weeks"), [])
    return tu.sort_weeks(weeks) if (tu is not None and weeks) else weeks


def tt_status(tt: dict) -> str:
    if tu is not None:
        return tu.status_label(tt)
    return str(tt.get("status") or "—").title()


def week_counts(week: dict) -> dict:
    if tu is not None:
        return tu.week_summary(week)
    out = Counter(d.get("loc") for d in week.get("days", []))
    return {"At CLC": out.get("CLC", 0), "Home school – part day": out.get("HOME", 0),
            "Home school – full day": 0, "Other": out.get("OTHER", 0)}


def day_text(day: dict, home_short: str) -> str:
    if tu is not None:
        return tu.short_day_text(day, home_short or "Home school")
    return f"{day.get('loc', '')} {day.get('start', '')}–{day.get('end', '')}"


# ──────────────────────────────────────────────────────────────────────────
# Small render helpers
# ──────────────────────────────────────────────────────────────────────────
def tile(col, label: str, big: str, note: str = "", kind: str = "info"):
    col.markdown(
        f'<div class="tile {kind}"><div class="lbl">{label}</div>'
        f'<div class="big">{big}</div><div class="note">{note}</div></div>',
        unsafe_allow_html=True,
    )


def _esc(s) -> str:
    return (str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def _short(s, n=110) -> str:
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[: n - 1] + "…"


def top_table(title: str, counter: Counter, total: int, n: int = 5):
    items = top(counter, n)
    st.markdown(f"**{title}**")
    if not items:
        st.markdown('<span class="muted">No data</span>', unsafe_allow_html=True)
        return
    df = pd.DataFrame([{"": k, "Count": v, "%": f"{round(v / total * 100)}%" if total else ""}
                       for k, v in items])
    st.dataframe(df, hide_index=True, use_container_width=True)


def plot_layout(fig, height=300):
    fig.update_layout(
        height=height, margin=dict(l=10, r=10, t=30, b=10),
        plot_bgcolor="white", paper_bgcolor="white",
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
        font=dict(size=12),
    )
    fig.update_xaxes(showgrid=False)
    fig.update_yaxes(gridcolor="#eef3f3")
    return fig


# ──────────────────────────────────────────────────────────────────────────
# Word summary (initials only)
# ──────────────────────────────────────────────────────────────────────────
def build_summary_docx(stu: dict, period_lbl: str, start: date, end: date,
                       goals: list[tuple[dict, dict]], bs: dict, sc_stats: dict,
                       tt: dict | None) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    h = doc.add_heading(f"Student Snapshot — {stu['initials']}", level=1)
    h.runs[0].font.color.rgb = RGBColor(0x1A, 0x7A, 0x7A)
    doc.add_paragraph(
        f"Program: {stu.get('program') or '—'}   ·   Year level: {stu.get('year_level') or '—'}   ·   "
        f"Placement: {_fmt(_d(stu.get('start_date')), True)} → {_fmt(_d(stu.get('planned_exit_date')), True)}"
    )
    doc.add_paragraph(f"Data period: {period_lbl} ({_fmt(start, True)} – {_fmt(end, True)})   ·   "
                      f"Prepared {_fmt(TODAY, True)}")

    doc.add_heading("SMARTAR goals", level=2)
    if not goals:
        doc.add_paragraph("No active SMARTAR goals.")
    else:
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for c, txt in zip(t.rows[0].cells, ["Goal", "Expected or better", "Entries", "Trend"]):
            c.text = txt
        for g, s in goals:
            r = t.add_row().cells
            r[0].text = g.get("smartar_goal") or "—"
            r[1].text = f"{s['pct_on']}%" if s["pct_on"] is not None else "—"
            r[2].text = str(s["total"])
            r[3].text = TREND_TXT[s["trend"]][0].split(" ", 1)[-1]

    doc.add_heading("Behaviour", level=2)
    n = len(bs["inc"])
    doc.add_paragraph(
        f"Incidents: {n} (previous equal period: {bs['prev_n']})   ·   Critical incidents: {bs['crit_n']}   ·   "
        f"Average severity: {bs['avg_sev'] if bs['avg_sev'] is not None else '—'}   ·   "
        f"Absences: {len(bs['absent'])}"
    )
    for title, counter in (("Most frequent behaviours", bs["behaviours"]),
                           ("Locations", bs["locations"]),
                           ("Sessions", bs["sessions"]),
                           ("Hypothesised function", bs["functions"])):
        items = top(counter, 3)
        if items:
            doc.add_paragraph(f"{title}: " + "; ".join(f"{k} ({v})" for k, v in items))
    if sc_stats["total"]:
        peak = max(sc_stats["grid"].items(), key=lambda kv: kv[1])
        doc.add_paragraph(
            f"Scatter plot ({', '.join(sc_stats['behaviours']) or 'target behaviour'}): "
            f"{sc_stats['total']} recorded across {sc_stats['weeks']} week(s); "
            f"peak {peak[0][0]} {peak[0][1]} ({peak[1]})."
        )

    doc.add_heading("Transition", level=2)
    if not tt:
        doc.add_paragraph("No transition timetable on file.")
    else:
        view = tt_view(tt)
        weeks = tt_weeks(view)
        status = re.sub(r"^[^\w]+", "", tt_status(tt))
        doc.add_paragraph(
            f"Status: {status}   ·   Home school: {view.get('home_school') or '—'}   ·   "
            f"Last CLC day: {_fmt(tu.last_clc_day(weeks), True) if (tu and weeks) else '—'}"
        )
        if weeks:
            t = doc.add_table(rows=1, cols=5)
            t.style = "Light Grid Accent 1"
            for c, txt in zip(t.rows[0].cells, ["Week", "At CLC", "Home – part", "Home – full", "Other"]):
                c.text = txt
            for w in weeks:
                wc = week_counts(w)
                r = t.add_row().cells
                r[0].text = f"T{w['term']} W{w['week']} {w['year']}"
                r[1].text = str(wc["At CLC"])
                r[2].text = str(wc["Home school – part day"])
                r[3].text = str(wc["Home school – full day"])
                r[4].text = str(wc["Other"])

    doc.add_paragraph()
    p = doc.add_paragraph("Initials only — no names, EDID or DOB. Generated from the CLC-LBU portal.")
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor(0x6B, 0x7C, 0x7C)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════
# PAGE
# ══════════════════════════════════════════════════════════════════════════
top_l, top_r = st.columns([4, 1])
with top_l:
    st.markdown(
        f'<p style="margin:0;color:#5a8a8a;font-size:.78rem;font-weight:600;'
        f'text-transform:uppercase;letter-spacing:.07em;">🧭 CLC Learning &amp; Behaviour Unit</p>'
        f'<h2 style="margin:.1rem 0 0;color:{TEAL} !important;">Student Snapshot</h2>'
        f'<p class="muted" style="margin:.1rem 0 .6rem;">Goals, behaviour and transition for one student, '
        f'on one screen — read-only.</p>',
        unsafe_allow_html=True,
    )
with top_r:
    st.write("")
    if st.button("🔄 Refresh data", use_container_width=True):
        st.cache_data.clear()
        st.rerun()

# ── Student picker ────────────────────────────────────────────────────────
c1, c2, c3, c4 = st.columns([1, 2, 1.4, 1])
with c4:
    st.write("")
    show_archived = st.toggle("Include ended placements", value=False)
register = load_register(show_archived)
if not register:
    st.warning("No students on the Student Register yet. Add them in the Student Register first.")
    st.stop()

programs = sorted({r.get("program") or "—" for r in register})
with c1:
    prog = st.selectbox("Program", ["All"] + programs, key="snap_prog")
pool = [r for r in register if prog == "All" or (r.get("program") or "—") == prog]
if not pool:
    st.info("No students in this program.")
    st.stop()

labels = {}
for r in pool:
    lbl = f"{r['initials']} · {r.get('program') or '—'}"
    if r.get("status") and r.get("status") != "active":
        lbl += " (ended)"
    labels[lbl] = r

# Remember the student from the URL (?student=JT) so a link can open straight to them
qp_student = (st.query_params.get("student") or "").strip().upper()
label_list = list(labels)
default_idx = next((i for i, l in enumerate(label_list) if labels[l]["initials"] == qp_student), 0)
with c2:
    chosen = st.selectbox("Student", label_list, index=default_idx, key="snap_student")
with c3:
    period = st.selectbox("Data period", PERIODS, key="snap_period")

stu = labels[chosen]
ini = stu["initials"].strip().upper()
try:
    if st.query_params.get("student") != ini:
        st.query_params["student"] = ini
except Exception:
    pass

p_start_d = _d(stu.get("start_date"))
start, end, period_lbl = period_range(period, p_start_d)
if p_start_d and start < p_start_d and period != "Whole placement":
    start = max(start, p_start_d)

# ── Load everything for this student ──────────────────────────────────────
with st.spinner(f"Loading {ini}…"):
    sm = load_smartar(ini)
    beh = load_behaviour(ini)
    sc = load_scatter(ini)
    tr = load_transition(ini)

active_goals = [g for g in sm["goals"] if g.get("active")]
past_goals = [g for g in sm["goals"] if not g.get("active")]
goal_rows = [(g, goal_stats(g, sm["tracking"], start, end)) for g in active_goals]
bs = behaviour_stats(beh, start, end)
scs = scatter_stats(sc, start, end)
tt = pick_timetable(tr["timetables"])

# ── Student header ────────────────────────────────────────────────────────
exit_d = _d(stu.get("planned_exit_date"))
days_in = (TODAY - p_start_d).days if p_start_d else None
chips = [
    f"Program {stu.get('program') or '—'}",
    f"Year {stu.get('year_level')}" if stu.get("year_level") not in (None, "") else None,
    f"Started {_fmt(p_start_d, True)}" if p_start_d else "Start date not set",
    f"{days_in // 7} weeks at CLC" if days_in is not None and days_in >= 0 else None,
    f"Planned exit {_fmt(exit_d, True)}" if exit_d else None,
    f"{(exit_d - TODAY).days} days to go" if exit_d and exit_d >= TODAY else None,
]
st.markdown(
    f'<div class="snap-head"><h2>{_esc(ini)}</h2>'
    f'<div class="sub">Showing {period_lbl}: {_fmt(start, True)} – {_fmt(end, True)}</div>'
    + "".join(f'<span class="chip">{_esc(c)}</span>' for c in chips if c)
    + "</div>",
    unsafe_allow_html=True,
)

# ── At a glance ───────────────────────────────────────────────────────────
t1, t2, t3, t4 = st.columns(4)

tracked = [s for _, s in goal_rows if s["pct_on"] is not None]
if not active_goals:
    tile(t1, "SMARTAR goals", "None active", "Set goals in the Goal Builder", "warn")
else:
    on_track = sum(1 for s in tracked if s["pct_on"] >= 60)
    declining = sum(1 for _, s in goal_rows if s["trend"] == "declining")
    note = f"{on_track} of {len(active_goals)} at Expected+ most of the time"
    if declining:
        note += f" · {declining} declining"
    avg = round(sum(s["pct_on"] for s in tracked) / len(tracked)) if tracked else None
    kind = "bad" if declining else ("good" if avg is not None and avg >= 60 else "warn")
    tile(t1, "SMARTAR goals", f"{avg}% Expected+" if avg is not None else "No tracking", note, kind)

n_inc = len(bs["inc"])
delta = n_inc - bs["prev_n"]
arrow = "▲" if delta > 0 else "▼" if delta < 0 else "="
kind = "bad" if bs["crit_n"] else ("warn" if delta > 0 else "good")
tile(t2, "Incidents", f"{n_inc}",
     f"{arrow} {abs(delta)} vs previous period · {bs['crit_n']} critical", kind)

school_days = school_days_between(start, end)
n_abs = len(bs["absent"])
tile(t3, "Absences", f"{n_abs}", f"of {school_days} school days in this period",
     "warn" if school_days and n_abs / school_days > 0.1 else "info")

if tt:
    view = tt_view(tt)
    weeks = tt_weeks(view)
    lcd = tu.last_clc_day(weeks) if (tu and weeks) else None
    tile(t4, "Transition", tt_status(tt).split(" ", 1)[-1],
         f"Last CLC day {_fmt(lcd, True)}" if lcd else (view.get("home_school") or ""),
         "good" if tt.get("status") in ("published", "active") else "warn")
else:
    tile(t4, "Transition", "No timetable", "Create one in Transition Timetables", "info")

# ── Actions ───────────────────────────────────────────────────────────────
st.write("")
a1, a2, a3, a4, a5, a6 = st.columns(6)
a1.link_button("🎯 Goal Builder", f"{PORTAL}/SMARTAR_Goal_Builder", use_container_width=True)
a2.link_button("⚡ Quick Track", f"{PORTAL}/SGoal_Quick_Track", use_container_width=True)
sid = beh["student_ids"][0] if beh["student_ids"] else ""
a3.link_button("📊 Behaviour Support",
               f"{PORTAL}/Behaviour_Support" + (f"?p=student_analysis&sid={quote(sid)}" if sid else ""),
               use_container_width=True)
a4.link_button("📍 Scatter Plot", f"{PORTAL}/scatter_plot_tracker", use_container_width=True)
a5.link_button("🗓️ Transition", f"{PORTAL}/Transition_Timetables", use_container_width=True)
try:
    summary = build_summary_docx(stu, period_lbl, start, end, goal_rows, bs, scs, tt)
    a6.download_button("⬇️ Meeting summary", summary,
                       file_name=f"Student_Snapshot_{ini}_{TODAY:%Y%m%d}.docx",
                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                       use_container_width=True, type="primary")
except Exception as exc:  # noqa: BLE001
    a6.caption(f"Summary unavailable: {exc}")

st.write("")
tab_g, tab_b, tab_t, tab_d = st.tabs(["🎯 Goals", "📊 Behaviour", "🗓️ Transition", "📁 Documents"])

# ══════════════════════════════════════════════════════════════════════════
# GOALS
# ══════════════════════════════════════════════════════════════════════════
with tab_g:
    if sm["error"]:
        st.error(f"Couldn't load SMARTAR data: {sm['error']}")
    elif not active_goals:
        st.info(f"{ini} has no active SMARTAR goals.")
    else:
        # Weekly % Expected-or-better, one line per goal
        fig = go.Figure()
        any_points = False
        for i, (g, s) in enumerate(goal_rows):
            if not s["weekly"]:
                continue
            any_points = True
            xs = list(s["weekly"].keys())
            fig.add_trace(go.Scatter(
                x=xs, y=list(s["weekly"].values()), mode="lines+markers",
                name=f"Goal {i + 1}", line=dict(color=GOAL_COLOURS[i % len(GOAL_COLOURS)], width=2.5),
                hovertemplate="Week of %{x|%-d %b}<br>%{y}% Expected or better<extra>Goal " + str(i + 1) + "</extra>",
            ))
        if any_points:
            fig.add_hline(y=60, line_dash="dot", line_color=GREY,
                          annotation_text="60%", annotation_position="right")
            fig.update_yaxes(range=[0, 105], ticksuffix="%")
            fig.update_layout(title=dict(text="Sessions at Expected or better, by week", font=dict(size=14)))
            st.plotly_chart(plot_layout(fig, 300), use_container_width=True)
        else:
            st.info("No goal tracking recorded in this period.")

        for i, (g, s) in enumerate(goal_rows):
            colour = GOAL_COLOURS[i % len(GOAL_COLOURS)]
            trend_txt, trend_col = TREND_TXT[s["trend"]]
            pct = f"{s['pct_on']}%" if s["pct_on"] is not None else "—"
            meta = [
                g.get("domain"),
                g.get("tracking_period") or "All Day",
                f"Started {_fmt(_d(g.get('start_date') or g.get('created_at')), True)}",
                f"Last tracked {_fmt(s['last'], True)}" if s["last"] else "Not tracked this period",
            ]
            levels = " ".join(
                f'<span class="pill" style="background:{c}">{LEVEL_EMOJI[l]} {l} {round(s["counts"][l] / s["total"] * 100)}%</span>'
                for l, c in (("EXCEEDING", "#27ae60"), ("ABOVE", "#2ecc71"),
                             ("EXPECTED", "#2980b9"), ("NOT YET", "#c0392b"))
                if s["total"] and s["counts"][l]
            )
            st.markdown(
                f'<div class="goal-card" style="border-left-color:{colour}">'
                f'<div style="display:flex;justify-content:space-between;gap:1rem;flex-wrap:wrap;">'
                f'<b style="color:{colour}">Goal {i + 1}</b>'
                f'<span><b>{pct}</b> Expected+ · {s["total"]} entries · '
                f'<span style="color:{trend_col};font-weight:600">{trend_txt}</span></span></div>'
                f'<div class="gtext" style="margin-top:.3rem">{_esc(g.get("smartar_goal") or "—")}</div>'
                f'<div class="meta">{" · ".join(_esc(m) for m in meta if m)}</div>'
                f'<div style="margin-top:.4rem">{levels}</div></div>',
                unsafe_allow_html=True,
            )
            with st.expander("Goal detail (exclusion goal, prompts, GAS levels)"):
                if g.get("exclusion_goal"):
                    st.markdown(f"**Exclusion goal:** {g['exclusion_goal']}")
                bits = [f"**{lbl}:** {g[k]}" for lbl, k in (
                    ("Skill", "skill"), ("Grouping", "grouping"), ("Task", "task_type"),
                    ("Setting", "setting"), ("Prompts", "prompts"), ("Rate", "rate")) if g.get(k)]
                if bits:
                    st.markdown("  \n".join(bits))
                gas = [(l, g.get(f"gas_{l.lower().replace(' ', '_')}")) for l in LEVEL_ORDER]
                gas = [(l, v) for l, v in gas if v]
                if gas:
                    st.markdown("**GAS levels**")
                    for l, v in gas:
                        st.markdown(f"{LEVEL_EMOJI[l]} **{l}** — {v}")

        if past_goals:
            with st.expander(f"Previous goals ({len(past_goals)})"):
                for g in reversed(past_goals):
                    end_d = _d(g.get("archived_at"))
                    st.markdown(
                        f"- {_short(g.get('smartar_goal'), 140)}  \n"
                        f"  <span class='muted'>{_fmt(_d(g.get('start_date')), True)} → {_fmt(end_d, True)}"
                        f"{' · ' + _esc(g.get('archive_reason')) if g.get('archive_reason') else ''}</span>",
                        unsafe_allow_html=True,
                    )

# ══════════════════════════════════════════════════════════════════════════
# BEHAVIOUR
# ══════════════════════════════════════════════════════════════════════════
with tab_b:
    if beh["error"]:
        st.error(f"Couldn't load Behaviour Support data: {beh['error']}")
    elif not beh["student_ids"]:
        st.info(f"{ini} isn't set up in the Behaviour Support Tool yet.")
    else:
        if beh["multiple"]:
            st.caption(f"ℹ️ {len(beh['student_ids'])} Behaviour Support records match {ini}; "
                       "their data is combined here.")
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Incidents", n_inc, delta=f"{delta:+d} vs previous", delta_color="inverse")
        m2.metric("Critical incidents", bs["crit_n"])
        m3.metric("Average severity", bs["avg_sev"] if bs["avg_sev"] is not None else "—",
                  help="1 = minor/coping · 3 = elevated · 5 = crisis")
        m4.metric("Absences", n_abs)

        if bs["inc"]:
            # Weekly incidents, critical stacked on top
            wk = Counter()
            wk_c = Counter()
            for i in bs["inc"]:
                m = _monday(_d(i["date"]))
                (wk_c if i.get("is_critical") else wk)[m] += 1
            mondays = []
            m = _monday(start)
            while m <= end:
                mondays.append(m)
                m += timedelta(days=7)
            fig = go.Figure()
            fig.add_bar(x=mondays, y=[wk[m] for m in mondays], name="Incidents", marker_color=TEAL)
            fig.add_bar(x=mondays, y=[wk_c[m] for m in mondays], name="Critical", marker_color=RED)
            fig.update_layout(barmode="stack",
                              title=dict(text="Incidents by week", font=dict(size=14)))
            fig.update_xaxes(tickformat="%-d %b")
            st.plotly_chart(plot_layout(fig, 280), use_container_width=True)

            p1, p2, p3 = st.columns(3)
            with p1:
                top_table("Behaviours", bs["behaviours"], n_inc)
            with p2:
                top_table("Locations", bs["locations"], n_inc)
            with p3:
                top_table("Hypothesised function", bs["functions"], n_inc)

            # When — session × day heat strip
            st.markdown("**When incidents happen**")
            z = [[sum(1 for i in bs["inc"]
                      if i.get("session") == s
                      and (i.get("day_of_week") or (_d(i["date"]).strftime("%A"))) == d)
                  for d in DAY_NAMES] for s in SESSIONS]
            hm = go.Figure(go.Heatmap(
                z=z, x=[d[:3] for d in DAY_NAMES], y=SESSIONS,
                colorscale=[[0, "#f4f8f8"], [0.5, "#e8a54b"], [1, RED]],
                showscale=False, text=z, texttemplate="%{text}",
                hovertemplate="%{y}, %{x}: %{z}<extra></extra>",
            ))
            hm.update_yaxes(autorange="reversed")
            st.plotly_chart(plot_layout(hm, 200), use_container_width=True)
        else:
            st.success("No incidents recorded in this period.")

        # Critical incidents
        crit = sorted(bs["crit"], key=lambda c: ci_date(c) or date.min, reverse=True)
        st.markdown(f"**Critical incidents in this period ({len(crit)})**")
        if not crit:
            st.markdown('<span class="muted">None recorded.</span>', unsafe_allow_html=True)
        for c in crit[:6]:
            fn = " ".join(x for x in (c.get("primary_hypothesis_function"),
                                      (c.get("primary_hypothesis_item") or "").lower()) if x)
            st.markdown(
                f'<div class="ci-card"><b>{_fmt(ci_date(c), True)}</b> · severity {c.get("severity") or "—"}'
                f' · {_esc(c.get("primary_location") or "location not set")}'
                f'{" · " + _esc(fn) if fn else ""}'
                f'{" · <i>" + _esc(c.get("status")) + "</i>" if c.get("status") else ""}<br>'
                f'{_esc(_short(c.get("primary_behaviour"), 220))}</div>',
                unsafe_allow_html=True,
            )
        if len(crit) > 6:
            st.caption(f"+ {len(crit) - 6} more — open Behaviour Support for the full list.")

    # Scatter plot tracker
    st.divider()
    st.markdown("#### 📍 Scatter Plot Tracker")
    if sc["error"]:
        st.error(f"Couldn't load scatter plot data: {sc['error']}")
    elif not sc["weeks"]:
        st.info(f"No scatter plot weeks recorded for {ini}.")
    else:
        for b in scs["behaviours"] or list(sc["definitions"].keys()):
            definition = sc["definitions"].get(b)
            st.markdown(f"**Target behaviour:** {b}" + (f"  \n<span class='muted'>{_esc(definition)}</span>"
                                                       if definition else ""), unsafe_allow_html=True)
        if not scs["total"]:
            st.info("No tallies recorded in this period.")
        else:
            peak = max(scs["grid"].items(), key=lambda kv: kv[1])
            s1, s2, s3 = st.columns(3)
            s1.metric("Tallies", scs["total"])
            s2.metric("Weeks recorded", scs["weeks"])
            s3.metric("Peak", f"{peak[0][0][:3]} {peak[0][1]}")
            z = [[scs["grid"].get((d, t), 0) for d in DAY_NAMES] for t in scs["slots"]]
            hm = go.Figure(go.Heatmap(
                z=z, x=[d[:3] for d in DAY_NAMES], y=scs["slots"],
                colorscale=[[0, "#f4f8f8"], [0.5, "#e8a54b"], [1, RED]],
                showscale=False, text=[[v or "" for v in row] for row in z], texttemplate="%{text}",
                hovertemplate="%{y}, %{x}: %{z}<extra></extra>",
            ))
            hm.update_yaxes(autorange="reversed")
            hm.update_layout(title=dict(text="Tallies by time slot and day (all weeks in period)",
                                        font=dict(size=14)))
            st.plotly_chart(plot_layout(hm, max(220, 34 * len(scs["slots"]) + 80)),
                            use_container_width=True)

# ══════════════════════════════════════════════════════════════════════════
# TRANSITION
# ══════════════════════════════════════════════════════════════════════════
with tab_t:
    if tr["error"]:
        st.error(f"Couldn't load transition timetables: {tr['error']}")
    elif not tt:
        st.info(f"No transition timetable for {ini} yet.")
    else:
        view = tt_view(tt)
        weeks = tt_weeks(view)
        home_short = view.get("home_school_short") or (
            tu.school_short_name(view.get("home_school") or "") if tu else "")
        if tt.get("status") == "draft":
            st.warning("This timetable is still a **draft** — it hasn't been published to staff.")
        elif tu is not None and tu.has_unpublished_changes(tt):
            st.warning("Showing the **published** version — there are saved edits that haven't been published.")
        lcd = tu.last_clc_day(weeks) if (tu and weeks) else None
        rec = _d(view.get("recommencement_date"))
        st.markdown(
            f"**Status:** {tt_status(tt)}  \n"
            f"**Home school:** {view.get('home_school') or '—'}  \n"
            f"**Contact:** {view.get('home_school_contact') or '—'} · **BSC:** {view.get('bsc') or '—'}  \n"
            f"**Last CLC day:** {_fmt(lcd, True)}"
            + (f" · **Recommences at home school:** {_fmt(rec, True)}" if rec else "")
        )
        if view.get("notes"):
            st.markdown(f"<span class='muted'>📝 {_esc(view['notes'])}</span>", unsafe_allow_html=True)

        if weeks:
            # Step-down chart
            labels_w = [f"T{w['term']} W{w['week']}" for w in weeks]
            wc = [week_counts(w) for w in weeks]
            fig = go.Figure()
            for key, colour in (("At CLC", TEAL), ("Home school – part day", "#e8a54b"),
                                ("Home school – full day", GREEN), ("Other", "#9aa9a9")):
                vals = [c[key] for c in wc]
                if any(vals):
                    fig.add_bar(x=labels_w, y=vals, name=key, marker_color=colour)
            fig.update_layout(barmode="stack",
                              title=dict(text="Step-down: days per week", font=dict(size=14)))
            fig.update_yaxes(range=[0, 5], dtick=1)
            st.plotly_chart(plot_layout(fig, 260), use_container_width=True)

            # This week + next week
            this_mon = _monday(TODAY)
            next_mon = this_mon + timedelta(days=7)
            for title, mon in (("This week", this_mon), ("Next week", next_mon)):
                days = [tu.day_for(weeks, mon + timedelta(days=i)) if tu else None for i in range(5)]
                if not any(days):
                    continue
                st.markdown(f"**{title}** — week of {_fmt(mon, True)}")
                cols = st.columns(5)
                for i, d in enumerate(days):
                    dd = mon + timedelta(days=i)
                    if not d:
                        bg, txt = "#f4f8f8", "—"
                    else:
                        loc = d.get("loc")
                        bg = TEAL_LIGHT if loc == "CLC" else AMBER_LIGHT if loc == "HOME" else "#eef0f0"
                        txt = _esc(day_text(d, home_short))
                    border = f"2px solid {TEAL}" if dd == TODAY else "1px solid #d6e4e4"
                    cols[i].markdown(
                        f'<div class="day-cell" style="background:{bg};border:{border}">'
                        f'<div class="dn">{dd:%a %-d}</div>{txt}</div>',
                        unsafe_allow_html=True,
                    )

            with st.expander("Full timetable, week by week"):
                rows = []
                for w in weeks:
                    row = {"Week": f"T{w['term']} W{w['week']} {w['year']}"}
                    for d in w.get("days", []):
                        dd = _d(d.get("date"))
                        row[dd.strftime("%a") if dd else "?"] = day_text(d, home_short)
                    rows.append(row)
                st.dataframe(pd.DataFrame(rows), hide_index=True, use_container_width=True)

        if len(tr["timetables"]) > 1:
            st.caption(f"{len(tr['timetables'])} timetables on file for {ini}; showing the most current.")

# ══════════════════════════════════════════════════════════════════════════
# DOCUMENTS
# ══════════════════════════════════════════════════════════════════════════
with tab_d:
    docs = load_documents(ini)
    if not docs:
        st.info(f"No documents filed for {ini} in the Document Archive.")
    else:
        df = pd.DataFrame([{
            "Date": _fmt(_d(d.get("created_at")), True),
            "Type": d.get("doc_type") or "",
            "Title": d.get("title") or d.get("file_name") or "",
            "From": d.get("source_tool") or "",
        } for d in docs])
        st.dataframe(df, hide_index=True, use_container_width=True)
        pick = st.selectbox("Download a document", range(len(docs)),
                            format_func=lambda i: f"{df.iloc[i]['Date']} · {df.iloc[i]['Type']} · {df.iloc[i]['Title']}",
                            key="snap_doc_pick")
        if st.button("Prepare download", key="snap_doc_prep"):
            d = docs[pick]
            data = download_bytes(d.get("storage_path") or "")
            if data:
                st.download_button("⬇️ Download", data,
                                   file_name=d.get("file_name") or "document",
                                   key="snap_doc_dl")
            else:
                st.error("Couldn't fetch that file.")
        st.link_button("📁 Open Document Archive", f"{PORTAL}/Document_Archive")

st.divider()
st.caption("🧭 Student Snapshot · read-only view of SMARTAR, Behaviour Support, Scatter Plot Tracker, "
           "Transition Timetables and the Document Archive · initials only.")
