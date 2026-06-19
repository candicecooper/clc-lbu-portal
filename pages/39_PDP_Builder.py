from theme import apply_clc_theme
from auth import require_auth, logout_button, get_staff_name
from confidential_auth import require_confidential_auth
import streamlit as st
from supabase import create_client, Client
from datetime import datetime, date
import json
import io
import os

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(page_title="CLC PDP Builder", page_icon="🎯", layout="wide",
                   initial_sidebar_state="collapsed")
apply_clc_theme()
require_auth("PDP Builder", "🎯")
logout_button("sidebar")
require_confidential_auth("PDP Builder", "🎯")

TEAL = "#1A7A7A"
AMBER = "#C8760A"
PORTAL_URL = "https://16cefb77.clc-portal.pages.dev/"

st.sidebar.markdown(
    f'<a href="{PORTAL_URL}" target="_self" style="display:block;text-align:center;'
    f'background:{TEAL};color:white;font-weight:600;font-size:0.85rem;'
    f'padding:0.5rem;border-radius:8px;text-decoration:none;margin-bottom:0.5rem;">'
    f'🏠 Back to Staff Portal</a>',
    unsafe_allow_html=True
)

# ── Supabase ──────────────────────────────────────────────────────────────────
@st.cache_resource
def get_client() -> Client:
    url = st.secrets["SUPABASE_URL"]
    key = st.secrets["SUPABASE_KEY"]
    return create_client(url, key)

supabase = get_client()

# ── Groq ──────────────────────────────────────────────────────────────────────
try:
    GROQ_KEY = st.secrets.get("GROQ_API_KEY", "")
except Exception:
    GROQ_KEY = os.environ.get("GROQ_API_KEY", "")

DEFAULT_PRINCIPLES = (
    "Collective responsibility\nLearning system\nEvaluate for impact\n"
    "Tight and flexible\nTrust and verify\nImprovement and responsibility"
)
AREAS_OF_IMPACT = "Wellbeing, Equity and excellence, Learner agency, Effective learners"

# Maps the AI's flat keys → the page's widget keys
KEY_MAP = {
    "goal1_goal": "w_goal1_goal", "goal1_strengths": "w_goal1_strengths",
    "goal1_learn_change": "w_goal1_learn", "goal1_support": "w_goal1_support",
    "goal1_monitor": "w_goal1_monitor", "goal1_evaluate": "w_goal1_evaluate",
    "goal2_goal": "w_goal2_goal", "goal2_strengths": "w_goal2_strengths",
    "goal2_learn_change": "w_goal2_learn", "goal2_support": "w_goal2_support",
    "goal2_monitor": "w_goal2_monitor", "goal2_evaluate": "w_goal2_evaluate",
    "career_goal": "w_career_goal", "career_actions": "w_career_actions",
    "wellbeing_goal": "w_wb_goal", "wellbeing_actions": "w_wb_actions",
    "cultural_commitment": "w_cultural_commitment",
}

# ── Field definitions ─────────────────────────────────────────────────────────
FIELD_KEYS = [
    "w_emp", "w_ldr", "w_date",
    "w_goal1_goal", "w_goal1_strengths", "w_goal1_learn", "w_goal1_support",
    "w_goal1_monitor", "w_goal1_evaluate", "w_goal1_r6", "w_goal1_r12",
    "w_goal2_goal", "w_goal2_strengths", "w_goal2_learn", "w_goal2_support",
    "w_goal2_monitor", "w_goal2_evaluate", "w_goal2_r6", "w_goal2_r12",
    "w_career_goal", "w_career_actions", "w_career_r6", "w_career_r12",
    "w_wb_goal", "w_wb_actions", "w_wb_r6", "w_wb_r12",
    "w_cultural_commitment",
    "w_ref_emp6", "w_ref_ldr6", "w_ref_emp12", "w_ref_ldr12",
]
ADD_KEYS = [wk + "__add" for wk in KEY_MAP.values()]
# Every key whose value is part of the plan and must survive reruns
PLAN_KEYS = FIELD_KEYS + ADD_KEYS + ["w_cultural_focus"]

def seed_state():
    # Restore any plan values that Streamlit garbage-collected when a widget
    # wasn't rendered (e.g. after a rerun fired mid-script). The store is the
    # source of truth and is refreshed at the end of every run.
    store = st.session_state.setdefault("plan_store", {})
    for k in PLAN_KEYS:
        if k not in st.session_state and k in store:
            st.session_state[k] = store[k]
    for k in FIELD_KEYS:
        if k not in st.session_state:
            st.session_state[k] = ""
    for k in ADD_KEYS:
        if k not in st.session_state:
            st.session_state[k] = ""
    if "w_cultural_focus" not in st.session_state:
        st.session_state["w_cultural_focus"] = "Knowing"
    if "w_emp" in st.session_state and not st.session_state["w_emp"]:
        st.session_state["w_emp"] = get_staff_name() or ""
    if "w_date" in st.session_state and not st.session_state["w_date"]:
        st.session_state["w_date"] = date.today().strftime("%d/%m/%Y")
    if "transcript" not in st.session_state:
        st.session_state["transcript"] = ""
    if "principles" not in st.session_state:
        st.session_state["principles"] = DEFAULT_PRINCIPLES
    if "prior_pdp_text" not in st.session_state:
        st.session_state["prior_pdp_text"] = ""
    # Apply a pending saved-plan load before any widgets are instantiated
    pid = st.session_state.pop("_load_plan_id", None)
    if pid:
        try:
            res = supabase.table("pdp_plans").select("data").eq("id", pid).single().execute()
            for k, v in ((res.data or {}).get("data", {}) or {}).items():
                st.session_state[k] = v
        except Exception:
            pass

seed_state()

# ── Prior PDP extraction ──────────────────────────────────────────────────────
def extract_text(uploaded) -> str:
    name = uploaded.name.lower()
    data = uploaded.read()
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")
    if name.endswith(".docx"):
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        def wt(tag): return f"{{{W}}}{tag}"
        doc = Document(io.BytesIO(data))
        out, seen = [], set()
        def add(block):
            block = block.strip()
            if block and block not in seen:
                seen.add(block); out.append(block)
        # 1) text boxes (where the DfE template stores typed answers)
        for tx in doc.element.body.iter(wt('txbxContent')):
            paras = []
            for p in tx.iter(wt('p')):
                txt = "".join(n.text for n in p.iter(wt('t')) if n.text).strip()
                if txt:
                    paras.append(txt)
            add("\n".join(paras))
        # 2) body paragraphs and tables (python-docx text excludes text boxes, so no overlap)
        for child in doc.element.body.iterchildren():
            if child.tag == wt('p'):
                add(Paragraph(child, doc).text)
            elif child.tag == wt('tbl'):
                last = None
                for row in Table(child, doc).rows:
                    for cell in row.cells:
                        tc = cell.text.strip()
                        if tc and tc != last:
                            add(tc); last = tc
        return "\n".join(out)
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((pg.extract_text() or "") for pg in reader.pages)
    return data.decode("utf-8", errors="ignore")

# ── Strip DfE template boilerplate, keep section markers + the person's answers ──
_PDP_HEADINGS = {"goals", "career fulfilment", "my wellbeing", "my cultural responsiveness journey",
                 "reflection overview", "goal 1", "goal 2", "knowing", "doing", "being",
                 "at 6 months", "at 12 months"}
_PDP_SKIP_EXACT = {"feedback and goal reflections", "areas of impact", "guiding principles",
                   "preparing your pdp", "reflective practice", "employee summary", "leader summary",
                   "resources on edi", "employee name", "leader name", "employee signature",
                   "leader signature", "performance and development plan"}
_PDP_BOILER = [
    "the performance and development", "the pdp is", "the pdp framework", "your goals may relate",
    "this goal provides", "knowing what you need to learn", "we are building a culturally",
    "the 4 areas of impact", "each site will consider", "as you consider the goals",
    "our 5 guiding principles", "as you work through", "you may even choose", "it's a deliberate shift",
    "it is an ongoing cycle", "the 6 and 12 month", "outside of the goals", "use these questions",
    "refer to page", "to what extent", "what (if anything)", "you can update your goals",
    "identify goals that are achievable", "this area of the pdp", "being able to identify areas",
    "it provides an opportunity", "it might help you formulate", "see how you can embed",
    "when i feel at my best", "what's important for my own", "how can i contribute",
    "what will i deliver", "what actions will i take", "what does career fulfilment",
    "what aspects of my role", "how can you broaden", "how can you respond", "how can you strengthen",
    "this template provides", "please ensure you create", "select 1 way", "how the pdp aligns",
    "select the links", "what support do i need", "what do i need to learn", "what are my strengths",
    "how will i monitor", "how will i evaluate", "what people, processes", "what current behaviours",
    "what constructive feedback", "how will i know", "cultural responsiveness", "aboriginal learner",
    "what did you", "where do you think", "what other achievements", "what have you learned",
]

def clean_pdp_text(raw: str) -> str:
    kept, last_marker, seen = [], None, set()
    for line in raw.split("\n"):
        line = line.strip()
        if not line:
            continue
        low = line.lower().replace("\u2019", "'").rstrip(":")
        if low in _PDP_HEADINGS:
            marker = "## " + line.upper().rstrip(":")
            if marker != last_marker:
                kept.append(marker); last_marker = marker
            continue
        if line.endswith("?"):
            continue
        if low in _PDP_SKIP_EXACT:
            continue
        if any(low.startswith(s) for s in _PDP_BOILER):
            continue
        if len(line) < 12:
            continue
        if line in seen:
            continue
        seen.add(line); kept.append(line); last_marker = None  # allow a marker to recur after content
    return "\n".join(kept)

# ── Token-budget trimming (Groq free tier = 12k tokens/min) ───────────────────
def trim(text: str, max_chars: int):
    if len(text) <= max_chars:
        return text, False
    head = int(max_chars * 0.55)
    tail = max_chars - head
    marker = "\n\n[… middle trimmed to fit the free-tier limit …]\n\n"
    return text[:head] + marker + text[-tail:], True

# (model, max_output_tokens, input_char_budget) — 8B has a smaller per-minute window
GROQ_MODELS = [
    ("llama-3.3-70b-versatile", 2500, 26000),
    ("llama-3.1-8b-instant", 1400, 11000),
]

def groq_json(system: str, user: str) -> dict:
    from groq import Groq
    client = Groq(api_key=GROQ_KEY)
    last_err = None
    for model, max_tok, char_budget in GROQ_MODELS:
        u, _ = trim(user, char_budget)
        try:
            resp = client.chat.completions.create(
                model=model,
                max_tokens=max_tok,
                temperature=0.3,
                response_format={"type": "json_object"},
                messages=[{"role": "system", "content": system},
                          {"role": "user", "content": u}],
            )
            txt = (resp.choices[0].message.content or "").strip()
            txt = txt.replace("```json", "").replace("```", "").strip()
            s, e = txt.find("{"), txt.rfind("}")
            return json.loads(txt[s:e + 1])
        except Exception as ex:
            last_err = ex
            msg = str(ex).lower()
            if any(t in msg for t in ("rate_limit", "429", "413", "request too large", "tokens per")):
                continue  # limit hit on this model — try the next (smaller) one
            raise
    raise last_err

# ── Fresh draft from transcript only ──────────────────────────────────────────
def align_transcript(transcript: str, principles: str) -> dict:
    system = (
        "You are helping a South Australian Department for Education staff member draft a "
        "Performance and Development Plan (PDP) from a transcript of a development conversation. "
        "Use only what is grounded in the transcript — never invent achievements, goals or quotes. "
        "Where a section is not covered, return a short honest prompt to complete it "
        "(e.g. 'Not discussed — consider what success looks like here'). "
        "Write in Australian English, first person, in a warm, practical, coaching-aligned voice. "
        "Frame goals to be specific and measurable with a timeframe (SMARTAR-friendly) where supported. "
        f"Where natural, connect goals to the 4 areas of impact ({AREAS_OF_IMPACT}) and these guiding principles: {principles}. "
        "Keep each field to 1-4 sentences. Return ONLY a JSON object with this shape: "
        '{"employee_name":"","leader_name":"",'
        '"goal1":{"goal":"","strengths":"","learn_change":"","support":"","monitor":"","evaluate":""},'
        '"goal2":{"goal":"","strengths":"","learn_change":"","support":"","monitor":"","evaluate":""},'
        '"career":{"goal":"","actions":""},'
        '"wellbeing":{"goal":"","actions":""},'
        '"cultural":{"focus":"Knowing|Doing|Being","commitment":""}}'
    )
    return groq_json(system, "Development conversation transcript:\n\n" + transcript)

def apply_alignment(parsed: dict):
    m = {
        "goal1": {"goal": "w_goal1_goal", "strengths": "w_goal1_strengths", "learn_change": "w_goal1_learn",
                  "support": "w_goal1_support", "monitor": "w_goal1_monitor", "evaluate": "w_goal1_evaluate"},
        "goal2": {"goal": "w_goal2_goal", "strengths": "w_goal2_strengths", "learn_change": "w_goal2_learn",
                  "support": "w_goal2_support", "monitor": "w_goal2_monitor", "evaluate": "w_goal2_evaluate"},
        "career": {"goal": "w_career_goal", "actions": "w_career_actions"},
        "wellbeing": {"goal": "w_wb_goal", "actions": "w_wb_actions"},
    }
    for sect, mapping in m.items():
        obj = parsed.get(sect, {}) or {}
        for src, key in mapping.items():
            if obj.get(src):
                st.session_state[key] = obj[src]
    cult = parsed.get("cultural", {}) or {}
    if cult.get("focus") in ("Knowing", "Doing", "Being"):
        st.session_state["w_cultural_focus"] = cult["focus"]
    if cult.get("commitment"):
        st.session_state["w_cultural_commitment"] = cult["commitment"]

FIELD_KEY_LIST = (
    "goal1_goal, goal1_strengths, goal1_learn_change, goal1_support, goal1_monitor, goal1_evaluate, "
    "goal2_goal, goal2_strengths, goal2_learn_change, goal2_support, goal2_monitor, goal2_evaluate, "
    "career_goal, career_actions, wellbeing_goal, wellbeing_actions, cultural_commitment"
)

REFLECTION_MAP = {
    "goal1_6m": "w_goal1_r6", "goal1_12m": "w_goal1_r12",
    "goal2_6m": "w_goal2_r6", "goal2_12m": "w_goal2_r12",
    "career_6m": "w_career_r6", "career_12m": "w_career_r12",
    "wellbeing_6m": "w_wb_r6", "wellbeing_12m": "w_wb_r12",
}

# ── Step 1: load the uploaded PDP into the plan fields ────────────────────────
def import_pdp(prior_text: str) -> dict:
    system = (
        "You are given the text of a completed Performance and Development Plan (PDP), extracted from a "
        "Word template. Lines beginning with '##' are MARKERS: section markers (GOAL 1, GOAL 2, CAREER "
        "FULFILMENT, MY WELLBEING, MY CULTURAL RESPONSIVENESS JOURNEY) and reflection markers (AT 6 MONTHS, "
        "AT 12 MONTHS). Content after a marker belongs to it; an 'AT 6 MONTHS' block belongs to the most "
        "recent goal/section. Map the staff member's written responses into the structured fields. "
        "Use only what is in the document; if a section has no response, leave its fields empty; never invent. "
        "Keep wording close to the original, lightly tidied. "
        "Mapping hints: content about a teaching/learning program, student outcomes or skills is a goal; "
        "content about future role, position, uncertainty about next year or career direction is career "
        "fulfilment; content about feeling supported, lonely, grateful, or what someone enjoys at work is "
        "wellbeing; content about Aboriginal students, cultural training or responsive practice is cultural. "
        "Australian English. "
        f"Goal/section field keys: {FIELD_KEY_LIST}. "
        "Reflection field keys: goal1_6m, goal1_12m, goal2_6m, goal2_12m, career_6m, career_12m, "
        "wellbeing_6m, wellbeing_12m. "
        "Also return cultural_focus as Knowing|Doing|Being if indicated, else empty. "
        'Return ONLY JSON: {"cultural_focus":"","fields":{"<key>":""...}}'
    )
    return groq_json(system, "COMPLETED PDP DOCUMENT:\n\n" + prior_text)

def apply_import(parsed: dict):
    fields = parsed.get("fields", {}) or {}
    for k, wk in KEY_MAP.items():
        if fields.get(k):
            st.session_state[wk] = fields[k]
    for k, wk in REFLECTION_MAP.items():
        if fields.get(k):
            st.session_state[wk] = fields[k]
    if parsed.get("cultural_focus") in ("Knowing", "Doing", "Being"):
        st.session_state["w_cultural_focus"] = parsed["cultural_focus"]

# ── Step 2: extend each existing section with detail from the meeting ──────────
# The amber addition sits under the MAIN field of each section.
SECTION_MAIN = {
    "goal1": "w_goal1_goal", "goal2": "w_goal2_goal", "career": "w_career_goal",
    "wellbeing": "w_wb_goal", "cultural": "w_cultural_commitment",
}

def section_baseline() -> dict:
    g = st.session_state.get
    def sect(main, *extra):
        parts = [g(main, "")] + [g(e, "") for e in extra]
        return " ".join(p for p in parts if p).strip()
    return {
        "goal1": sect("w_goal1_goal", "w_goal1_strengths", "w_goal1_learn", "w_goal1_support",
                      "w_goal1_monitor", "w_goal1_evaluate"),
        "goal2": sect("w_goal2_goal", "w_goal2_strengths", "w_goal2_learn", "w_goal2_support",
                      "w_goal2_monitor", "w_goal2_evaluate"),
        "career": sect("w_career_goal", "w_career_actions"),
        "wellbeing": sect("w_wb_goal", "w_wb_actions"),
        "cultural": sect("w_cultural_commitment"),
    }

def add_meeting(baseline: dict, transcript: str, principles: str) -> dict:
    system = (
        "You are given (A) a staff member's EXISTING PDP content for each section, and (B) a transcript of a "
        "development conversation about that plan. For EACH section, write an EXTENSION that builds on and "
        "deepens the existing content using what was discussed — adding specific detail, examples, actions, "
        "progress, timeframes, evidence or nuance that enrich and extend the existing goal. This should make "
        "the goal more detailed and concrete. Do NOT restate the existing content; write the ADDITIONAL "
        "detail that extends it, so it reads naturally as a continuation. Also fold in any genuinely new "
        "points the conversation raised for that section. If the conversation adds nothing for a section, "
        "return an empty string for it. Ground everything in the transcript; never invent. First person, "
        "Australian English, warm and practical, 1-4 sentences each. "
        f"Where natural, connect to the 4 areas of impact ({AREAS_OF_IMPACT}) and guiding principles: {principles}. "
        "Sections: goal1, goal2, career, wellbeing, cultural. "
        'Return ONLY JSON: {"additions":{"goal1":"","goal2":"","career":"","wellbeing":"","cultural":""}}'
    )
    existing = "\n".join(f"[{k}] {v if v else '(empty)'}" for k, v in baseline.items())
    user = "EXISTING PDP CONTENT BY SECTION:\n" + existing + "\n\nCONVERSATION TRANSCRIPT:\n" + transcript
    return groq_json(system, user)

def apply_additions(parsed: dict):
    additions = parsed.get("additions", {}) or {}
    for sect, wk in SECTION_MAIN.items():
        if additions.get(sect):
            st.session_state[wk + "__add"] = additions[sect]

# ── Save / load / export ──────────────────────────────────────────────────────
def collect_plan() -> dict:
    keys = FIELD_KEYS + ["w_cultural_focus"] + ADD_KEYS
    return {k: st.session_state.get(k, "") for k in keys}

def save_plan():
    payload = {
        "staff_name": get_staff_name() or "",
        "employee_name": st.session_state.get("w_emp", ""),
        "leader_name": st.session_state.get("w_ldr", ""),
        "plan_date": st.session_state.get("w_date", ""),
        "data": collect_plan(),
        "updated_at": datetime.utcnow().isoformat(),
    }
    supabase.table("pdp_plans").insert(payload).execute()

def load_my_plans() -> list:
    name = get_staff_name() or ""
    res = (supabase.table("pdp_plans")
           .select("id, employee_name, plan_date, updated_at")
           .eq("staff_name", name).order("updated_at", desc=True).execute())
    return res.data or []

def load_plan(plan_id: str):
    res = supabase.table("pdp_plans").select("data").eq("id", plan_id).single().execute()
    for k, v in ((res.data or {}).get("data", {}) or {}).items():
        st.session_state[k] = v

def delete_plan(plan_id: str):
    supabase.table("pdp_plans").delete().eq("id", plan_id).execute()

def plan_to_markdown() -> str:
    g = st.session_state.get
    def block(wk):
        out = g(wk, "")
        add = g(wk + "__add", "")
        if add:
            out += f"\n\n> ➕ **Added from the meeting:** {add}"
        return out
    md = ["# Performance and Development Plan\n",
          f"**Employee:** {g('w_emp','')}  ", f"**Leader:** {g('w_ldr','')}  ",
          f"**Date:** {g('w_date','')}\n"]
    for n in ("1", "2"):
        md.append(f"\n## Goal {n}\n")
        md.append(f"**Goal & purpose:** {block(f'w_goal{n}_goal')}\n")
        md.append(f"- **Strengths to leverage:** {block(f'w_goal{n}_strengths')}")
        md.append(f"- **Learn / change:** {block(f'w_goal{n}_learn')}")
        md.append(f"- **Support needed:** {block(f'w_goal{n}_support')}")
        md.append(f"- **Monitor change:** {block(f'w_goal{n}_monitor')}")
        md.append(f"- **Evaluate impact:** {block(f'w_goal{n}_evaluate')}")
        md.append(f"\n*6 months:* {g(f'w_goal{n}_r6','')}")
        md.append(f"\n*12 months:* {g(f'w_goal{n}_r12','')}")
    md.append(f"\n## Career fulfilment\n{block('w_career_goal')}\n")
    md.append(f"**Actions:** {block('w_career_actions')}")
    md.append(f"\n## My wellbeing\n{block('w_wb_goal')}\n")
    md.append(f"**Actions:** {block('w_wb_actions')}")
    md.append(f"\n## Cultural responsiveness journey\n**Focus:** {g('w_cultural_focus','')}\n\n{block('w_cultural_commitment')}")
    md.append(f"\n## Reflection overview\n")
    md.append(f"**Employee — 6 months:** {g('w_ref_emp6','')}")
    md.append(f"**Leader — 6 months:** {g('w_ref_ldr6','')}")
    md.append(f"**Employee — 12 months:** {g('w_ref_emp12','')}")
    md.append(f"**Leader — 12 months:** {g('w_ref_ldr12','')}")
    return "\n".join(md)

def build_docx() -> bytes:
    """Build a Word document styled to match the DfE PDP proforma."""
    from docx import Document
    from docx.shared import Pt, RGBColor, Mm
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    g = st.session_state.get
    TEAL = "1A7A7A"; TEAL_DARK = "13605A"; TEAL_PALE = "E2F0ED"
    AMBER = "C8760A"; GREY = "5B6B68"; BORDER = "CFE0DC"
    W_MM = 184.6  # content width on A4 with 12.7mm side margins

    def _shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
        tcPr.append(shd)

    def _bd(cell, color=BORDER, sz=6):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement('w:tcBorders')
        for side in ("top", "bottom", "left", "right"):
            el = OxmlElement(f'w:{side}')
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
            el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)
            b.append(el)
        tcPr.append(b)

    def _mar(cell, t=80, bo=80, l=140, r=140):
        tcPr = cell._tc.get_or_add_tcPr()
        m = OxmlElement('w:tcMar')
        for side, val in (("top", t), ("bottom", bo), ("left", l), ("right", r)):
            el = OxmlElement(f'w:{side}'); el.set(qn('w:w'), str(val)); el.set(qn('w:type'), 'dxa')
            m.append(el)
        tcPr.append(m)

    def _tw(tbl, width_mm=W_MM):
        tbl.autofit = False
        tblPr = tbl._tbl.tblPr
        for ex in tblPr.findall(qn('w:tblW')):
            tblPr.remove(ex)
        tblW = OxmlElement('w:tblW')
        tblW.set(qn('w:w'), str(int(width_mm * 1440 / 25.4))); tblW.set(qn('w:type'), 'dxa')
        tblPr.append(tblW)

    def _run(p, text, bold=False, size=10.5, color="222222", italic=False):
        r = p.add_run("" if text is None else str(text)); r.bold = bold; r.italic = italic
        r.font.size = Pt(size); r.font.name = "Arial"; r.font.color.rgb = RGBColor.from_string(color)
        return r

    doc = Document()
    s = doc.sections[0]
    s.page_width = Mm(210); s.page_height = Mm(297)
    s.left_margin = s.right_margin = Mm(12.7); s.top_margin = s.bottom_margin = Mm(14)
    base = doc.styles["Normal"]; base.font.name = "Arial"; base.font.size = Pt(10.5)

    def gap(pt=2):
        doc.add_paragraph().paragraph_format.space_after = Pt(pt)

    def band(text, size=15, fill=TEAL):
        t = doc.add_table(rows=1, cols=1); _tw(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0); c.width = Mm(W_MM); _shade(c, fill); _mar(c, 90, 90, 160, 160)
        _run(c.paragraphs[0], text, bold=True, size=size, color="FFFFFF")
        gap()

    def subhead(text):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
        _run(p, text, bold=True, size=11, color=TEAL_DARK)

    def prompts(items):
        for it in items:
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(0)
            _run(p, it, size=9, color=GREY, italic=True)

    def box(fields, add_pairs=None):
        t = doc.add_table(rows=1, cols=1); _tw(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0); c.width = Mm(W_MM); _bd(c); _mar(c)
        first = True
        for label, value in fields:
            if not value:
                continue
            p = c.paragraphs[0] if first else c.add_paragraph(); first = False
            p.paragraph_format.space_after = Pt(3)
            if label:
                _run(p, label + ": ", bold=True, color=TEAL_DARK)
            _run(p, value)
        for label, add in (add_pairs or []):
            if not add:
                continue
            p = c.paragraphs[0] if first else c.add_paragraph(); first = False
            p.paragraph_format.space_after = Pt(1)
            _run(p, "➕ Added from the meeting", bold=True, size=10, color=AMBER)
            if label:
                _run(p, f"  ({label})", size=9, color=AMBER)
            p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(3)
            _run(p2, add, color=AMBER)
        if first:
            _run(c.paragraphs[0], "")
        gap()

    def reflections(v6, v12):
        cap = doc.add_paragraph(); cap.paragraph_format.space_after = Pt(2)
        _run(cap, "Feedback and goal reflections — to what extent has the goal been achieved? "
                  "What needs adjusting?", italic=True, size=9, color=GREY)
        t = doc.add_table(rows=2, cols=2); _tw(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, head in enumerate(("AT 6 MONTHS", "AT 12 MONTHS")):
            hc = t.cell(0, j); hc.width = Mm(W_MM / 2); _shade(hc, TEAL_DARK); _mar(hc, 50, 50, 140, 140)
            _run(hc.paragraphs[0], head, bold=True, size=9, color="FFFFFF")
        for j, val in enumerate((v6, v12)):
            cc = t.cell(1, j); cc.width = Mm(W_MM / 2); _bd(cc); _mar(cc)
            _run(cc.paragraphs[0], val or "", size=10)
        gap()

    # Title
    band("PERFORMANCE AND DEVELOPMENT PLAN", size=17)
    gov = doc.add_paragraph(); gov.paragraph_format.space_after = Pt(4)
    _run(gov, "Government of South Australia  ·  Department for Education  ·  "
              "Cowandilla Learning Centre — LBU", size=8.5, color=GREY)
    idt = doc.add_table(rows=3, cols=2); _tw(idt); idt.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (lbl, key) in enumerate([("Employee Name", "w_emp"), ("Leader Name", "w_ldr"), ("Date", "w_date")]):
        lc = idt.cell(i, 0); lc.width = Mm(45); _shade(lc, TEAL_PALE); _bd(lc); _mar(lc)
        _run(lc.paragraphs[0], lbl, bold=True, size=10, color=TEAL_DARK)
        vc = idt.cell(i, 1); vc.width = Mm(W_MM - 45); _bd(vc); _mar(vc)
        _run(vc.paragraphs[0], g(key, ""))
    sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(4); sp.paragraph_format.space_after = Pt(6)
    _run(sp, "When you are learning and thriving, everyone benefits.", bold=True, color=TEAL)

    GP = ["What's important for my own development as a professional?",
          "How can I contribute to the guiding principles and 4 areas of impact?",
          "What will I deliver, what actions will I take, by when and how?"]

    band("GOALS")
    prompts(GP)
    for n in ("1", "2"):
        subhead(f"GOAL {n}")
        box([("Goal & purpose", g(f"w_goal{n}_goal", "")),
             ("Strengths to leverage", g(f"w_goal{n}_strengths", "")),
             ("What I need to learn / change", g(f"w_goal{n}_learn", "")),
             ("Support I need", g(f"w_goal{n}_support", "")),
             ("How I'll monitor change", g(f"w_goal{n}_monitor", "")),
             ("How I'll evaluate impact", g(f"w_goal{n}_evaluate", ""))],
            [("Goal & purpose", g(f"w_goal{n}_goal__add", "")),
             ("Strengths", g(f"w_goal{n}_strengths__add", "")),
             ("Learn / change", g(f"w_goal{n}_learn__add", "")),
             ("Support", g(f"w_goal{n}_support__add", "")),
             ("Monitor", g(f"w_goal{n}_monitor__add", "")),
             ("Evaluate", g(f"w_goal{n}_evaluate__add", ""))])
        reflections(g(f"w_goal{n}_r6", ""), g(f"w_goal{n}_r12", ""))

    band("CAREER FULFILMENT")
    prompts(["What does career fulfilment look like for me?",
             "What aspects of my role do I enjoy and want to do more of?",
             "What actions will I take, by when and how?"])
    box([("Goal", g("w_career_goal", "")), ("Actions", g("w_career_actions", ""))],
        [("Goal", g("w_career_goal__add", "")), ("Actions", g("w_career_actions__add", ""))])
    reflections(g("w_career_r6", ""), g("w_career_r12", ""))

    band("MY WELLBEING")
    prompts(["When I'm at my best, what factors exist and how are they enabled?",
             "What actions will I take, by when and how?"])
    box([("Goal", g("w_wb_goal", "")), ("Actions", g("w_wb_actions", ""))],
        [("Goal", g("w_wb_goal__add", "")), ("Actions", g("w_wb_actions__add", ""))])
    reflections(g("w_wb_r6", ""), g("w_wb_r12", ""))

    band("MY CULTURAL RESPONSIVENESS JOURNEY")
    prompts(["Knowing — broaden knowledge of histories, cultures and languages of Aboriginal people.",
             "Doing — respond to the diverse needs, backgrounds, experiences and knowledge.",
             "Being — strengthen self-awareness of diverse cultural and linguistic needs."])
    fp = doc.add_paragraph(); fp.paragraph_format.space_after = Pt(2)
    _run(fp, "Focus: ", bold=True, color=TEAL_DARK); _run(fp, g("w_cultural_focus", ""))
    box([(None, g("w_cultural_commitment", ""))], [(None, g("w_cultural_commitment__add", ""))])

    band("REFLECTION OVERVIEW")
    prompts(["What did you enjoy most? Where were you most successful, and why?",
             "What did you find challenging? What have you learned?"])
    refmap = {("6 months", 0): "w_ref_emp6", ("6 months", 1): "w_ref_ldr6",
              ("12 months", 0): "w_ref_emp12", ("12 months", 1): "w_ref_ldr12"}
    for period in ("6 months", "12 months"):
        t = doc.add_table(rows=2, cols=2); _tw(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, who in enumerate(("Employee summary", "Leader summary")):
            hc = t.cell(0, j); hc.width = Mm(W_MM / 2); _shade(hc, TEAL_DARK); _mar(hc, 50, 50, 140, 140)
            _run(hc.paragraphs[0], f"{who} — at {period}", bold=True, size=9, color="FFFFFF")
        for j in range(2):
            cc = t.cell(1, j); cc.width = Mm(W_MM / 2); _bd(cc); _mar(cc)
            _run(cc.paragraphs[0], g(refmap[(period, j)], ""), size=10)
            sig = cc.add_paragraph(); sig.paragraph_format.space_before = Pt(14)
            _run(sig, "Signature: ______________________   Date: __________", size=8.5, color=GREY)
        gap()

    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()

SAMPLE = (
    "Leader: Thanks for making time. What's been going well this term?\n"
    "Staff: The interval support tracking has helped me see patterns earlier. I've gotten better at noticing "
    "when a student is moving out of their window of tolerance and stepping in before things escalate.\n"
    "Leader: Where do you want to grow?\n"
    "Staff: I'd like to build confidence leading restorative conversations after incidents. I'd like to be "
    "running structured restorative debriefs with students myself by the end of term 3, and supporting one "
    "colleague to do the same.\n"
    "Leader: What would help?\n"
    "Staff: Coaching from you, observing a couple of debriefs, and time carved out in the week. I'd know it's "
    "working if incident re-escalation drops and students re-engage faster afterwards.\n"
    "Leader: And wellbeing?\n"
    "Staff: The back-to-back days wear me down. I'm at my best when I get a proper reset at lunch and I've "
    "prepped the morning before. I want to protect that more deliberately.\n"
    "Leader: Career-wise?\n"
    "Staff: Longer term I'm interested in a coordinator role — leading practice across the centre.\n"
    "Leader: Cultural responsiveness?\n"
    "Staff: I want to deepen how I respond to our Aboriginal students — not just knowledge, but the doing, "
    "embedding more responsive practice into how I plan and run sessions."
)

# ── Header ────────────────────────────────────────────────────────────────────
st.markdown(f"""
<div style="background:linear-gradient(135deg,{TEAL},#0f5f5f);color:white;border-radius:12px;
            padding:1.25rem 1.5rem;margin-bottom:1rem;">
  <div style="font-size:1.5rem;font-weight:800;">🎯 PDP Builder</div>
  <div style="opacity:0.85;font-size:0.9rem;margin-top:0.2rem;">
    Draft a Performance and Development Plan from a development conversation, or layer a meeting onto a
    plan already drafted — every field editable.</div>
</div>
""", unsafe_allow_html=True)

st.info(
    "This supports the collaborative employee-and-leader PDP process. Drafts are a starting point — review and "
    "edit every field.", icon="💬")

# ── Identity ──────────────────────────────────────────────────────────────────
c1, c2, c3 = st.columns([2, 2, 1])
c1.text_input("Employee name", key="w_emp")
c2.text_input("Leader name", key="w_ldr")
c3.text_input("Date", key="w_date")

# ── Engine ────────────────────────────────────────────────────────────────────
st.markdown("### Build the plan")

has_prior = bool(st.session_state["prior_pdp_text"].strip())

# ---- Step 1: prior PDP → fields ----
st.markdown("**Step 1 — Start from the existing PDP** (optional)")
up = st.file_uploader(
    "Upload the PDP drafted before the meeting (Word, PDF or text). Skip this to draft fresh from the transcript.",
    type=["docx", "pdf", "txt", "md"])
if up is not None:
    try:
        st.session_state["prior_pdp_text"] = extract_text(up)
        has_prior = bool(st.session_state["prior_pdp_text"].strip())
        st.caption(f"✅ File read — {len(st.session_state['prior_pdp_text'])} characters. "
                   f"Now load it into the plan.")
    except Exception as ex:
        st.error(f"Couldn't read that file: {ex}. A .docx or .txt usually works best.")

load_clicked = st.button("📥 Load prior PDP into the plan", use_container_width=False, disabled=not has_prior)
if not has_prior:
    st.caption("Upload a file above to enable this.")

# ---- Step 2: transcript → additions ----
st.markdown("**Step 2 — Add the development conversation**")
st.text_area("Conversation transcript or notes", key="transcript", height=180,
             placeholder="Paste the development conversation transcript or notes here…")

plan_has_content = any(str(st.session_state.get(wk, "")).strip() for wk in KEY_MAP.values())

bc1, bc2, bc3 = st.columns([1.7, 1.7, 1.6])
add_clicked = bc1.button("➕ Add meeting to the plan", type="primary", use_container_width=True)
fresh_label = "✨ Draft fresh (replaces plan)" if plan_has_content else "✨ Draft fresh from transcript"
fresh_disabled = plan_has_content and not st.session_state.get("confirm_replace", False)
fresh_clicked = bc2.button(fresh_label, use_container_width=True, disabled=fresh_disabled)
def _load_sample_cb():
    st.session_state["transcript"] = SAMPLE
bc3.button("Load sample transcript", use_container_width=True, on_click=_load_sample_cb)

st.caption("**Add meeting** keeps everything already in the plan and layers the new points from the "
           "conversation on top (flagged in amber). This is the step to use after loading a PDP.")
if plan_has_content:
    st.checkbox("⚠️ Allow ‘Draft fresh’ to replace the loaded plan (this erases the imported content)",
                key="confirm_replace")
else:
    st.caption("**Draft fresh** drafts every field from the transcript alone — use it when starting with "
               "no prior PDP.")

def _guard():
    if not GROQ_KEY:
        st.error("No Groq API key found in secrets — add GROQ_API_KEY. You can still complete the plan manually.")
        return False
    return True

if load_clicked and _guard():
    cleaned = clean_pdp_text(st.session_state["prior_pdp_text"])
    if len(cleaned) < 30:  # cleaning stripped too much — fall back to raw text
        cleaned = st.session_state["prior_pdp_text"]
    p_in, p_trim = trim(cleaned, 18000)
    if p_trim:
        st.info("The uploaded PDP was long, so it was trimmed to fit Groq's free-tier limit.")
    with st.spinner("Reading the uploaded PDP and filling the plan…"):
        try:
            apply_import(import_pdp(p_in))
            st.success("Loaded. The existing plan content is in each field below — review it, then add the meeting.")
        except Exception as ex:
            st.error(f"Couldn't read that PDP into the plan: {ex}. Try a .docx, or shorten it.")

if add_clicked and _guard():
    if not st.session_state["transcript"].strip():
        st.warning("Paste a transcript first.")
    else:
        t_in, t_trim = trim(st.session_state["transcript"], 20000)
        if t_trim:
            st.info("The transcript was long, so it was trimmed to fit Groq's free-tier limit.")
        with st.spinner("Reading the conversation and finding what's new…"):
            try:
                apply_additions(add_meeting(section_baseline(), t_in, st.session_state["principles"]))
                st.success("Added. Each section now has an amber extension drawing detail from the meeting.")
            except Exception as ex:
                st.error(f"Couldn't add the meeting: {ex}. Try again or shorten the transcript.")

if fresh_clicked and _guard():
    if not st.session_state["transcript"].strip():
        st.warning("Paste a transcript first.")
    else:
        t_in, t_trim = trim(st.session_state["transcript"], 28000)
        if t_trim:
            st.info("The transcript was long, so it was trimmed to fit Groq's free-tier limit.")
        with st.spinner("Drafting each section from the transcript…"):
            try:
                apply_alignment(align_transcript(t_in, st.session_state["principles"]))
                st.success("Drafted. Review and edit every field — it's your plan.")
            except Exception as ex:
                st.error(f"Couldn't draft from that transcript: {ex}. Try again or shorten it.")

with st.expander("Strategy hooks (used by the alignment)"):
    st.markdown(f"**4 areas of impact:** {AREAS_OF_IMPACT}")
    st.text_area("Guiding principles — confirm wording against EDi", key="principles", height=120)

st.divider()

# ── Addition box helper ───────────────────────────────────────────────────────
def addition_box(widget_key: str):
    add_key = widget_key + "__add"
    if st.session_state.get(add_key):
        st.markdown(
            f"<div style='color:{AMBER};font-weight:600;font-size:0.8rem;margin:-0.4rem 0 0.1rem;'>"
            f"➕ Added from the meeting</div>", unsafe_allow_html=True)
        st.text_area("Added from the meeting", key=add_key, height=70, label_visibility="collapsed")

# ── GOALS ─────────────────────────────────────────────────────────────────────
def goal_block(n: str):
    st.subheader(f"Goal {n}")
    st.text_area("Goal & purpose — what I'll deliver, by when and how", key=f"w_goal{n}_goal", height=90)
    addition_box(f"w_goal{n}_goal")
    with st.expander("Coaching scaffold"):
        for label, suffix in [("My strengths to leverage", "strengths"),
                              ("What I need to learn / change", "learn"),
                              ("Support I need (people, processes, resources, tools)", "support"),
                              ("How I'll monitor change (feedback I'll seek)", "monitor"),
                              ("How I'll evaluate impact (how I'll know it worked)", "evaluate")]:
            st.text_area(label, key=f"w_goal{n}_{suffix}", height=70)
            addition_box(f"w_goal{n}_{suffix}")
    rc1, rc2 = st.columns(2)
    rc1.text_area("At 6 months", key=f"w_goal{n}_r6", height=90)
    rc2.text_area("At 12 months", key=f"w_goal{n}_r12", height=90)

st.markdown("## 🎯 Goals")
st.caption("Achievable within 6–12 months · tied to practice, standards or site priorities")
goal_block("1")
st.markdown("---")
goal_block("2")

# ── CAREER FULFILMENT ─────────────────────────────────────────────────────────
st.markdown("## 🏅 Career fulfilment")
st.caption("An area of focus or future aspiration · may take up to 2 years")
st.text_area("My career fulfilment goal", key="w_career_goal", height=90)
addition_box("w_career_goal")
st.text_area("Actions — what I'll do, by when and how", key="w_career_actions", height=70)
addition_box("w_career_actions")
cc1, cc2 = st.columns(2)
cc1.text_area("At 6 months", key="w_career_r6", height=80)
cc2.text_area("At 12 months", key="w_career_r12", height=80)

# ── WELLBEING ─────────────────────────────────────────────────────────────────
st.markdown("## 🌱 My wellbeing")
st.caption("What helps me learn and thrive at work")
st.text_area("My wellbeing goal", key="w_wb_goal", height=90)
addition_box("w_wb_goal")
st.text_area("Actions", key="w_wb_actions", height=70)
addition_box("w_wb_actions")
wc1, wc2 = st.columns(2)
wc1.text_area("At 6 months", key="w_wb_r6", height=80)
wc2.text_area("At 12 months", key="w_wb_r12", height=80)

# ── CULTURAL RESPONSIVENESS ───────────────────────────────────────────────────
st.markdown("## 🧭 My cultural responsiveness journey")
st.caption("Choose one way to build your commitment over 12 months")
st.radio("Focus", ["Knowing", "Doing", "Being"], key="w_cultural_focus", horizontal=True,
         help="Knowing: broaden knowledge of histories, cultures and languages. "
              "Doing: respond to diverse needs and knowledge. "
              "Being: strengthen self-awareness of cultural and linguistic needs.")
st.text_area("My commitment — what I'll learn or do over the next 12 months",
             key="w_cultural_commitment", height=90)
addition_box("w_cultural_commitment")

# ── REFLECTION OVERVIEW ───────────────────────────────────────────────────────
st.markdown("## 📝 Reflection overview")
st.caption("Completed together at the 6 and 12 month review")
r1, r2 = st.columns(2)
r1.text_area("Employee summary — 6 months", key="w_ref_emp6", height=90)
r2.text_area("Leader summary — 6 months", key="w_ref_ldr6", height=90)
r3, r4 = st.columns(2)
r3.text_area("Employee summary — 12 months", key="w_ref_emp12", height=90)
r4.text_area("Leader summary — 12 months", key="w_ref_ldr12", height=90)

st.divider()

# ── Save / load / export ──────────────────────────────────────────────────────
st.markdown("### Save & export")
sc1, sc2, sc3 = st.columns([1, 2, 1])

if sc1.button("💾 Save to my plans", type="primary", use_container_width=True):
    try:
        save_plan()
        st.success("Saved.")
    except Exception as ex:
        st.error(f"Save failed: {ex}")

try:
    plans = load_my_plans()
except Exception:
    plans = []

if plans:
    opts = {f"{p.get('employee_name','(no name)')} · {p.get('plan_date','')} "
            f"· saved {str(p.get('updated_at',''))[:16]}": p["id"] for p in plans}
    chosen = sc2.selectbox("My saved plans", list(opts.keys()), index=None,
                           placeholder="Choose a saved plan…")
    lc1, lc2 = sc2.columns(2)
    if lc1.button("Load", use_container_width=True, disabled=chosen is None):
        st.session_state["_load_plan_id"] = opts[chosen]; st.rerun()
    if lc2.button("Delete", use_container_width=True, disabled=chosen is None):
        delete_plan(opts[chosen]); st.rerun()
else:
    sc2.caption("No saved plans yet.")

_fname = (st.session_state.get("w_emp", "").strip().replace(" ", "_") or "PDP")
try:
    sc3.download_button(
        "⬇️ Download Word", data=build_docx(),
        file_name=f"PDP_{_fname}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True)
except Exception as ex:
    sc3.error(f"Word export error: {ex}")
sc3.download_button("⬇️ Download Markdown", data=plan_to_markdown(),
                    file_name=f"PDP_{_fname}.md", mime="text/markdown", use_container_width=True)

st.caption("Plans are stored privately against your staff name. Additions from the meeting are saved and "
           "exported with their amber 'Added from the meeting' label.")

# Keep the persistent store in sync with the current field values (runs after all
# widgets have rendered, so it captures both AI output and any manual edits).
_store = st.session_state.setdefault("plan_store", {})
for _k in PLAN_KEYS:
    _store[_k] = st.session_state.get(_k, "")
