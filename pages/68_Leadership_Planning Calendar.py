"""
CLC-LBU Leadership Planning Calendar 2027
-----------------------------------------
Confidential tile (shared confidential_auth.py password).

Leadership can log meetings, PD, SDDs and other commitments as
Proposed / Scheduled / Confirmed / Cancelled, with where they will occur,
then view each 2027 term week-by-week to plan ahead (starting with Term 1).

Supabase table: leadership_planning_events  (see leadership_planning_schema.sql)
Uses the main portal Supabase project (SUPABASE_URL / SUPABASE_KEY).
"""
from __future__ import annotations

import os
import io
import html
import uuid
from datetime import date, datetime, timedelta

import streamlit as st
from supabase import create_client, Client

from theme import apply_clc_theme
from confidential_auth import require_confidential_auth

try:
    from confidential_auth import confidential_logout_button
except Exception:  # keep the page working if the helper is renamed
    confidential_logout_button = None

st.set_page_config(
    page_title="Leadership Planning Calendar",
    page_icon="🗓️",
    layout="wide",
    initial_sidebar_state="collapsed",
)
apply_clc_theme()
require_confidential_auth("Leadership Planning Calendar", "🗓️")

PORTAL_URL = "https://clc-portal.pages.dev/"
st.sidebar.markdown(
    f'<a href="{PORTAL_URL}" target="_self" style="display:block;text-align:center;'
    f'background:#1A7A7A;color:white;font-weight:600;font-size:0.85rem;'
    f'padding:0.5rem;border-radius:8px;text-decoration:none;margin-bottom:0.5rem;">'
    f'🏠 Back to Staff Portal</a>',
    unsafe_allow_html=True,
)
if confidential_logout_button:
    try:
        with st.sidebar:
            confidential_logout_button()
    except Exception:
        pass

# ══════════════════════════════════════════════════════════════════════════════
# CONSTANTS
# ══════════════════════════════════════════════════════════════════════════════
TABLE = "leadership_planning_events"
YEAR = 2027

# SA Department for Education government school term dates 2027
TERMS = {
    1: (date(2027, 1, 27), date(2027, 4, 9)),
    2: (date(2027, 4, 26), date(2027, 7, 2)),
    3: (date(2027, 7, 19), date(2027, 9, 24)),
    4: (date(2027, 10, 11), date(2027, 12, 10)),
}

# SA public holidays 2027 that fall on weekdays (check/extend as needed)
PUBLIC_HOLIDAYS = {
    date(2027, 1, 1): "New Year's Day",
    date(2027, 1, 26): "Australia Day",
    date(2027, 3, 8): "Adelaide Cup Day",
    date(2027, 3, 26): "Good Friday",
    date(2027, 3, 29): "Easter Monday",
    date(2027, 6, 14): "King's Birthday",
    date(2027, 10, 4): "Labour Day",
}

EVENT_TYPES = [
    "Leadership meeting",
    "Staff meeting",
    "Professional development",
    "School development day",
    "Network / partner meeting",
    "External meeting",
    "Planning / admin",
    "Other",
]
TYPE_ICON = {
    "Leadership meeting": "👔",
    "Staff meeting": "👥",
    "Professional development": "🎓",
    "School development day": "🏫",
    "Network / partner meeting": "🤝",
    "External meeting": "🏢",
    "Planning / admin": "📋",
    "Other": "📌",
}

STATUSES = ["Proposed", "Scheduled", "Confirmed", "Cancelled"]
# background, text, border
STATUS_STYLE = {
    "Proposed": ("#FEF3C7", "#92400E", "#F59E0B"),
    "Scheduled": ("#DBEAFE", "#1E3A8A", "#3B82F6"),
    "Confirmed": ("#D1FAE5", "#065F46", "#10B981"),
    "Cancelled": ("#F1F5F9", "#64748B", "#94A3B8"),
}

LOCATION_MODES = ["On site – CLC", "Off site", "Online / Teams", "Hybrid", "TBC"]
LOCATION_ICON = {
    "On site – CLC": "🏠",
    "Off site": "📍",
    "Online / Teams": "💻",
    "Hybrid": "🔀",
    "TBC": "❔",
}

AUDIENCES = [
    "Leadership", "All staff", "Teachers", "SSOs",
    "JP team", "PY team", "SY team", "Admin", "External partners",
]

NO_TIME = "—"
TIME_SLOTS = [NO_TIME] + [
    f"{h:02d}:{m:02d}" for h in range(7, 19) for m in (0, 15, 30, 45)
]

WEEKDAY_NAMES = ["Mon", "Tue", "Wed", "Thu", "Fri"]

# ══════════════════════════════════════════════════════════════════════════════
# STYLES
# ══════════════════════════════════════════════════════════════════════════════
st.markdown("""
<style>
  .lp-header { background:linear-gradient(135deg,#1A7A7A,#0f5f5f);color:white;
    border-radius:12px;padding:1rem 1.4rem;margin-bottom:1rem; }
  .lp-header h1 { color:white;font-size:1.5rem;margin:0; }
  .lp-header p { margin:0.25rem 0 0 0;opacity:0.9;font-size:0.9rem; }
  .lp-legend { display:flex;flex-wrap:wrap;gap:0.5rem;margin:0.25rem 0 0.75rem 0; }
  .lp-pill { font-size:0.75rem;padding:2px 10px;border-radius:999px;border:1px solid;font-weight:600; }
  .lp-wrap { overflow-x:auto;width:100%; }
  table.lp-grid { border-collapse:collapse;width:100%;min-width:820px;table-layout:fixed; }
  table.lp-grid th { background:#0f5f5f;color:white;font-size:0.8rem;padding:6px;text-align:left; }
  table.lp-grid td { border:1px solid #e2e8f0;vertical-align:top;padding:4px 5px;
    min-height:70px;height:78px;font-size:0.78rem;background:white; }
  td.lp-week { background:#f0f7f7 !important;font-weight:700;color:#0f5f5f;width:92px; }
  td.lp-week span { display:block;font-weight:400;font-size:0.72rem;color:#475569; }
  td.lp-out { background:#f8fafc !important;color:#94a3b8; }
  td.lp-hol { background:#FDF2F8 !important; }
  .lp-daynum { font-size:0.72rem;color:#64748b;font-weight:600; }
  .lp-holname { font-size:0.7rem;color:#9D174D;font-weight:600; }
  .lp-chip { display:block;border-left:4px solid;border-radius:5px;padding:3px 5px;
    margin-top:3px;line-height:1.25;word-wrap:break-word; }
  .lp-chip b { font-weight:600; }
  .lp-chip .lp-meta { display:block;font-size:0.68rem;opacity:0.85; }
  .lp-cancel { text-decoration:line-through;opacity:0.7; }
</style>
""", unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# DATA LAYER
# ══════════════════════════════════════════════════════════════════════════════
@st.cache_resource
def get_supabase() -> Client | None:
    url = os.environ.get("SUPABASE_URL", "")
    key = os.environ.get("SUPABASE_KEY", "")
    if not (url and key):
        try:
            url = st.secrets.get("SUPABASE_URL", "") or url
            key = st.secrets.get("SUPABASE_KEY", "") or key
        except Exception:
            pass
    if not (url and key):
        return None
    return create_client(url, key)


def _to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except Exception:
        return None


@st.cache_data(ttl=60, show_spinner=False)
def load_events() -> tuple[list[dict], str | None]:
    sb = get_supabase()
    if sb is None:
        return [], "no_credentials"
    try:
        res = sb.table(TABLE).select("*").order("start_date").execute()
    except Exception as e:
        return [], str(e)
    rows = []
    for r in res.data or []:
        r["start_date"] = _to_date(r.get("start_date"))
        r["end_date"] = _to_date(r.get("end_date")) or r["start_date"]
        if r["start_date"] is None:
            continue
        rows.append(r)
    return rows, None


def save_event(data: dict, event_id: str | None = None) -> bool:
    sb = get_supabase()
    if sb is None:
        st.error("Supabase credentials not found.")
        return False
    payload = dict(data)
    payload["updated_at"] = datetime.utcnow().isoformat()
    try:
        if event_id:
            sb.table(TABLE).update(payload).eq("id", event_id).execute()
        else:
            payload["id"] = str(uuid.uuid4())
            payload["created_at"] = payload["updated_at"]
            sb.table(TABLE).insert(payload).execute()
    except Exception as e:
        st.error(f"Could not save: {e}")
        return False
    load_events.clear()
    return True


def delete_event(event_id: str) -> bool:
    sb = get_supabase()
    if sb is None:
        return False
    try:
        sb.table(TABLE).delete().eq("id", event_id).execute()
    except Exception as e:
        st.error(f"Could not delete: {e}")
        return False
    load_events.clear()
    return True


# ══════════════════════════════════════════════════════════════════════════════
# HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def term_weeks(term: int, include_pre_term: bool) -> list[tuple[str, date]]:
    """Return [(label, monday), ...] for a term."""
    start, end = TERMS[term]
    monday = start - timedelta(days=start.weekday())
    last_monday = end - timedelta(days=end.weekday())
    weeks = []
    if include_pre_term:
        weeks.append(("Pre-term", monday - timedelta(days=7)))
    n = 1
    while monday <= last_monday:
        weeks.append((f"Week {n}", monday))
        monday += timedelta(days=7)
        n += 1
    return weeks


def term_and_week_for(d: date) -> str:
    for t, (start, end) in TERMS.items():
        first_monday = start - timedelta(days=start.weekday())
        if first_monday - timedelta(days=7) <= d <= end:
            if d < first_monday:
                return f"T{t} pre-term"
            wk = (d - first_monday).days // 7 + 1
            return f"T{t} Wk {wk}"
    return "Holidays"


def fmt_date_range(r: dict) -> str:
    s, e = r["start_date"], r["end_date"]
    if e and e != s:
        return f"{s.strftime('%a %d %b')} – {e.strftime('%a %d %b')}"
    return s.strftime("%a %d %b %Y")


def fmt_time(r: dict) -> str:
    st_, en = r.get("start_time") or "", r.get("end_time") or ""
    if st_ and en:
        return f"{st_}–{en}"
    return st_ or ""


def fmt_location(r: dict) -> str:
    mode = r.get("location_mode") or ""
    loc = r.get("location") or ""
    if mode and loc:
        return f"{mode}: {loc}"
    return mode or loc or "TBC"


def chip_html(r: dict) -> str:
    bg, fg, border = STATUS_STYLE.get(r.get("status"), STATUS_STYLE["Proposed"])
    icon = TYPE_ICON.get(r.get("event_type"), "📌")
    loc_icon = LOCATION_ICON.get(r.get("location_mode"), "📍")
    cancel = " lp-cancel" if r.get("status") == "Cancelled" else ""
    time_txt = fmt_time(r)
    meta_bits = []
    if time_txt:
        meta_bits.append(html.escape(time_txt))
    meta_bits.append(f"{loc_icon} {html.escape(r.get('location') or r.get('location_mode') or 'TBC')}")
    title_attr = html.escape(
        f"{r.get('status')} · {r.get('event_type')}\n{fmt_location(r)}\n{r.get('summary') or ''}"
    )
    return (
        f'<span class="lp-chip{cancel}" title="{title_attr}" '
        f'style="background:{bg};color:{fg};border-color:{border};">'
        f'{icon} <b>{html.escape(r.get("title") or "")}</b>'
        f'<span class="lp-meta">{" · ".join(meta_bits)}</span></span>'
    )


def events_on(events: list[dict], d: date) -> list[dict]:
    hits = [r for r in events if r["start_date"] <= d <= r["end_date"]]
    return sorted(hits, key=lambda r: (r.get("start_time") or "99", r.get("title") or ""))


def render_term_grid(term: int, events: list[dict], include_pre_term: bool) -> set[date]:
    start, end = TERMS[term]
    weeks = term_weeks(term, include_pre_term)
    shown: set[date] = set()
    rows = []
    for label, monday in weeks:
        cells = [
            f'<td class="lp-week">{label}<span>{monday.strftime("%d %b")}</span></td>'
        ]
        for i in range(5):
            d = monday + timedelta(days=i)
            shown.add(d)
            classes = []
            in_term = start <= d <= end
            if not in_term:
                classes.append("lp-out")
            hol = PUBLIC_HOLIDAYS.get(d)
            if hol:
                classes.append("lp-hol")
            inner = f'<div class="lp-daynum">{d.day} {d.strftime("%b")}'
            if not in_term and label != "Pre-term":
                inner += " · outside term"
            inner += "</div>"
            if hol:
                inner += f'<div class="lp-holname">🎌 {html.escape(hol)}</div>'
            inner += "".join(chip_html(r) for r in events_on(events, d))
            cls = f' class="{" ".join(classes)}"' if classes else ""
            cells.append(f"<td{cls}>{inner}</td>")
        rows.append("<tr>" + "".join(cells) + "</tr>")

    head = "<tr><th style='width:92px'>Week</th>" + "".join(
        f"<th>{n}</th>" for n in WEEKDAY_NAMES
    ) + "</tr>"
    st.markdown(
        f'<div class="lp-wrap"><table class="lp-grid">{head}{"".join(rows)}</table></div>',
        unsafe_allow_html=True,
    )
    return shown


def legend():
    pills = "".join(
        f'<span class="lp-pill" style="background:{bg};color:{fg};border-color:{bd};">{s}</span>'
        for s, (bg, fg, bd) in STATUS_STYLE.items()
    )
    pills += ('<span class="lp-pill" style="background:#FDF2F8;color:#9D174D;'
              'border-color:#F9A8D4;">Public holiday</span>')
    st.markdown(f'<div class="lp-legend">{pills}</div>', unsafe_allow_html=True)


def build_docx(term: int, events: list[dict]) -> bytes | None:
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.section import WD_ORIENT
    except Exception:
        return None

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9.5)

    start, end = TERMS[term]
    h = doc.add_heading(f"CLC-LBU Leadership Planning — Term {term} {YEAR}", level=1)
    h.runs[0].font.color.rgb = RGBColor(0x1A, 0x7A, 0x7A)
    doc.add_paragraph(
        f"Term dates: {start.strftime('%a %d %b')} – {end.strftime('%a %d %b %Y')}   ·   "
        f"CONFIDENTIAL — leadership only   ·   Exported {date.today().strftime('%d %b %Y')}"
    )

    if not events:
        doc.add_paragraph("No items entered for this term yet.")
    else:
        cols = ["When", "Time", "Item", "Type", "Status", "Where", "Audience", "Lead", "Summary / next action"]
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Light Grid Accent 1"
        for i, c in enumerate(cols):
            table.rows[0].cells[i].text = c
            for p in table.rows[0].cells[i].paragraphs:
                for run in p.runs:
                    run.font.bold = True
        for r in events:
            summary = r.get("summary") or ""
            if r.get("next_action"):
                owner = f" ({r['action_owner']})" if r.get("action_owner") else ""
                summary = (summary + "\n" if summary else "") + f"Next: {r['next_action']}{owner}"
            vals = [
                f"{term_and_week_for(r['start_date'])}\n{fmt_date_range(r)}",
                fmt_time(r),
                r.get("title") or "",
                r.get("event_type") or "",
                r.get("status") or "",
                fmt_location(r),
                r.get("audience") or "",
                r.get("lead") or "",
                summary,
            ]
            cells = table.add_row().cells
            for i, v in enumerate(vals):
                cells[i].text = v

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_csv(events: list[dict]) -> bytes:
    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Term/Week", "Start date", "End date", "Start time", "End time", "Title", "Type",
                "Status", "Location type", "Location", "Audience", "Lead", "Summary",
                "Next action", "Action owner", "Entered by"])
    for r in events:
        w.writerow([
            term_and_week_for(r["start_date"]), r["start_date"], r["end_date"],
            r.get("start_time") or "", r.get("end_time") or "", r.get("title") or "",
            r.get("event_type") or "", r.get("status") or "", r.get("location_mode") or "",
            r.get("location") or "", r.get("audience") or "", r.get("lead") or "",
            r.get("summary") or "", r.get("next_action") or "", r.get("action_owner") or "",
            r.get("created_by") or "",
        ])
    return buf.getvalue().encode("utf-8-sig")


# ══════════════════════════════════════════════════════════════════════════════
# FORM (used for both new entries and edits)
# ══════════════════════════════════════════════════════════════════════════════
def render_event_form(prefix: str, existing: dict | None = None):
    ex = existing or {}
    default_start = ex.get("start_date") or TERMS[1][0]
    default_end = ex.get("end_date") or default_start

    with st.form(f"{prefix}_form", clear_on_submit=existing is None):
        title = st.text_input("Title *", value=ex.get("title", ""),
                              placeholder="e.g. Term 1 leadership planning day",
                              key=f"{prefix}_title")
        c1, c2 = st.columns(2)
        with c1:
            etype = st.selectbox(
                "Type", EVENT_TYPES,
                index=EVENT_TYPES.index(ex["event_type"]) if ex.get("event_type") in EVENT_TYPES else 0,
                key=f"{prefix}_type",
            )
        with c2:
            status = st.selectbox(
                "Status", STATUSES,
                index=STATUSES.index(ex["status"]) if ex.get("status") in STATUSES else 0,
                key=f"{prefix}_status",
            )

        c1, c2, c3, c4 = st.columns(4)
        with c1:
            sdate = st.date_input("Start date *", value=default_start, format="DD/MM/YYYY",
                                  key=f"{prefix}_sdate")
        with c2:
            edate = st.date_input("End date (multi-day)", value=default_end, format="DD/MM/YYYY",
                                  key=f"{prefix}_edate",
                                  help="Leave the same as the start date for a one-day item.")
        with c3:
            stime_val = ex.get("start_time") or NO_TIME
            stime = st.selectbox("Start time", TIME_SLOTS,
                                 index=TIME_SLOTS.index(stime_val) if stime_val in TIME_SLOTS else 0,
                                 key=f"{prefix}_stime")
        with c4:
            etime_val = ex.get("end_time") or NO_TIME
            etime = st.selectbox("End time", TIME_SLOTS,
                                 index=TIME_SLOTS.index(etime_val) if etime_val in TIME_SLOTS else 0,
                                 key=f"{prefix}_etime")

        c1, c2 = st.columns([1, 2])
        with c1:
            lmode = st.selectbox(
                "Where", LOCATION_MODES,
                index=LOCATION_MODES.index(ex["location_mode"]) if ex.get("location_mode") in LOCATION_MODES else 0,
                key=f"{prefix}_lmode",
            )
        with c2:
            location = st.text_input("Venue / room / link details", value=ex.get("location", "") or "",
                                     placeholder="e.g. Staff room, LBU head office, Teams link to follow",
                                     key=f"{prefix}_loc")

        existing_aud = [a.strip() for a in (ex.get("audience") or "").split(",") if a.strip() in AUDIENCES]
        c1, c2 = st.columns([2, 1])
        with c1:
            audience = st.multiselect("Who's involved", AUDIENCES, default=existing_aud,
                                      key=f"{prefix}_aud")
        with c2:
            lead = st.text_input("Lead / facilitator", value=ex.get("lead", "") or "",
                                 key=f"{prefix}_lead")

        summary = st.text_area("Summary details", value=ex.get("summary", "") or "", height=100,
                               placeholder="Purpose, focus, key details, bookings or costs to note…",
                               key=f"{prefix}_summary")
        c1, c2 = st.columns([2, 1])
        with c1:
            next_action = st.text_input("Next action to progress this", value=ex.get("next_action", "") or "",
                                        placeholder="e.g. Confirm venue with presenter",
                                        key=f"{prefix}_next")
        with c2:
            action_owner = st.text_input("Action owner", value=ex.get("action_owner", "") or "",
                                         key=f"{prefix}_owner")

        created_by = ex.get("created_by") or ""
        if existing is None:
            created_by = st.text_input("Entered by", value=st.session_state.get("lp_entered_by", ""),
                                       key=f"{prefix}_by")

        submitted = st.form_submit_button("💾 Save changes" if existing else "➕ Add to calendar",
                                          type="primary")

    if submitted:
        if not title.strip():
            st.warning("Please give the item a title.")
            return False
        if edate < sdate:
            st.warning("End date can't be before the start date.")
            return False
        if existing is None and created_by:
            st.session_state["lp_entered_by"] = created_by
        data = {
            "title": title.strip(),
            "event_type": etype,
            "status": status,
            "start_date": sdate.isoformat(),
            "end_date": edate.isoformat(),
            "start_time": None if stime == NO_TIME else stime,
            "end_time": None if etime == NO_TIME else etime,
            "location_mode": lmode,
            "location": location.strip() or None,
            "audience": ", ".join(audience) or None,
            "lead": lead.strip() or None,
            "summary": summary.strip() or None,
            "next_action": next_action.strip() or None,
            "action_owner": action_owner.strip() or None,
            "created_by": (created_by or "").strip() or None,
        }
        ok = save_event(data, ex.get("id"))
        if ok:
            st.success("Saved." if existing else f"Added “{title.strip()}”.")
            return True
    return False


# ══════════════════════════════════════════════════════════════════════════════
# PAGE
# ══════════════════════════════════════════════════════════════════════════════
st.markdown(
    f'<div class="lp-header"><h1>🗓️ Leadership Planning Calendar {YEAR}</h1>'
    f'<p>Confidential — meetings, PD and key dates being proposed, scheduled or confirmed, '
    f'and where they will happen.</p></div>',
    unsafe_allow_html=True,
)

events, err = load_events()
if err == "no_credentials":
    st.error("Supabase credentials (SUPABASE_URL / SUPABASE_KEY) were not found in Render environment variables.")
    st.stop()
elif err:
    if "does not exist" in err or "PGRST205" in err or "schema cache" in err:
        st.error(f"The `{TABLE}` table hasn't been created yet. Run `leadership_planning_schema.sql` "
                 "in the Supabase SQL editor (main portal project), then refresh this page.")
    else:
        st.error(f"Could not load the planning calendar: {err}")
    st.stop()

# ── Term + filters ────────────────────────────────────────────────────────────
c1, c2, c3 = st.columns([1, 2, 2])
with c1:
    term = st.selectbox("Term", [1, 2, 3, 4], format_func=lambda t: f"Term {t} {YEAR}", key="lp_term")
with c2:
    status_filter = st.multiselect("Show status", STATUSES, default=["Proposed", "Scheduled", "Confirmed"],
                                   key="lp_status_filter")
with c3:
    type_filter = st.multiselect("Show types", EVENT_TYPES, default=EVENT_TYPES, key="lp_type_filter")

filtered = [r for r in events
            if r.get("status") in status_filter and r.get("event_type") in type_filter]

t_start, t_end = TERMS[term]
pre_term_start = t_start - timedelta(days=t_start.weekday() + 7)
term_events = [r for r in filtered if r["end_date"] >= pre_term_start and r["start_date"] <= t_end]
all_term_events = [r for r in events if r["end_date"] >= pre_term_start and r["start_date"] <= t_end]

m1, m2, m3, m4 = st.columns(4)
m1.metric("Confirmed", sum(r.get("status") == "Confirmed" for r in all_term_events))
m2.metric("Scheduled", sum(r.get("status") == "Scheduled" for r in all_term_events))
m3.metric("Proposed", sum(r.get("status") == "Proposed" for r in all_term_events))
m4.metric("Location TBC", sum(
    (r.get("location_mode") == "TBC") and r.get("status") != "Cancelled" for r in all_term_events))

tab_cal, tab_add, tab_list, tab_export = st.tabs(
    ["📅 Term calendar", "➕ Add item", "📝 Review & edit", "⬇️ Export"]
)

# ── Calendar ──────────────────────────────────────────────────────────────────
with tab_cal:
    st.caption(f"Term {term}: {t_start.strftime('%a %d %b')} – {t_end.strftime('%a %d %b %Y')}. "
               "Hover an item for its summary.")
    include_pre = st.toggle("Show the week before term starts", value=True, key="lp_pre",
                            help="Useful for planning leadership days and SDDs before students return.")
    legend()
    shown_days = render_term_grid(term, term_events, include_pre)

    hidden = [r for r in term_events
              if not any(r["start_date"] <= d <= r["end_date"] for d in shown_days)]
    if hidden:
        st.markdown("**Also in this period (weekend or outside the grid)**")
        for r in hidden:
            st.markdown(f"- {fmt_date_range(r)} — {r.get('title')} ({r.get('status')}, {fmt_location(r)})")

    holiday_events = [r for r in filtered
                      if term_and_week_for(r["start_date"]) == "Holidays"]
    if holiday_events:
        with st.expander(f"School holiday items ({len(holiday_events)})"):
            for r in holiday_events:
                st.markdown(f"- {fmt_date_range(r)} — {r.get('title')} ({r.get('status')}, {fmt_location(r)})")

# ── Add ───────────────────────────────────────────────────────────────────────
with tab_add:
    st.caption("Start with what you know — an item can be Proposed with the location TBC and updated later.")
    if render_event_form("lp_new"):
        st.rerun()

# ── Review & edit ─────────────────────────────────────────────────────────────
with tab_list:
    scope = st.radio("Show", [f"Term {term} (incl. pre-term week)", "All of 2027"], horizontal=True,
                     key="lp_scope")
    source = [r for r in events if r.get("event_type") in type_filter and r.get("status") in STATUSES]
    if scope.startswith("Term"):
        source = [r for r in source if r["end_date"] >= pre_term_start and r["start_date"] <= t_end]
    search = st.text_input("Search", placeholder="Title, venue, lead, notes…", key="lp_search").strip().lower()
    if search:
        source = [r for r in source if search in " ".join(
            str(r.get(k) or "") for k in ("title", "location", "lead", "summary", "next_action", "audience")
        ).lower()]

    if not source:
        st.info("Nothing here yet — add the first item in the ➕ Add item tab.")
    for r in source:
        bg, fg, bd = STATUS_STYLE.get(r.get("status"), STATUS_STYLE["Proposed"])
        icon = TYPE_ICON.get(r.get("event_type"), "📌")
        label = (f"{icon} {r['start_date'].strftime('%d %b')} · {r.get('title')} "
                 f"— {r.get('status')} · {term_and_week_for(r['start_date'])}")
        with st.expander(label):
            st.markdown(
                f'<span class="lp-pill" style="background:{bg};color:{fg};border-color:{bd};">'
                f'{r.get("status")}</span> &nbsp; **{fmt_date_range(r)}** {fmt_time(r)} · '
                f'{LOCATION_ICON.get(r.get("location_mode"), "📍")} {html.escape(fmt_location(r))}',
                unsafe_allow_html=True,
            )
            if r.get("next_action"):
                owner = f" — {r['action_owner']}" if r.get("action_owner") else ""
                st.markdown(f"➡️ **Next:** {r['next_action']}{owner}")
            if r.get("created_by"):
                st.caption(f"Entered by {r['created_by']}")

            if render_event_form(f"lp_edit_{r['id']}", existing=r):
                st.rerun()

            confirm_key = f"lp_del_confirm_{r['id']}"
            if not st.session_state.get(confirm_key):
                if st.button("🗑️ Delete", key=f"lp_del_{r['id']}"):
                    st.session_state[confirm_key] = True
                    st.rerun()
            else:
                st.warning("Delete this item permanently? (Use status ‘Cancelled’ to keep a record instead.)")
                d1, d2 = st.columns(2)
                if d1.button("Yes, delete", key=f"lp_del_yes_{r['id']}", type="primary"):
                    st.session_state.pop(confirm_key, None)
                    if delete_event(r["id"]):
                        st.rerun()
                if d2.button("Keep it", key=f"lp_del_no_{r['id']}"):
                    st.session_state.pop(confirm_key, None)
                    st.rerun()

# ── Export ────────────────────────────────────────────────────────────────────
with tab_export:
    st.caption("Exports use the current status/type filters. Marked confidential — share with leadership only.")
    export_events = sorted(term_events, key=lambda r: (r["start_date"], r.get("start_time") or "99"))
    docx_bytes = build_docx(term, export_events)
    c1, c2 = st.columns(2)
    with c1:
        if docx_bytes:
            st.download_button(
                f"📄 Term {term} planning overview (Word)",
                data=docx_bytes,
                file_name=f"CLC_Leadership_Planning_T{term}_{YEAR}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )
        else:
            st.info("Word export needs python-docx in requirements.txt.")
    with c2:
        st.download_button(
            f"📊 Term {term} items (CSV / Excel)",
            data=build_csv(export_events),
            file_name=f"CLC_Leadership_Planning_T{term}_{YEAR}.csv",
            mime="text/csv",
            use_container_width=True,
        )
