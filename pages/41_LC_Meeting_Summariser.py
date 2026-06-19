"""
LC Meeting Summariser
=====================
Drop in your LC leader meeting agenda (optional) and the meeting transcript,
and get back: a summary, agenda-item breakdown, actions, and a classified
"share with staff" section that you approve before anything is exported.
Exports to a Word document that mirrors the DfE LBu LC leader proforma.

Share-with-staff handling:
  - Every candidate item is classified Share / Hold / Your call, each with a reason.
  - A named guardrail list keeps confidential matters (named staff
    performance/conduct/health, IR, role-security before official comms,
    pre-decisional or OFFICIAL: Sensitive material) on Hold by default.
  - A "staff-safe wording" is offered so you can relay the gist without the
    confidential detail.
  - You tick what actually goes out. Only ticked items reach the staff section
    of the Word doc; held items go to a separate "Confidential — manager record
    only" section marked OFFICIAL: Sensitive.

Designed for the CLC-LBU portal. Single page, single AI call, graceful
fallbacks if Supabase / Claude secrets aren't configured.

Stack: Streamlit + Claude (claude-sonnet-4-6) + python-docx + Supabase (optional)
"""

import streamlit as st
from confidential_auth import require_confidential_auth
import json
import io
import os
import re
from datetime import datetime, date

# ---------------------------------------------------------------------------
# Optional dependencies — handled gracefully so the page never hard-crashes
# ---------------------------------------------------------------------------
try:
    from anthropic import Anthropic
    ANTHROPIC_LIB = True
except Exception:
    ANTHROPIC_LIB = False

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    DOCX_LIB = True
except Exception:
    DOCX_LIB = False

try:
    from supabase import create_client
    SUPABASE_LIB = True
except Exception:
    SUPABASE_LIB = False

try:
    import docx as _docx_read  # python-docx, reused for reading uploads
    READ_DOCX = True
except Exception:
    READ_DOCX = False


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
MODEL = "claude-sonnet-4-6"
MAX_TRANSCRIPT_CHARS = 400000  # Claude handles ~200k tokens; this is a safety net only
SUPABASE_TABLE = "lc_meetings"

# This week's agenda, captured from your 19.06.2026 proforma, so you can
# load it with one click and test the tool straight away.
SAMPLE_AGENDA = """LBu LC leader Agenda
Office: Schools and Preschools | Date of meeting: 19.06.2026
Directorate: Priority Learners | Unit: LBu
Time: 8.00am | Location: Microsoft Teams
Chair: Susan Furga
Attendees: Beck; Candice, Sarah; Andrew

Actions from previous meetings:
(none recorded)

Agenda items:
- Susan: Reform - injury report
- Susan: Managing transitions re reform discussion and reflection
- Everyone: Elephant ED
- Susan: Partial freeze impacts and changes in corporate
- Susan: Bus Driver role update
- EMS - no updates
- Resource plans for new financial year
- AoB
"""

try:
    st.set_page_config(page_title="LC Meeting Summariser", page_icon="🗒️", layout="wide")
except Exception:
    pass


# ---------------------------------------------------------------------------
# Helpers: secrets / clients
# ---------------------------------------------------------------------------
def get_secret(name):
    # Check Streamlit secrets (secrets.toml) first, then OS environment
    # variables (e.g. Render's Environment tab) so either location works.
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:
        pass
    return os.environ.get(name)


def get_anthropic_client():
    key = get_secret("ANTHROPIC_API_KEY")
    if not (ANTHROPIC_LIB and key):
        return None
    try:
        return Anthropic(api_key=key)
    except Exception:
        return None


@st.cache_resource(show_spinner=False)
def get_supabase():
    url = get_secret("SUPABASE_URL")
    key = get_secret("SUPABASE_KEY") or get_secret("SUPABASE_ANON_KEY")
    if not (SUPABASE_LIB and url and key):
        return None
    try:
        return create_client(url, key)
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Helpers: reading uploads
# ---------------------------------------------------------------------------
def read_upload(uploaded_file):
    """Return plain text from an uploaded .docx / .txt / .vtt / .md file."""
    if uploaded_file is None:
        return ""
    name = (uploaded_file.name or "").lower()
    data = uploaded_file.getvalue()

    if name.endswith(".docx") and READ_DOCX:
        try:
            doc = _docx_read.Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for tbl in doc.tables:
                for row in tbl.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as e:
            return f"[Could not read .docx: {e}]"

    try:
        text = data.decode("utf-8", errors="ignore")
    except Exception:
        text = str(data)
    return clean_vtt(text)


def clean_vtt(text):
    """Strip WEBVTT timestamps/cue numbers so Teams .vtt transcripts read cleanly."""
    if "WEBVTT" not in text and "-->" not in text:
        return text
    out = []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if s == "WEBVTT":
            continue
        if "-->" in s:
            continue
        if re.fullmatch(r"\d+", s):
            continue
        out.append(s)
    return "\n".join(out)


# ---------------------------------------------------------------------------
# Helpers: AI analysis
# ---------------------------------------------------------------------------
CONFIDENTIAL_CATEGORIES = """- A named individual's performance, conduct, capability, discipline or PDP issues
- Health, medical, wellbeing or personal circumstances of named staff or students
- Industrial relations, union, grievance or complaint matters
- Restructure, redeployment, role security or freeze specifics affecting named roles, BEFORE official communication
- Anything pre-decisional / not yet endorsed or announced
- Anything that reads as OFFICIAL: Sensitive, legal, or privacy-bound"""


def build_prompt(agenda_text, transcript_text):
    has_agenda = bool(agenda_text.strip())
    agenda_block = (
        f"AGENDA (use this to structure the discussion breakdown):\n{agenda_text}\n\n"
        if has_agenda else
        "AGENDA: none supplied — derive the topics from the transcript.\n\n"
    )

    schema = """{
  "meeting_meta": {"date": "", "chair": "", "attendees": "", "apologies": ""},
  "summary": "3-6 sentence plain-English overview of the meeting",
  "agenda_items": [
    {"lead": "who led it", "topic": "topic name", "discussion": "what was discussed", "outcome": "decision/result or 'noted'"}
  ],
  "decisions": ["clear decisions made"],
  "actions": [
    {"owner": "person responsible", "action": "what to do", "due": "date/timeframe or 'TBC'"}
  ],
  "share_with_staff": [
    {
      "item": "short headline",
      "raw_note": "what was actually said (for the manager's eyes)",
      "staff_safe_wording": "a version with personal/confidential detail removed, ready to relay; or null if nothing can be safely shared",
      "classification": "share | hold | your_call",
      "reason": "one line explaining the classification",
      "audience": "all_staff | leadership | individual"
    }
  ]
}"""

    instructions = f"""You are an executive assistant for the Manager of an alternative education unit (LBu / Cowandilla Learning Centre) in the South Australian Department for Education. You are summarising her line-management ("LC leader") meeting.

{agenda_block}TRANSCRIPT:
{transcript_text}

Produce a faithful, concise summary. Rules:
- If an agenda was supplied, create one entry in "agenda_items" for EACH agenda topic, in order. If a topic wasn't actually discussed, set outcome to "not discussed". Add any extra topics raised (AoB) at the end.
- If no agenda was supplied, group the transcript into the natural topics discussed.
- "actions" = anything someone committed to do, or that was assigned. Capture the owner and any timeframe. Do not invent owners — use "TBC" if unclear.
- "decisions" = clear decisions made.
- "do not fabricate". If something is uncertain, keep it tentative.

SHARE-WITH-STAFF — classify, do not decide for her:
- List every item that *could* be relevant to communicate to her team. For EACH, set "classification":
  * "share"   = operationally useful to the team and clears all the confidential categories below.
  * "hold"    = touches ANY confidential category below. Default to hold when unsure.
  * "your_call" = genuinely depends on timing or whether it has been officially announced yet — flag for her judgment.
- Treat these categories as HOLD-by-default:
{CONFIDENTIAL_CATEGORIES}
- Always provide "staff_safe_wording": a version that strips names and confidential detail and conveys only what the team needs. If nothing can be safely shared, set it to null.
- "raw_note" is for the manager's record only and may contain the sensitive detail.
- Err on the side of caution: it is far better to mark something "hold" or "your_call" than to leak a personnel matter.

Return ONLY valid JSON matching this exact schema, with no prose, no markdown, no code fences:
{schema}"""
    return instructions


def _repair_truncated_json(t):
    """Best-effort recovery when the model's JSON was cut off mid-output."""
    s = t[t.find("{"):] if "{" in t else t
    stack, in_str, esc = [], False, False
    for ch in s:
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
        else:
            if ch == '"':
                in_str = True
            elif ch in "{[":
                stack.append("}" if ch == "{" else "]")
            elif ch in "}]":
                if stack:
                    stack.pop()
    repaired = s
    if in_str:                       # close an open string
        repaired += '"'
    repaired = re.sub(r"[,\s]+$", "", repaired)   # drop trailing comma/space
    repaired += "".join(reversed(stack))          # close open braces/brackets
    try:
        return json.loads(repaired)
    except Exception:
        return None


def parse_json_response(text):
    if not text:
        return None
    t = text.strip()
    t = re.sub(r"^```(?:json)?", "", t).strip()
    t = re.sub(r"```$", "", t).strip()
    start, end = t.find("{"), t.rfind("}")
    if start != -1 and end != -1 and end > start:
        candidate = t[start:end + 1]
    else:
        candidate = t
    try:
        return json.loads(candidate)
    except Exception:
        pass
    # fall back to repairing a truncated / partial object
    return _repair_truncated_json(t)


def analyse_meeting(agenda_text, transcript_text):
    client = get_anthropic_client()
    if client is None:
        return None, "Claude is not configured. Add ANTHROPIC_API_KEY to your app secrets."

    transcript_text = transcript_text.strip()
    truncated = False
    if len(transcript_text) > MAX_TRANSCRIPT_CHARS:
        transcript_text = transcript_text[:MAX_TRANSCRIPT_CHARS]
        truncated = True

    prompt = build_prompt(agenda_text, transcript_text)
    try:
        resp = client.messages.create(
            model=MODEL,
            max_tokens=8192,
            temperature=0.2,
            system="You return only valid JSON. No commentary, no markdown, no code fences.",
            messages=[{"role": "user", "content": prompt}],
        )
        raw = "".join(
            block.text for block in resp.content if getattr(block, "type", "") == "text"
        )
        stop_reason = getattr(resp, "stop_reason", None)
    except Exception as e:
        return None, f"AI request failed: {e}"

    if not raw.strip():
        return None, "Claude returned an empty response. Please try again."

    data = parse_json_response(raw)
    if data is None:
        # keep the raw reply so the page can show it for troubleshooting
        st.session_state["last_raw_response"] = raw
        if stop_reason == "max_tokens":
            return None, ("The response was cut off because the meeting produced more "
                          "output than the limit allowed. Try splitting a very long "
                          "transcript, or trimming obvious chit-chat, then run it again.")
        return None, ("Could not read Claude's response as structured data. It's saved "
                      "below so you can see what came back — usually re-running fixes it.")

    data.setdefault("meeting_meta", {})
    data.setdefault("summary", "")
    data.setdefault("agenda_items", [])
    data.setdefault("decisions", [])
    data.setdefault("actions", [])
    data.setdefault("share_with_staff", [])
    # normalise classification values
    for s in data["share_with_staff"]:
        c = str(s.get("classification", "your_call")).lower().strip()
        if c not in ("share", "hold", "your_call"):
            c = "your_call"
        s["classification"] = c
    note = "Transcript was long and was trimmed for analysis." if truncated else ""
    return data, note


# ---------------------------------------------------------------------------
# Helpers: Word export (mirrors the LBu LC leader proforma)
# ---------------------------------------------------------------------------
TEAL = RGBColor(0x00, 0x86, 0x7D) if DOCX_LIB else None


def _shade_cell(cell, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text, bold=False, size=10, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text) if text is not None else "")
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _heading(doc, text, color=None, size=14):
    h = doc.add_paragraph()
    hr = h.add_run(text)
    hr.bold = True
    hr.font.size = Pt(size)
    if color is not None:
        hr.font.color.rgb = color
    return h


def build_docx(data, meta, approved_indices):
    """approved_indices: set of share_with_staff indices the manager ticked to share."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    # OFFICIAL banner
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("OFFICIAL")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r.font.size = Pt(9)

    # Title
    title = doc.add_paragraph()
    tr = title.add_run("LBu LC leader — Meeting Summary")
    tr.bold = True
    tr.font.size = Pt(22)
    tr.font.color.rgb = TEAL

    mm = data.get("meeting_meta", {}) or {}
    mdate = meta.get("date") or mm.get("date") or ""
    chair = meta.get("chair") or mm.get("chair") or "Susan Furga"
    attendees = meta.get("attendees") or mm.get("attendees") or ""
    apologies = meta.get("apologies") or mm.get("apologies") or ""

    # Header info table
    info = doc.add_table(rows=3, cols=4)
    info.style = "Table Grid"
    rows = [
        ("Office", "Schools and Preschools", "Date of meeting", mdate),
        ("Directorate", "Priority Learners", "Unit", "LBu"),
        ("Chair", chair, "Location", "Microsoft Teams"),
    ]
    for i, (a, b, c, d) in enumerate(rows):
        _set_cell_text(info.rows[i].cells[0], a, bold=True)
        _set_cell_text(info.rows[i].cells[1], b)
        _set_cell_text(info.rows[i].cells[2], c, bold=True)
        _set_cell_text(info.rows[i].cells[3], d)

    doc.add_paragraph("")

    att = doc.add_table(rows=2, cols=2)
    att.style = "Table Grid"
    _set_cell_text(att.rows[0].cells[0], "Attendees", bold=True)
    _set_cell_text(att.rows[0].cells[1], attendees)
    _set_cell_text(att.rows[1].cells[0], "Apologies", bold=True)
    _set_cell_text(att.rows[1].cells[1], apologies)

    # Summary
    _heading(doc, "Summary", TEAL)
    doc.add_paragraph(data.get("summary", ""))

    # Discussion by item
    items = data.get("agenda_items", []) or []
    if items:
        _heading(doc, "Discussion by item", TEAL)
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        for c, label in zip(t.rows[0].cells, ["Lead", "Topic", "Discussion & outcome"]):
            _set_cell_text(c, label, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            _shade_cell(c, "00867D")
        for it in items:
            row = t.add_row().cells
            _set_cell_text(row[0], it.get("lead", ""))
            _set_cell_text(row[1], it.get("topic", ""), bold=True)
            disc = it.get("discussion", "")
            outc = it.get("outcome", "")
            _set_cell_text(row[2], disc + (f"\nOutcome: {outc}" if outc else ""))

    # Decisions
    decisions = data.get("decisions", []) or []
    if decisions:
        _heading(doc, "Decisions", TEAL)
        for d in decisions:
            doc.add_paragraph(d, style="List Bullet")

    # Actions
    actions = data.get("actions", []) or []
    _heading(doc, "Actions", TEAL)
    if actions:
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        for c, label in zip(t.rows[0].cells, ["Owner", "Action", "Due"]):
            _set_cell_text(c, label, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            _shade_cell(c, "00867D")
        for a in actions:
            row = t.add_row().cells
            _set_cell_text(row[0], a.get("owner", ""))
            _set_cell_text(row[1], a.get("action", ""))
            _set_cell_text(row[2], a.get("due", "TBC"))
    else:
        doc.add_paragraph("No actions recorded.")

    # Split share items into approved (to share) vs withheld (manager record)
    share = data.get("share_with_staff", []) or []
    approved = [(i, s) for i, s in enumerate(share) if i in approved_indices]
    withheld = [(i, s) for i, s in enumerate(share) if i not in approved_indices]

    # To share with staff (approved only)
    _heading(doc, "To share with staff", TEAL)
    if approved:
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        for c, label in zip(t.rows[0].cells, ["Item", "What to tell the team"]):
            _set_cell_text(c, label, bold=True)
            _shade_cell(c, "FFE08A")
        for _, s in approved:
            row = t.add_row().cells
            _set_cell_text(row[0], s.get("item", ""), bold=True)
            wording = s.get("staff_safe_wording") or s.get("item", "")
            _set_cell_text(row[1], wording)
            _shade_cell(row[0], "FFF6D6")
            _shade_cell(row[1], "FFF6D6")
    else:
        doc.add_paragraph("Nothing approved to share with staff.")

    # Confidential — manager record only
    if withheld:
        doc.add_paragraph("")
        sp = doc.add_paragraph()
        spr = sp.add_run("OFFICIAL: Sensitive — manager record only · do not circulate")
        spr.bold = True
        spr.font.size = Pt(9)
        spr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        _heading(doc, "Confidential — not shared", RGBColor(0x80, 0x80, 0x80), size=13)
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        for c, label in zip(t.rows[0].cells, ["Item", "Note", "Why held"]):
            _set_cell_text(c, label, bold=True)
            _shade_cell(c, "EDEDED")
        for _, s in withheld:
            row = t.add_row().cells
            _set_cell_text(row[0], s.get("item", ""), bold=True)
            _set_cell_text(row[1], s.get("raw_note") or "")
            tag = s.get("classification", "")
            _set_cell_text(row[2], f"[{tag}] {s.get('reason','')}")

    # Footer
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("OFFICIAL")
    fr.bold = True
    fr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    fr.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Helpers: Supabase persistence (optional)
# ---------------------------------------------------------------------------
def save_meeting(meta, data, approved_indices, week_label=None):
    sb = get_supabase()
    if sb is None:
        return False, "Supabase not configured — saved to this session only."
    try:
        payload = dict(data)
        payload["_approved_indices"] = sorted(list(approved_indices))
        payload["_week_label"] = week_label or ""
        sb.table(SUPABASE_TABLE).insert({
            "meeting_date": meta.get("date") or None,
            "week_label": week_label or None,
            "chair": meta.get("chair") or None,
            "summary": data.get("summary", ""),
            "payload": payload,
            "created_at": datetime.utcnow().isoformat(),
        }).execute()
        return True, "Saved to the archive."
    except Exception as e:
        return False, f"Could not save: {e}"


def load_recent_meetings(limit=25):
    sb = get_supabase()
    if sb is None:
        return []
    try:
        res = (sb.table(SUPABASE_TABLE)
               .select("id, meeting_date, week_label, summary, created_at")
               .order("created_at", desc=True)
               .limit(limit).execute())
        return res.data or []
    except Exception:
        return []


def load_meeting_full(meeting_id):
    """Return the full saved row (incl. payload) for one archived meeting."""
    sb = get_supabase()
    if sb is None:
        return None
    try:
        res = (sb.table(SUPABASE_TABLE)
               .select("*")
               .eq("id", meeting_id)
               .limit(1).execute())
        rows = res.data or []
        return rows[0] if rows else None
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Archive password gate
# ---------------------------------------------------------------------------
def get_archive_password():
    """Password for viewing archived meetings. Set ARCHIVE_PASSWORD in secrets."""
    return get_secret("ARCHIVE_PASSWORD") or "change-me"


def parse_meeting_date(s):
    """Parse a meeting-date string (various formats) into a date; fall back to today."""
    if s:
        for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d %B %Y", "%d %b %Y"):
            try:
                return datetime.strptime(str(s).strip(), fmt).date()
            except Exception:
                continue
    return date.today()


def suggest_week():
    """Default the term week to the last saved week + 2 (meetings are fortnightly)."""
    try:
        recent = load_recent_meetings(limit=1)
        if recent:
            label = recent[0].get("week_label") or ""
            m = re.search(r"\d+", label)
            if m:
                return min(int(m.group()) + 2, 11)
    except Exception:
        pass
    return 8


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------
require_confidential_auth("LC Meeting Summariser", "🗒️")

st.title("🗒️ LC Meeting Summariser")
st.caption("Drop in your agenda (optional) and the transcript. Get a summary, "
           "actions, and a classified share-with-staff list you approve before export.")

missing = []
if get_anthropic_client() is None:
    missing.append("Claude (AI) — add **ANTHROPIC_API_KEY** to secrets")
if not DOCX_LIB:
    missing.append("python-docx — add `python-docx` to requirements.txt")
if missing:
    st.warning("Some features need setup:\n\n- " + "\n- ".join(missing))

# Session seeds (set BEFORE widgets are instantiated)
if "seed_agenda" not in st.session_state:
    st.session_state.seed_agenda = ""
if "result" not in st.session_state:
    st.session_state.result = None
if "result_note" not in st.session_state:
    st.session_state.result_note = ""

left, right = st.columns(2)

with left:
    st.subheader("1 · Agenda (optional)")
    a_up = st.file_uploader("Upload agenda (.docx / .txt)", type=["docx", "txt", "md"],
                            key="agenda_upload")
    if st.button("↩ Load this week's agenda (19.06.2026)"):
        st.session_state.seed_agenda = SAMPLE_AGENDA
        st.rerun()
    agenda_text_pasted = st.text_area(
        "…or paste the agenda here",
        value=st.session_state.seed_agenda, height=220, key="agenda_paste",
        placeholder="Leave blank if you don't have an agenda — the transcript is enough.")

with right:
    st.subheader("2 · Transcript (required)")
    t_up = st.file_uploader("Upload transcript (.docx / .txt / .vtt)",
                            type=["docx", "txt", "vtt", "md"], key="transcript_upload")
    transcript_text_pasted = st.text_area(
        "…or paste the transcript here", height=220, key="transcript_paste",
        placeholder="Paste the Teams transcript or your notes.")

agenda_text = read_upload(a_up) if a_up else agenda_text_pasted
transcript_text = read_upload(t_up) if t_up else transcript_text_pasted

c1, c2, _ = st.columns([1, 1, 4])
with c1:
    go = st.button("✨ Summarise meeting", type="primary")
with c2:
    if st.button("🗑 Clear"):
        for k in list(st.session_state.keys()):
            if k.startswith("share_ok_"):
                del st.session_state[k]
        st.session_state.seed_agenda = ""
        st.session_state.result = None
        st.session_state.result_note = ""
        st.rerun()

if go:
    if not transcript_text.strip():
        st.error("Please add a transcript (upload or paste) before summarising.")
    else:
        # clear any stale approval state from a previous run
        for k in list(st.session_state.keys()):
            if k.startswith("share_ok_"):
                del st.session_state[k]
        with st.spinner("Reading the meeting and pulling out actions…"):
            data, note = analyse_meeting(agenda_text, transcript_text)
        if data is None:
            st.error(note)
            raw_dbg = st.session_state.get("last_raw_response")
            if raw_dbg:
                with st.expander("See what Claude returned (for troubleshooting)"):
                    st.code(raw_dbg)
        else:
            st.session_state.pop("last_raw_response", None)
            st.session_state.result = data
            st.session_state.result_note = note

# ---------------------------------------------------------------------------
# Results
# ---------------------------------------------------------------------------
data = st.session_state.result
if data:
    if st.session_state.result_note:
        st.info(st.session_state.result_note)

    st.divider()
    st.header("Summary")
    st.write(data.get("summary", ""))

    items = data.get("agenda_items", []) or []
    if items:
        st.header("Discussion by item")
        for it in items:
            label = f"📌 {it.get('topic','(untitled)')}  ·  {it.get('lead','')}".strip(" ·")
            with st.expander(label):
                st.markdown(f"**Discussion:** {it.get('discussion','')}")
                if it.get("outcome"):
                    st.markdown(f"**Outcome:** {it.get('outcome')}")

    decisions = data.get("decisions", []) or []
    if decisions:
        st.header("Decisions")
        for d in decisions:
            st.markdown(f"- {d}")

    st.header("✅ Actions")
    actions = data.get("actions", []) or []
    if actions:
        st.table([
            {"Owner": a.get("owner", ""),
             "Action": a.get("action", ""),
             "Due": a.get("due", "TBC")}
            for a in actions
        ])
    else:
        st.caption("No actions recorded.")

    # ---- Share with staff: classified + approve before export ----
    st.header("📣 To share with staff")
    st.caption("Nothing leaves this tool until you tick it. Held items stay in a "
               "confidential, manager-only section of the export.")
    share = data.get("share_with_staff", []) or []

    if not share:
        st.caption("Nothing flagged as potentially shareable.")
    else:
        buckets = {"share": [], "your_call": [], "hold": []}
        for i, s in enumerate(share):
            buckets.get(s.get("classification", "your_call"), buckets["your_call"]).append((i, s))

        def render_item(i, s, default_tick, accent, bg):
            wording = s.get("staff_safe_wording")
            st.markdown(
                f"""<div style="background:{bg};border-left:5px solid {accent};
                padding:10px 14px;border-radius:6px;margin:6px 0;">
                <strong>{s.get('item','')}</strong>
                <span style="float:right;font-size:0.8em;color:#555;">
                {s.get('classification','').replace('_',' ')} · {s.get('audience','')}</span>
                <br><span style="font-size:0.9em;color:#444;">{s.get('reason','')}</span>
                {f"<br><em>Staff-safe wording:</em> {wording}" if wording else
                 "<br><em>No safe wording available — relay verbally if appropriate.</em>"}
                </div>""",
                unsafe_allow_html=True,
            )
            return st.checkbox("Include in staff section", value=default_tick,
                               key=f"share_ok_{i}")

        if buckets["share"]:
            st.subheader("✅ Recommended to share")
            for i, s in buckets["share"]:
                render_item(i, s, True, "#2E9E5B", "#E9F8EF")

        if buckets["your_call"]:
            st.subheader("🟡 Your call")
            for i, s in buckets["your_call"]:
                render_item(i, s, False, "#E0A500", "#FFF6D6")

        if buckets["hold"]:
            st.subheader("🔒 Hold — confidential")
            st.caption("Default off. These touch personnel/confidential categories. "
                       "Only override if you're certain it's appropriate.")
            for i, s in buckets["hold"]:
                render_item(i, s, False, "#C0392B", "#FBEAEA")

    st.divider()
    meta = {
        "date": (data.get("meeting_meta", {}) or {}).get("date", ""),
        "chair": (data.get("meeting_meta", {}) or {}).get("chair", "Susan Furga"),
        "attendees": (data.get("meeting_meta", {}) or {}).get("attendees", ""),
        "apologies": (data.get("meeting_meta", {}) or {}).get("apologies", ""),
    }

    st.markdown("**Meeting details for the archive**")
    st.caption("Meetings are fortnightly — the week defaults to the last saved week + 2.")
    d1, d2 = st.columns(2)
    with d1:
        week_num = st.number_input("Term week", min_value=1, max_value=11,
                                   value=suggest_week(), step=1)
    with d2:
        meeting_dt = st.date_input("Meeting date",
                                   value=parse_meeting_date(meta["date"]),
                                   format="DD.MM.YYYY")
    week_label = f"Week {int(week_num)}"
    meta["date"] = meeting_dt.strftime("%d.%m.%Y")

    e1, e2 = st.columns([1, 1])
    # current approval set, read from checkbox state
    approved_indices = {
        i for i in range(len(share))
        if st.session_state.get(f"share_ok_{i}", False)
    }

    with e1:
        if DOCX_LIB:
            buf = build_docx(data, meta, approved_indices)
            stamp = f"{week_label.replace(' ', '')}_{meta['date'].replace('.', '-')}"
            st.download_button(
                "⬇ Download Word summary", data=buf,
                file_name=f"LC_meeting_summary_{stamp}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            st.caption(f"{len(approved_indices)} item(s) will appear in the staff section.")
        else:
            st.caption("Add `python-docx` to requirements.txt to enable Word export.")
    with e2:
        if st.button("💾 Save to archive"):
            ok, msg = save_meeting(meta, data, approved_indices, week_label)
            (st.success if ok else st.info)(msg)

# ---------------------------------------------------------------------------
# Archived meetings — password protected
# ---------------------------------------------------------------------------
st.divider()
st.header("🔒 Archived meetings")

if get_supabase() is None:
    st.caption("Saving/archiving is off — add your Supabase keys to enable it.")
else:
    if "archive_unlocked" not in st.session_state:
        st.session_state.archive_unlocked = False

    if not st.session_state.archive_unlocked:
        st.caption("These records can include confidential, manager-only notes. "
                   "Enter the access code to view them.")
        pw = st.text_input("Access code", type="password", key="archive_pw")
        cols = st.columns([1, 5])
        with cols[0]:
            if st.button("Unlock"):
                if pw and pw == get_archive_password():
                    st.session_state.archive_unlocked = True
                    st.rerun()
                else:
                    st.error("Incorrect access code.")
        if get_archive_password() == "change-me":
            st.warning("No archive password set. Add **ARCHIVE_PASSWORD** to your "
                       "app secrets — until then the code is `change-me`.")
    else:
        top = st.columns([1, 5])
        with top[0]:
            if st.button("🔒 Lock"):
                st.session_state.archive_unlocked = False
                st.rerun()

        meetings = load_recent_meetings()
        if not meetings:
            st.caption("No saved meetings yet.")
        else:
            def label_for(m):
                when = m.get("meeting_date") or (m.get("created_at", "")[:10])
                wk = m.get("week_label")
                prefix = f"{wk} · {when}" if wk else when
                snip = (m.get("summary", "") or "")[:55]
                return f"{prefix} — {snip}…"

            options = {label_for(m): m["id"] for m in meetings}
            choice = st.selectbox("Select a meeting", list(options.keys()))
            row = load_meeting_full(options[choice]) if choice else None

            if row:
                payload = row.get("payload", {}) or {}
                approved = set(payload.get("_approved_indices", []))
                stored_date = row.get("meeting_date") or (payload.get("meeting_meta", {}) or {}).get("date", "")
                wk = row.get("week_label") or payload.get("_week_label", "")
                a_meta = {
                    "date": stored_date,
                    "chair": (payload.get("meeting_meta", {}) or {}).get("chair", "Susan Furga"),
                    "attendees": (payload.get("meeting_meta", {}) or {}).get("attendees", ""),
                    "apologies": (payload.get("meeting_meta", {}) or {}).get("apologies", ""),
                }

                header = " · ".join(x for x in [wk, stored_date] if x)
                if header:
                    st.markdown(f"### {header}")
                st.subheader("Summary")
                st.write(payload.get("summary", ""))

                a_actions = payload.get("actions", []) or []
                if a_actions:
                    st.subheader("Actions")
                    st.table([
                        {"Owner": a.get("owner", ""),
                         "Action": a.get("action", ""),
                         "Due": a.get("due", "TBC")}
                        for a in a_actions
                    ])

                a_share = payload.get("share_with_staff", []) or []
                shared = [s for i, s in enumerate(a_share) if i in approved]
                held = [s for i, s in enumerate(a_share) if i not in approved]
                if shared:
                    st.subheader("📣 Shared with staff")
                    for s in shared:
                        st.markdown(f"- **{s.get('item','')}** — "
                                    f"{s.get('staff_safe_wording') or s.get('item','')}")
                if held:
                    st.subheader("🔒 Confidential — not shared")
                    for s in held:
                        st.markdown(
                            f"""<div style="background:#FBEAEA;border-left:5px solid #C0392B;
                            padding:8px 12px;border-radius:6px;margin:5px 0;">
                            <strong>{s.get('item','')}</strong><br>
                            {s.get('raw_note') or ''}<br>
                            <em style="font-size:0.85em;color:#555;">Why held: {s.get('reason','')}</em>
                            </div>""",
                            unsafe_allow_html=True,
                        )

                if DOCX_LIB:
                    buf2 = build_docx(payload, a_meta, approved)
                    stamp2 = (a_meta["date"] or row.get("created_at", "")[:10] or
                              date.today().isoformat()).replace(".", "-").replace("/", "-")
                    st.download_button(
                        "⬇ Re-download Word summary", data=buf2,
                        file_name=f"LC_meeting_summary_{stamp2}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key="archive_download")
