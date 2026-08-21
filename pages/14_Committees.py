import streamlit as st
from supabase import create_client
from datetime import datetime, date
import json
import re
from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Page config
st.set_page_config(page_title="Committees", page_icon="📋", layout="wide")

# Hide sidebar
st.markdown("""
    <style>
    [data-testid="stSidebar"] {display: none;}
    </style>
""", unsafe_allow_html=True)

# Initialize Supabase
supabase = create_client(
    st.secrets["supabase"]["url"],
    st.secrets["supabase"]["key"]
)

# ══════════════════════════════════════════════════════════════════════════
# RESILIENT GROQ AI HELPER — portal-wide standard (Aug 2026)
#
# Fixes the recurring portal-wide Groq issue, which has three stacked causes:
#   1) A stale groq/httpx combo can raise
#      "Client.__init__() got an unexpected keyword argument 'proxies'" on
#      client creation — worked around by retrying with an explicit http_client.
#   2) Groq retired "llama-3.3-70b-versatile" for free/dev-tier accounts
#      (16 Aug 2026), so any call hardcoded to that model can start failing
#      with a model-not-found/decommissioned error — worked around by trying
#      a small chain of current models instead of one hardcoded name.
#   3) qwen/qwen3.6-27b is a hybrid-thinking model — left to its defaults it
#      spends the whole max_tokens budget on a <think> chain-of-thought block
#      before ever writing the actual answer, which shows up as "Extra data"/
#      JSON-parse errors or empty/truncated replies. Fixed by scoping
#      reasoning_effort="none" to that model only — gpt-oss and llama don't
#      accept that parameter, so it must NOT be sent to them.
# It also self-heals 413 "request too large" errors by halving max_tokens
# and retrying, instead of failing outright.
#
# Use groq_ai_call(prompt, ...) anywhere in a tool instead of building a
# Groq client or calling chat.completions.create directly.
# ══════════════════════════════════════════════════════════════════════════
_GROQ_MODEL_CHAIN = ["openai/gpt-oss-120b", "qwen/qwen3.6-27b", "llama-3.3-70b-versatile"]
def _get_groq_client(api_key):
    from groq import Groq
    try:
        return Groq(api_key=api_key)
    except TypeError:
        import httpx
        return Groq(api_key=api_key, http_client=httpx.Client())
def groq_ai_call(prompt, max_tokens=1500, temperature=0.1, api_key=None, system=None):
    """Resilient Groq chat-completion call. Returns the response text (str).
    Raises the last error encountered if every model/retry combination fails
    — wrap the call in try/except at the call site if you want a graceful
    fallback instead of Streamlit's default traceback."""
    api_key = api_key or st.secrets.get("GROQ_API_KEY", "")
    if not api_key:
        raise RuntimeError("Groq API key not configured.")
    client = _get_groq_client(api_key)
    messages = []
    if system:
        messages.append({"role": "system", "content": system})
    messages.append({"role": "user", "content": prompt})
    last_err = None
    for model in _GROQ_MODEL_CHAIN:
        tokens = max_tokens
        for _attempt in range(2):  # original size, then one halved retry
            try:
                kwargs = dict(model=model, messages=messages, max_tokens=tokens, temperature=temperature)
                if model == "qwen/qwen3.6-27b":
                    kwargs["reasoning_effort"] = "none"
                resp = client.chat.completions.create(**kwargs)
                return resp.choices[0].message.content.strip()
            except Exception as e:
                last_err = e
                msg = str(e).lower()
                if "413" in msg or "too large" in msg:
                    tokens = max(256, tokens // 2)
                    continue  # retry same model with a smaller token budget
                break  # any other error: move on to the next model in the chain
    raise last_err


# Initialize session state
if 'committee_view' not in st.session_state:
    st.session_state.committee_view = 'home'
if 'selected_committee' not in st.session_state:
    st.session_state.selected_committee = None
if 'selected_meeting' not in st.session_state:
    st.session_state.selected_meeting = None

# Navigation handler
def navigate_to(view, committee=None, meeting=None):
    st.session_state.committee_view = view
    if committee:
        st.session_state.selected_committee = committee
    if meeting:
        st.session_state.selected_meeting = meeting
    st.rerun()

# Committee configurations
COMMITTEES = {
    'finance': {
        'name': 'Finance Committee',
        'icon': '💰',
        'color': '#059669',
        'description': 'Budget · Expenditure · Reports'
    },
    'whs': {
        'name': 'WHS Committee',
        'icon': '🦺',
        'color': '#ea580c',
        'description': 'Safety · Hazards · Compliance'
    },
    'social': {
        'name': 'Social Club',
        'icon': '🎉',
        'color': '#7c3aed',
        'description': 'Events · Activities · Fun'
    },
    'special': {
        'name': 'Special Events',
        'icon': '✨',
        'color': '#ec4899',
        'description': 'Planning · Coordination · Celebration'
    },
    'personel': {
        'name': 'Personal Advisory Committee',
        'icon': '🧭',
        'color': '#2563eb',
        'description': 'Guidance · Support · Advocacy'
    }
}

# ========== MINUTES DOCX GENERATOR ==========

def _hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip('#')
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def _set_cell_shading(cell, hex_color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color.lstrip('#'))
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side, val in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def _add_bottom_border(paragraph, color="e2e8f0", size=6):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_page_number_field(paragraph, grey_rgb):
    run = paragraph.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = grey_rgb
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = "PAGE"
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)

    run2 = paragraph.add_run(" of ")
    run2.font.size = Pt(9)
    run2.font.color.rgb = grey_rgb

    run3 = paragraph.add_run()
    run3.font.size = Pt(9)
    run3.font.color.rgb = grey_rgb
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'begin')
    instr2 = OxmlElement('w:instrText')
    instr2.set(qn('xml:space'), 'preserve')
    instr2.text = "NUMPAGES"
    fld4 = OxmlElement('w:fldChar')
    fld4.set(qn('w:fldCharType'), 'end')
    run3._r.append(fld3)
    run3._r.append(instr2)
    run3._r.append(fld4)


_HEADING_PATTERN = re.compile(r'^\d+\.\s+[A-Z][A-Z\s&/\'\-]+$')
_ITEM_HEADING_PATTERN = re.compile(r'^\d+\.\s+.+$')
_LABEL_PATTERN = re.compile(r'^(Present|Apologies|Status|Discussion|Actions):\s*(.*)$', re.IGNORECASE)


def _parse_minutes_body(text):
    """Return the lines of the minutes text starting at the first
    ALL-CAPS numbered section heading (e.g. '1. ATTENDANCE'), skipping
    the raw metadata header that's already shown in the meta table."""
    lines = text.split('\n')
    start = 0
    for i, line in enumerate(lines):
        if _HEADING_PATTERN.match(line.strip()):
            start = i
            break
    return lines[start:]


def generate_minutes_docx(meeting, committee, formatted_minutes_text):
    """Build a professionally branded Word document for a set of committee
    meeting minutes and return it as a BytesIO buffer ready for download."""
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    accent_hex = committee['color']
    accent_rgb = _hex_to_rgb(accent_hex)
    navy_rgb = RGBColor(0x1e, 0x3a, 0x5f)
    grey_rgb = RGBColor(0x64, 0x74, 0x8b)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x1f, 0x29, 0x37)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    content_width = Cm(21.0) - section.left_margin - section.right_margin

    # ---- Letterhead ----
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    hcell = header_table.rows[0].cells[0]
    hcell.width = content_width
    _set_cell_shading(hcell, "1e3a5f")
    _set_cell_margins(hcell, top=200, bottom=200, left=300, right=300)
    hp = hcell.paragraphs[0]
    run = hp.add_run("COWANDILLA LEARNING CENTRE")
    run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
    run.font.size = Pt(15)
    run.font.bold = True
    hp2 = hcell.add_paragraph()
    run2 = hp2.add_run("Learning and Behaviour Unit")
    run2.font.color.rgb = RGBColor(0xcb, 0xd5, 0xe1)
    run2.font.size = Pt(10.5)
    run2.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ---- Title block ----
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(0)
    trun = title_p.add_run("MEETING MINUTES")
    trun.font.size = Pt(20)
    trun.font.bold = True
    trun.font.color.rgb = navy_rgb

    committee_p = doc.add_paragraph()
    committee_p.paragraph_format.space_after = Pt(10)
    crun = committee_p.add_run(committee['name'])
    crun.font.size = Pt(13)
    crun.font.bold = True
    crun.font.color.rgb = accent_rgb
    _add_bottom_border(committee_p, color=accent_hex.lstrip('#'), size=10)

    # ---- Meta table ----
    meta_items = [
        ("Meeting", meeting['meeting_title']),
        ("Date", meeting['meeting_date']),
        ("Time", meeting['meeting_time']),
        ("Location", meeting.get('location') or 'TBA'),
        ("Status", meeting.get('status', 'draft').capitalize()),
    ]
    meta_table = doc.add_table(rows=len(meta_items), cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_table.autofit = False
    label_w = Cm(3.6)
    value_w = content_width - label_w
    for r, (label, value) in enumerate(meta_items):
        lcell = meta_table.cell(r, 0)
        vcell = meta_table.cell(r, 1)
        lcell.width = label_w
        vcell.width = value_w
        _set_cell_shading(lcell, "f1f5f9")
        _set_cell_margins(lcell, top=80, bottom=80, left=150, right=150)
        _set_cell_margins(vcell, top=80, bottom=80, left=150, right=150)
        lp = lcell.paragraphs[0]
        lrun = lp.add_run(label)
        lrun.font.bold = True
        lrun.font.size = Pt(9.5)
        lrun.font.color.rgb = grey_rgb
        vp = vcell.paragraphs[0]
        vrun = vp.add_run(str(value))
        vrun.font.size = Pt(10.5)
        vcell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lcell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- Body ----
    body_lines = _parse_minutes_body(formatted_minutes_text)
    blank_run = False
    for raw_line in body_lines:
        line = raw_line.strip()

        if not line:
            blank_run = True
            continue

        if _HEADING_PATTERN.match(line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(line)
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.color.rgb = navy_rgb
            _add_bottom_border(p)
            blank_run = False
            continue

        label_match = _LABEL_PATTERN.match(line)
        if label_match:
            label, rest = label_match.group(1), label_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            lr = p.add_run(f"{label}: ")
            lr.font.bold = True
            lr.font.size = Pt(10.5)
            if rest:
                vr = p.add_run(rest)
                vr.font.size = Pt(10.5)
            blank_run = False
            continue

        if _ITEM_HEADING_PATTERN.match(line) and len(line) < 90 and not line.isupper():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(line)
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = accent_rgb
            blank_run = False
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6 if blank_run else 3)
        r = p.add_run(line)
        r.font.size = Pt(10.5)
        blank_run = False

    # ---- Footer ----
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = ""
    footer_p.paragraph_format.tab_stops.add_tab_stop(content_width, WD_TAB_ALIGNMENT.RIGHT)
    pPr = footer_p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '4')
    top.set(qn('w:space'), '4')
    top.set(qn('w:color'), 'e2e8f0')
    pBdr.append(top)
    pPr.append(pBdr)

    fr1 = footer_p.add_run("Cowandilla Learning Centre — Learning and Behaviour Unit")
    fr1.font.size = Pt(8.5)
    fr1.font.color.rgb = grey_rgb
    footer_p.add_run("\t")
    _add_page_number_field(footer_p, grey_rgb)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ========== HEADER ==========
st.markdown("""
    <div style="background: linear-gradient(135deg, #1e3a5f 0%, #2d5a8c 100%); 
                padding: 2rem; border-radius: 12px; margin-bottom: 2rem; 
                box-shadow: 0 4px 6px rgba(0,0,0,0.1);">
        <div style="display: flex; align-items: center; gap: 1rem;">
            <div style="font-size: 3rem;">🏛️</div>
            <div>
                <h1 style="margin: 0; color: white; font-size: 2.2rem;">Committees</h1>
                <p style="margin: 0; color: #e0e7ff; font-size: 1rem;">Cowandilla Learning Centre</p>
            </div>
        </div>
    </div>
""", unsafe_allow_html=True)

# ========== HOME VIEW ==========
if st.session_state.committee_view == 'home':
    
    st.markdown("<h3 style='color: #1e3a5f; margin-bottom: 1.5rem;'>Select a committee to manage meetings and documents.</h3>", unsafe_allow_html=True)
    
    # Badge colors (light bg / dark text) per committee, derived to complement each accent color
    BADGE_COLORS = {
        'finance': ('#d1fae5', '#065f46'),
        'whs': ('#fef3c7', '#92400e'),
        'social': ('#e9d5ff', '#6b21a8'),
        'special': ('#fce7f3', '#9f1239'),
        'personel': ('#dbeafe', '#1e40af'),
    }

    def render_committee_card(key, info):
        badge_bg, badge_text = BADGE_COLORS.get(key, ('#e2e8f0', '#334155'))
        st.markdown(f"""
            <div style="border: 3px solid {info['color']}; border-radius: 16px; padding: 2rem; 
                        background: white; margin-bottom: 1.5rem; min-height: 220px;
                        box-shadow: 0 2px 4px rgba(0,0,0,0.05);">
                <div style="font-size: 3.5rem; margin-bottom: 1rem;">{info['icon']}</div>
                <h3 style="color: #1e3a5f; margin-bottom: 0.5rem;">{info['name']}</h3>
                <p style="color: #64748b; margin-bottom: 1rem;">{info['description']}</p>
                <span style="background: {badge_bg}; color: {badge_text}; padding: 0.4rem 0.8rem; 
                            border-radius: 20px; font-size: 0.85rem; font-weight: 600;">● Active</span>
            </div>
        """, unsafe_allow_html=True)
        if st.button(f"Open {info['name']}", key=f"btn_{key}", use_container_width=True):
            navigate_to('committee_home', key)

    # Committee Cards - 3 per row
    committee_keys = list(COMMITTEES.keys())
    for row_start in range(0, len(committee_keys), 3):
        row_keys = committee_keys[row_start:row_start + 3]
        cols = st.columns(3)
        for col, key in zip(cols, row_keys):
            with col:
                render_committee_card(key, COMMITTEES[key])
    
    # Documents Card - Centered below
    st.divider()
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        st.markdown("""
            <div style="border: 3px solid #0891b2; border-radius: 16px; padding: 2rem; 
                        background: white; margin-bottom: 1.5rem; min-height: 220px;
                        box-shadow: 0 2px 4px rgba(0,0,0,0.05);">
                <div style="font-size: 3.5rem; margin-bottom: 1rem;">📄</div>
                <h3 style="color: #1e3a5f; margin-bottom: 0.5rem;">Documents Library</h3>
                <p style="color: #64748b; margin-bottom: 1rem;">Policies · Procedures · Resources</p>
                <span style="background: #cffafe; color: #0e7490; padding: 0.4rem 0.8rem; 
                            border-radius: 20px; font-size: 0.85rem; font-weight: 600;">● Active</span>
            </div>
        """, unsafe_allow_html=True)
        if st.button("📂 Open Documents Library", key="btn_docs", use_container_width=True):
            navigate_to('documents')

# ========== COMMITTEE HOME VIEW ==========
elif st.session_state.committee_view == 'committee_home':
    committee = COMMITTEES[st.session_state.selected_committee]
    
    st.subheader(f"{committee['icon']} {committee['name']}")
    
    # Tab navigation
    tab1, tab2 = st.tabs(["📅 Meetings", "➕ New Meeting"])
    
    with tab1:
        st.markdown("### Meetings")
        
        try:
            # Fetch meetings for this committee
            response = supabase.table("committee_meetings").select("*").eq(
                "committee_name", st.session_state.selected_committee
            ).order("meeting_date", desc=True).execute()
            
            if response.data:
                for meeting in response.data:
                    with st.container():
                        col1, col2, col3, col4 = st.columns([3, 2, 1.5, 1.5])
                        
                        with col1:
                            st.markdown(f"**{meeting['meeting_title']}**")
                            st.caption(f"📅 {meeting['meeting_date']}")
                        
                        with col2:
                            status_colors = {
                                'draft': ('#fef3c7', '#92400e'),
                                'published': ('#d1fae5', '#065f46'),
                                'archived': ('#e5e7eb', '#6b7280')
                            }
                            bg, text = status_colors.get(meeting['status'], ('#e5e7eb', '#6b7280'))
                            st.markdown(f"<span style='background:{bg};color:{text};padding:0.25rem 0.6rem;border-radius:12px;font-size:0.8rem;font-weight:600;'>{meeting['status'].upper()}</span>", unsafe_allow_html=True)
                        
                        with col3:
                            if st.button("View Meeting", key=f"view_{meeting['id']}", use_container_width=True):
                                navigate_to('meeting_detail', meeting=meeting['id'])
                        
                        with col4:
                            if meeting['status'] != 'archived':
                                if st.button("Archive", key=f"archive_{meeting['id']}", use_container_width=True):
                                    supabase.table("committee_meetings").update({"status": "archived"}).eq("id", meeting['id']).execute()
                                    st.rerun()
                        
                        st.divider()
            else:
                st.info("No meetings yet. Create your first meeting using the 'New Meeting' tab.")
        
        except Exception as e:
            st.error(f"Error loading meetings: {str(e)}")
    
    with tab2:
        st.markdown("### Create New Meeting")
        
        with st.form("new_meeting_form"):
            meeting_title = st.text_input("Meeting Title*", placeholder="e.g., Finance Committee Meeting - May 2026")
            meeting_date = st.date_input("Meeting Date*", value=date.today())
            meeting_time = st.time_input("Meeting Time*")
            location = st.text_input("Location", placeholder="e.g., Staff Room")
            
            st.markdown("#### Agenda Items")
            num_items = st.number_input("Number of agenda items", min_value=1, max_value=20, value=5)
            
            agenda_items = []
            for i in range(num_items):
                st.markdown(f"**Item {i+1}**")
                col1, col2 = st.columns([3, 1])
                with col1:
                    item_title = st.text_input(f"Title", key=f"item_title_{i}", placeholder="e.g., Budget Review")
                with col2:
                    duration = st.number_input(f"Duration (min)", key=f"duration_{i}", min_value=5, value=10)
                item_description = st.text_area(f"Description", key=f"item_desc_{i}", placeholder="Brief description...")
                
                if item_title:
                    agenda_items.append({
                        'item_number': i + 1,
                        'item_title': item_title,
                        'item_description': item_description,
                        'duration_minutes': duration
                    })
            
            submit = st.form_submit_button("Create Meeting", use_container_width=True)
            
            if submit:
                if not meeting_title:
                    st.error("Please provide a meeting title.")
                else:
                    try:
                        # Create meeting
                        meeting_response = supabase.table("committee_meetings").insert({
                            "committee_name": st.session_state.selected_committee,
                            "meeting_title": meeting_title,
                            "meeting_date": str(meeting_date),
                            "meeting_time": str(meeting_time),
                            "location": location,
                            "status": "draft",
                            "created_at": datetime.now().isoformat()
                        }).execute()
                        
                        meeting_id = meeting_response.data[0]['id']
                        
                        # Create agenda items
                        if agenda_items:
                            for item in agenda_items:
                                item['meeting_id'] = meeting_id
                            supabase.table("committee_agenda_items").insert(agenda_items).execute()
                        
                        st.success("Meeting created successfully!")
                        st.balloons()
                        navigate_to('meeting_detail', meeting=meeting_id)
                    
                    except Exception as e:
                        st.error(f"Error creating meeting: {str(e)}")
    
    # Return button
    st.divider()
    if st.button("↩ Return to Committees", use_container_width=True):
        navigate_to('home')

# ========== MEETING DETAIL VIEW ==========
elif st.session_state.committee_view == 'meeting_detail':
    
    try:
        # Fetch meeting details
        meeting_response = supabase.table("committee_meetings").select("*").eq(
            "id", st.session_state.selected_meeting
        ).execute()
        
        if not meeting_response.data:
            st.error("Meeting not found.")
            if st.button("↩ Return"):
                navigate_to('committee_home', st.session_state.selected_committee)
        else:
            meeting = meeting_response.data[0]
            committee = COMMITTEES[meeting['committee_name']]
            
            st.subheader(f"{committee['icon']} {meeting['meeting_title']}")
            st.caption(f"📅 {meeting['meeting_date']} | 🕐 {meeting['meeting_time']} | 📍 {meeting.get('location', 'TBA')}")
            
            # Edit Meeting Details Section
            with st.expander("✏️ Edit Meeting Details"):
                with st.form("edit_meeting_form"):
                    col1, col2 = st.columns(2)
                    
                    with col1:
                        edit_title = st.text_input(
                            "Meeting Title*", 
                            value=meeting['meeting_title'],
                            key="edit_title"
                        )
                        edit_date = st.date_input(
                            "Meeting Date*", 
                            value=datetime.strptime(meeting['meeting_date'], '%Y-%m-%d').date(),
                            key="edit_date"
                        )
                    
                    with col2:
                        edit_time = st.time_input(
                            "Meeting Time*", 
                            value=datetime.strptime(meeting['meeting_time'], '%H:%M:%S').time(),
                            key="edit_time"
                        )
                        edit_location = st.text_input(
                            "Location", 
                            value=meeting.get('location', ''),
                            key="edit_location"
                        )
                    
                    submit_edit = st.form_submit_button("💾 Update Meeting Details", use_container_width=True)
                    
                    if submit_edit:
                        try:
                            supabase.table("committee_meetings").update({
                                "meeting_title": edit_title,
                                "meeting_date": str(edit_date),
                                "meeting_time": str(edit_time),
                                "location": edit_location
                            }).eq("id", st.session_state.selected_meeting).execute()
                            
                            st.success("✅ Meeting details updated!")
                            st.rerun()
                        except Exception as e:
                            st.error(f"Error updating meeting: {str(e)}")
            
            st.divider()
            
            # Tab navigation
            tab1, tab2, tab3 = st.tabs(["📋 Agenda", "📝 Minutes", "⬇️ Download"])
            
            # ========== AGENDA TAB ==========
            with tab1:
                st.markdown("### Meeting Agenda")
                
                # Fetch agenda items
                agenda_response = supabase.table("committee_agenda_items").select("*").eq(
                    "meeting_id", st.session_state.selected_meeting
                ).order("item_number").execute()
                
                if agenda_response.data:
                    for item in agenda_response.data:
                        with st.expander(f"**{item['item_number']}. {item['item_title']}** ({item['duration_minutes']} min)"):
                            if item.get('item_description'):
                                st.markdown(item['item_description'])
                            else:
                                st.caption("No description provided")
                            
                            # Edit agenda item
                            st.divider()
                            st.markdown("**Edit This Item**")
                            
                            with st.form(f"edit_agenda_{item['id']}"):
                                col1, col2 = st.columns([3, 1])
                                with col1:
                                    new_title = st.text_input(
                                        "Item Title", 
                                        value=item['item_title'],
                                        key=f"title_{item['id']}"
                                    )
                                with col2:
                                    new_duration = st.number_input(
                                        "Duration (min)", 
                                        min_value=5, 
                                        value=item['duration_minutes'],
                                        key=f"dur_{item['id']}"
                                    )
                                
                                new_description = st.text_area(
                                    "Description", 
                                    value=item.get('item_description', ''),
                                    key=f"desc_{item['id']}"
                                )
                                
                                col_update, col_delete = st.columns(2)
                                
                                with col_update:
                                    if st.form_submit_button("💾 Update", use_container_width=True):
                                        try:
                                            supabase.table("committee_agenda_items").update({
                                                "item_title": new_title,
                                                "item_description": new_description,
                                                "duration_minutes": new_duration
                                            }).eq("id", item['id']).execute()
                                            st.success("Updated!")
                                            st.rerun()
                                        except Exception as e:
                                            st.error(f"Error: {str(e)}")
                                
                                with col_delete:
                                    if st.form_submit_button("🗑️ Delete", use_container_width=True):
                                        try:
                                            supabase.table("committee_agenda_items").delete().eq("id", item['id']).execute()
                                            st.success("Deleted!")
                                            st.rerun()
                                        except Exception as e:
                                            st.error(f"Error: {str(e)}")
                    
                    # Add new agenda item
                    st.divider()
                    st.markdown("### ➕ Add New Agenda Item")
                    with st.form("add_agenda_item"):
                        col1, col2 = st.columns([3, 1])
                        with col1:
                            new_item_title = st.text_input("Item Title*")
                        with col2:
                            new_item_duration = st.number_input("Duration (min)", min_value=5, value=10)
                        
                        new_item_desc = st.text_area("Description (optional)")
                        
                        if st.form_submit_button("➕ Add Item", use_container_width=True):
                            if new_item_title:
                                try:
                                    # Get next item number
                                    max_num = max([i['item_number'] for i in agenda_response.data]) if agenda_response.data else 0
                                    
                                    supabase.table("committee_agenda_items").insert({
                                        "meeting_id": st.session_state.selected_meeting,
                                        "item_number": max_num + 1,
                                        "item_title": new_item_title,
                                        "item_description": new_item_desc,
                                        "duration_minutes": new_item_duration
                                    }).execute()
                                    st.success("Item added!")
                                    st.rerun()
                                except Exception as e:
                                    st.error(f"Error: {str(e)}")
                            else:
                                st.warning("Please provide an item title.")
                
                else:
                    st.info("No agenda items yet. Add your first item below.")
                    
                    # Add first agenda item
                    st.markdown("### ➕ Add Agenda Item")
                    with st.form("add_first_agenda_item"):
                        col1, col2 = st.columns([3, 1])
                        with col1:
                            first_item_title = st.text_input("Item Title*")
                        with col2:
                            first_item_duration = st.number_input("Duration (min)", min_value=5, value=10)
                        
                        first_item_desc = st.text_area("Description (optional)")
                        
                        if st.form_submit_button("➕ Add First Item", use_container_width=True):
                            if first_item_title:
                                try:
                                    supabase.table("committee_agenda_items").insert({
                                        "meeting_id": st.session_state.selected_meeting,
                                        "item_number": 1,
                                        "item_title": first_item_title,
                                        "item_description": first_item_desc,
                                        "duration_minutes": first_item_duration
                                    }).execute()
                                    st.success("Item added!")
                                    st.rerun()
                                except Exception as e:
                                    st.error(f"Error: {str(e)}")
                            else:
                                st.warning("Please provide an item title.")
            
            # ========== MINUTES TAB ==========
            with tab2:
                st.markdown("### Meeting Minutes")
                
                # Check if minutes exist
                minutes_response = supabase.table("committee_minutes").select("*").eq(
                    "meeting_id", st.session_state.selected_meeting
                ).execute()
                
                existing_minutes = minutes_response.data[0] if minutes_response.data else None
                
                # Minutes input method selection
                input_method = st.radio(
                    "Choose minutes input method:",
                    ["AI-Assisted (Transcript)", "Manual Entry (DfE Format)"],
                    horizontal=True
                )
                
                if input_method == "AI-Assisted (Transcript)":
                    st.markdown("#### 🤖 AI-Assisted Minutes Generation")
                    st.caption("Paste your meeting transcript below and AI will convert it into structured minutes following DfE format.")
                    
                    transcript = st.text_area(
                        "Meeting Transcript",
                        height=300,
                        placeholder="Paste the meeting transcript here...",
                        value=existing_minutes.get('transcript', '') if existing_minutes else ''
                    )
                    
                    if st.button("🪄 Generate Minutes from Transcript", use_container_width=True):
                        if transcript:
                            with st.spinner("Generating minutes..."):
                                try:
                                    # AI prompt for minutes generation
                                    prompt = f"""You are a professional minute-taker for educational committees in South Australia. Convert this meeting transcript into formal meeting minutes following the Department for Education (DfE) format.

Meeting: {meeting['meeting_title']}
Date: {meeting['meeting_date']}
Committee: {committee['name']}

Format the minutes with these sections:
1. ATTENDANCE: List attendees (extract from transcript)
2. APOLOGIES: List apologies (if mentioned)
3. MINUTES OF PREVIOUS MEETING: Confirmation status
4. BUSINESS ARISING: Key points from previous meeting
5. AGENDA ITEMS: For each agenda item, provide:
   - Item title
   - Discussion summary (key points, decisions, actions)
   - Action items with responsible person
6. GENERAL BUSINESS: Any other business discussed
7. NEXT MEETING: Date and time if mentioned
8. MEETING CLOSED: Time if mentioned

Keep it professional, concise, and factual. Use past tense. Focus on decisions and actions.

TRANSCRIPT:
{transcript}

Generate the formatted minutes:"""
                                    
                                    generated_minutes = groq_ai_call(
                                        prompt,
                                        max_tokens=4000,
                                        temperature=0.3
                                    )
                                    
                                    # Save minutes
                                    if existing_minutes:
                                        supabase.table("committee_minutes").update({
                                            "transcript": transcript,
                                            "formatted_minutes": generated_minutes,
                                            "updated_at": datetime.now().isoformat()
                                        }).eq("meeting_id", st.session_state.selected_meeting).execute()
                                    else:
                                        supabase.table("committee_minutes").insert({
                                            "meeting_id": st.session_state.selected_meeting,
                                            "transcript": transcript,
                                            "formatted_minutes": generated_minutes,
                                            "created_at": datetime.now().isoformat()
                                        }).execute()
                                    
                                    st.success("✅ Minutes generated successfully!")
                                    st.rerun()
                                
                                except Exception as e:
                                    st.error(f"Error generating minutes: {str(e)}")
                        else:
                            st.warning("Please paste a transcript first.")
                    
                    # Display generated minutes
                    if existing_minutes and existing_minutes.get('formatted_minutes'):
                        st.divider()
                        st.markdown("#### Generated Minutes")
                        
                        edited_minutes = st.text_area(
                            "Review and edit minutes if needed:",
                            value=existing_minutes['formatted_minutes'],
                            height=400,
                            key="edit_ai_minutes"
                        )
                        
                        col1, col2, col3 = st.columns(3)
                        with col1:
                            if st.button("💾 Save Edited Minutes", use_container_width=True):
                                supabase.table("committee_minutes").update({
                                    "formatted_minutes": edited_minutes,
                                    "updated_at": datetime.now().isoformat()
                                }).eq("meeting_id", st.session_state.selected_meeting).execute()
                                st.success("Minutes saved!")
                                st.rerun()
                        
                        with col2:
                            if st.button("📢 Publish Minutes", use_container_width=True):
                                supabase.table("committee_meetings").update({
                                    "status": "published"
                                }).eq("id", st.session_state.selected_meeting).execute()
                                st.success("Minutes published!")
                                st.rerun()
                        
                        with col3:
                            try:
                                quick_docx = generate_minutes_docx(meeting, committee, edited_minutes)
                                st.download_button(
                                    "📄 Download Word",
                                    data=quick_docx,
                                    file_name=f"{meeting['meeting_title']}_{meeting['meeting_date']}_Minutes.docx",
                                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                    use_container_width=True,
                                    key="quick_download_ai"
                                )
                            except Exception as e:
                                st.caption(f"Preview unavailable: {str(e)[:40]}")
                
                else:  # Manual Entry
                    st.markdown("#### ✍️ Manual Minutes Entry (DfE Format)")
                    
                    with st.form("manual_minutes_form"):
                        st.markdown("**ATTENDANCE**")
                        attendance = st.text_area("Present", placeholder="List attendees...", height=80)
                        apologies = st.text_input("Apologies", placeholder="List apologies...")
                        
                        st.markdown("**MINUTES OF PREVIOUS MEETING**")
                        previous_confirmed = st.selectbox("Status", ["Confirmed", "Amended", "Not available"])
                        
                        st.markdown("**BUSINESS ARISING**")
                        business_arising = st.text_area("Business Arising from Previous Meeting", height=100, placeholder="Key points from previous meeting...")
                        
                        st.markdown("**AGENDA ITEMS**")
                        # Fetch agenda items
                        agenda_response = supabase.table("committee_agenda_items").select("*").eq(
                            "meeting_id", st.session_state.selected_meeting
                        ).order("item_number").execute()
                        
                        agenda_minutes = {}
                        if agenda_response.data:
                            for item in agenda_response.data:
                                st.markdown(f"**{item['item_number']}. {item['item_title']}**")
                                discussion = st.text_area(
                                    f"Discussion & Decisions",
                                    key=f"disc_{item['id']}",
                                    placeholder="Summarize discussion, decisions made, and outcomes...",
                                    height=100
                                )
                                actions = st.text_area(
                                    f"Action Items",
                                    key=f"action_{item['id']}",
                                    placeholder="Action item | Responsible person | Due date",
                                    height=60
                                )
                                agenda_minutes[item['item_number']] = {
                                    'title': item['item_title'],
                                    'discussion': discussion,
                                    'actions': actions
                                }
                        
                        st.markdown("**GENERAL BUSINESS**")
                        general_business = st.text_area("Other Business Discussed", height=100)
                        
                        st.markdown("**NEXT MEETING**")
                        next_meeting = st.text_input("Next Meeting Date/Time")
                        
                        st.markdown("**MEETING CLOSED**")
                        closed_time = st.text_input("Meeting Closed At")
                        
                        submit_manual = st.form_submit_button("💾 Save Minutes", use_container_width=True)
                        
                        if submit_manual:
                            # Format minutes in DfE structure
                            formatted_manual = f"""MEETING MINUTES

Committee: {committee['name']}
Meeting Title: {meeting['meeting_title']}
Date: {meeting['meeting_date']}
Time: {meeting['meeting_time']}
Location: {meeting.get('location', 'TBA')}

1. ATTENDANCE
Present: {attendance}
Apologies: {apologies}

2. MINUTES OF PREVIOUS MEETING
Status: {previous_confirmed}

3. BUSINESS ARISING
{business_arising}

4. AGENDA ITEMS
"""
                            for num, item_data in agenda_minutes.items():
                                formatted_manual += f"\n{num}. {item_data['title']}\n"
                                formatted_manual += f"Discussion: {item_data['discussion']}\n"
                                if item_data['actions']:
                                    formatted_manual += f"Actions: {item_data['actions']}\n"
                            
                            formatted_manual += f"""
5. GENERAL BUSINESS
{general_business}

6. NEXT MEETING
{next_meeting}

7. MEETING CLOSED
{closed_time}
"""
                            
                            try:
                                if existing_minutes:
                                    supabase.table("committee_minutes").update({
                                        "formatted_minutes": formatted_manual,
                                        "updated_at": datetime.now().isoformat()
                                    }).eq("meeting_id", st.session_state.selected_meeting).execute()
                                else:
                                    supabase.table("committee_minutes").insert({
                                        "meeting_id": st.session_state.selected_meeting,
                                        "formatted_minutes": formatted_manual,
                                        "created_at": datetime.now().isoformat()
                                    }).execute()
                                
                                st.success("✅ Minutes saved successfully!")
                                st.rerun()
                            
                            except Exception as e:
                                st.error(f"Error saving minutes: {str(e)}")
                    
                    # Show existing manual minutes
                    if existing_minutes and existing_minutes.get('formatted_minutes') and not existing_minutes.get('transcript'):
                        st.divider()
                        st.markdown("#### Current Minutes")
                        st.text_area("", value=existing_minutes['formatted_minutes'], height=400, disabled=True, key="view_manual")
                        
                        col1, col2 = st.columns(2)
                        with col1:
                            if st.button("📢 Publish Minutes", use_container_width=True):
                                supabase.table("committee_meetings").update({
                                    "status": "published"
                                }).eq("id", st.session_state.selected_meeting).execute()
                                st.success("Minutes published!")
                                st.rerun()
                        
                        with col2:
                            try:
                                quick_docx_manual = generate_minutes_docx(
                                    meeting, committee, existing_minutes['formatted_minutes']
                                )
                                st.download_button(
                                    "📄 Download Word",
                                    data=quick_docx_manual,
                                    file_name=f"{meeting['meeting_title']}_{meeting['meeting_date']}_Minutes.docx",
                                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                    use_container_width=True,
                                    key="quick_download_manual"
                                )
                            except Exception as e:
                                st.caption(f"Preview unavailable: {str(e)[:40]}")
            
            # ========== DOWNLOAD TAB ==========
            with tab3:
                st.markdown("### Download Meeting Documents")
                
                # FIX: Re-fetch minutes data in this tab instead of relying on tab2's variable
                download_minutes_response = supabase.table("committee_minutes").select("*").eq(
                    "meeting_id", st.session_state.selected_meeting
                ).execute()
                
                if download_minutes_response.data:
                    minutes_data = download_minutes_response.data[0]
                    
                    # Download minutes as a professionally formatted Word document
                    if minutes_data.get('formatted_minutes'):
                        try:
                            docx_buffer = generate_minutes_docx(
                                meeting, committee, minutes_data['formatted_minutes']
                            )
                            st.download_button(
                                label="📄 Download Minutes (Word)",
                                data=docx_buffer,
                                file_name=f"{meeting['meeting_title']}_{meeting['meeting_date']}_Minutes.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                use_container_width=True
                            )
                        except Exception as e:
                            st.error(f"Error generating Word document: {str(e)}")
                    
                    # Download transcript if exists (kept as plain text - it's a raw transcript)
                    if minutes_data.get('transcript'):
                        st.download_button(
                            label="📝 Download Transcript (TXT)",
                            data=minutes_data['transcript'],
                            file_name=f"{meeting['meeting_title']}_{meeting['meeting_date']}_Transcript.txt",
                            mime="text/plain",
                            use_container_width=True
                        )
                else:
                    st.info("No minutes available for download yet.")
            
            # Return button
            st.divider()
            if st.button("↩ Return to Committee", use_container_width=True):
                navigate_to('committee_home', meeting['committee_name'])
    
    except Exception as e:
        st.error(f"Error loading meeting: {str(e)}")
        if st.button("↩ Return"):
            navigate_to('home')

# ========== DOCUMENTS VIEW ==========
elif st.session_state.committee_view == 'documents':
    
    tab1, tab2 = st.tabs(["📄 View Documents", "⬆️ Upload Document"])
    
    with tab1:
        st.subheader("Committee Documents")
        
        col1, col2 = st.columns([3, 1])
        with col1:
            search_term = st.text_input("🔍 Search documents", placeholder="Search by name or description...")
        with col2:
            filter_category = st.selectbox("Filter by category", ["All", "Policy", "Procedure", "Meeting Minutes", "Report", "Other"])
        
        try:
            query = supabase.table("committees_documents").select("*").order("upload_date", desc=True)
            
            if filter_category != "All":
                query = query.eq("category", filter_category)
            
            response = query.execute()
            
            if response.data:
                filtered_docs = response.data
                if search_term:
                    filtered_docs = [
                        doc for doc in response.data 
                        if search_term.lower() in doc['document_name'].lower() 
                        or (doc.get('description') and search_term.lower() in doc['description'].lower())
                    ]
                
                if filtered_docs:
                    for doc in filtered_docs:
                        with st.container():
                            col1, col2, col3, col4, col5 = st.columns([3, 1.5, 1.5, 1.5, 1.5])
                            
                            with col1:
                                pin = "📌 " if doc.get('is_current') else ""
                                st.markdown(f"**{pin}{doc['document_name']}**")
                                if doc.get('description'):
                                    st.caption(doc['description'])
                            
                            with col2:
                                if doc.get('category'):
                                    st.caption(f"🏷️ {doc['category']}")
                            
                            with col3:
                                st.caption(f"📅 {doc['upload_date'][:10]}")
                            
                            with col4:
                                st.caption(f"👤 {doc['uploaded_by']}")
                            
                            with col5:
                                try:
                                    signed_url = supabase.storage.from_("committees").create_signed_url(
                                        doc['file_path'], 
                                        3600
                                    )
                                    
                                    if signed_url:
                                        st.link_button(
                                            "👁️",
                                            signed_url['signedURL'],
                                            use_container_width=True
                                        )
                                    
                                    file_data = supabase.storage.from_("committees").download(doc['file_path'])
                                    st.download_button(
                                        label="⬇️",
                                        data=file_data,
                                        file_name=doc['document_name'],
                                        mime=doc.get('mime_type', 'application/octet-stream'),
                                        key=f"download_{doc['id']}",
                                        use_container_width=True
                                    )
                                except Exception as e:
                                    st.error(f"Error: {str(e)[:30]}")
                            
                            st.divider()
                else:
                    st.warning(f"No documents found matching '{search_term}'")
            else:
                st.info("📭 No documents uploaded yet.")
        
        except Exception as e:
            st.error(f"Error loading documents: {str(e)}")
    
    with tab2:
        st.subheader("Upload Committee Document")
        
        with st.form("upload_form"):
            uploaded_file = st.file_uploader(
                "Choose a file",
                type=['pdf', 'docx', 'doc', 'xlsx', 'xls', 'pptx', 'txt']
            )
            
            col1, col2 = st.columns(2)
            
            with col1:
                doc_name = st.text_input("Document Name*", placeholder="e.g., Committees Policy 2026")
                category = st.selectbox("Document Category*", ["Policy", "Procedure", "Meeting Minutes", "Report", "Other"])
            
            with col2:
                uploaded_by = st.text_input("Your Name*", placeholder="e.g., Candice")
                is_current_policy = st.checkbox("📌 Mark as current/active policy")
            
            description = st.text_area("Description (optional)", placeholder="Brief description...")
            
            submit = st.form_submit_button("📤 Upload Document", use_container_width=True)
            
            if submit:
                if not uploaded_file or not doc_name or not uploaded_by:
                    st.error("Please complete all required fields.")
                else:
                    try:
                        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                        file_path = f"{timestamp}_{uploaded_file.name}"
                        file_bytes = uploaded_file.read()
                        
                        supabase.storage.from_("committees").upload(file_path, file_bytes, {"content-type": uploaded_file.type})
                        
                        supabase.table("committees_documents").insert({
                            "document_name": doc_name,
                            "description": description,
                            "file_path": file_path,
                            "uploaded_by": uploaded_by,
                            "upload_date": datetime.now().isoformat(),
                            "category": category,
                            "mime_type": uploaded_file.type,
                            "file_size": len(file_bytes),
                            "is_current": is_current_policy
                        }).execute()
                        
                        st.success(f"✅ Document '{doc_name}' uploaded successfully!")
                        st.balloons()
                    
                    except Exception as e:
                        st.error(f"Error uploading document: {str(e)}")
    
    st.divider()
    if st.button("↩ Return to Committees Home", use_container_width=True):
        navigate_to('home')

# ========== FOOTER ==========
st.divider()
st.caption("💡 Tip: Use AI-assisted minutes generation to save time, or enter minutes manually following DfE format.")
