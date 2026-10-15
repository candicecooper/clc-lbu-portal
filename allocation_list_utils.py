"""
allocation_list_utils.py
Writes newly allocated students into the CLC Allocation List workbook
(e.g. CLC_Allocation_List_Week_2.xlsx) without breaking its layout.

Layout it expects (sheet "Student Information"):
  Row 2            column headings
  "JUNIOR PRIMARY" section header row  -> student rows -> spare blank rows
  "PRIMARY"        section header row  -> student rows
  "SENIORS"        section header row  -> student rows
  "TOTAL STUDENT ALLOCATED" row (count in column E)

Sections are found by scanning column A, so the tool keeps working if rows
are added or removed by hand.
"""

from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass
from datetime import date, datetime

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter, range_boundaries

SHEET_NAME = "Student Information"

# Section header text in column A  ->  program code
SECTION_HEADERS = {
    "JUNIOR PRIMARY": "JP",
    "PRIMARY": "PY",
    "SENIORS": "SY",
}
TOTAL_LABEL = "TOTAL STUDENT ALLOCATED"

# Program aliases accepted from the allocation data
PROGRAM_ALIASES = {
    "JP": "JP", "JUNIOR PRIMARY": "JP", "JUNIOR": "JP",
    "PY": "PY", "PRIMARY": "PY", "PRIMARY YEARS": "PY",
    "SY": "SY", "SENIOR": "SY", "SENIORS": "SY", "SENIOR YEARS": "SY",
}

# Column letters in the allocation list
COL = {
    "num": "A", "name": "B", "edid": "C", "dob": "D", "school": "E",
    "year": "F", "male": "G", "female": "H",
    "start_term": "I", "start_week": "J", "start_date": "K",
    "review": "L",
    "exit_term": "M", "exit_week": "N", "exit_date": "O",
    "ext_term_week": "P", "ext_end_date": "Q",
    "prev": "R",
    "ealD": "S", "atsi": "T", "dcp": "U", "sic": "V", "iesp": "W",
    "swd": "X", "speech_language": "Y", "asd": "Z", "adhd": "AA",
    "intellectual": "AB", "language_difficulty": "AC",
    "sensory_processing": "AD", "sc": "AE", "medication": "AF",
    "focus_teacher": "AG", "bsc": "AH",
}
FLAG_FIELDS = [
    "ealD", "atsi", "dcp", "sic", "swd", "speech_language", "asd", "adhd",
    "intellectual", "language_difficulty", "sensory_processing", "sc",
    "medication",
]
FLAG_LABELS = {
    "ealD": "EALD", "atsi": "ATSI", "dcp": "DCP Involve", "sic": "SIC",
    "swd": "SWD", "speech_language": "Speech and Language", "asd": "ASD",
    "adhd": "ADHD", "intellectual": "Intellectual",
    "language_difficulty": "LNG Diff", "sensory_processing": "Sensory Proces",
    "sc": "SC", "medication": "Medication",
}
LAST_COL = "AH"

# ---------------------------------------------------------------------------
# SA Department for Education term dates.  UPDATE EACH YEAR.
# (start = first day of Week 1, end = last day of term)
# ---------------------------------------------------------------------------
SA_TERMS = {
    2026: [
        (date(2026, 1, 27), date(2026, 4, 10)),
        (date(2026, 4, 27), date(2026, 7, 3)),
        (date(2026, 7, 20), date(2026, 9, 25)),
        (date(2026, 10, 12), date(2026, 12, 11)),
    ],
    2027: [
        (date(2027, 1, 27), date(2027, 4, 9)),
        (date(2027, 4, 26), date(2027, 7, 2)),
        (date(2027, 7, 19), date(2027, 9, 24)),
        (date(2027, 10, 11), date(2027, 12, 10)),
    ],
}


def term_and_week(d: date | None) -> tuple[str, str]:
    """Return ("T4", "Wk 1") for a date, or ("", "") if outside a known term."""
    if not d:
        return "", ""
    for terms in SA_TERMS.values():
        for i, (start, end) in enumerate(terms, start=1):
            if start <= d <= end:
                # Week 1 starts on the Monday of the term's first week
                monday = start.toordinal() - start.weekday()
                week = (d.toordinal() - monday) // 7 + 1
                return f"T{i}", f"Wk {week}"
    return "", ""


def current_term_week(today: date | None = None) -> tuple[str, str]:
    return term_and_week(today or date.today())


# ---------------------------------------------------------------------------
# Value helpers
# ---------------------------------------------------------------------------
def to_date(value) -> date | None:
    """Accept date/datetime/ISO string/'d/m/yyyy' and return a date."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s[:19] if "T" in s else s, fmt).date()
        except ValueError:
            continue
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})", s)  # '20/11' -> this year
    if m:
        try:
            return date(date.today().year, int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def is_yes(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"y", "yes", "true", "1", "x", "✓"}


def norm_edid(value) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper()


def norm_program(value) -> str:
    """Accepts JP / 'Junior Primary' / 'Junior Primary (JP)' etc."""
    v = str(value or "").strip().upper()
    m = re.search(r"\((JP|PY|SY)\)", v)
    if m:
        return m.group(1)
    return PROGRAM_ALIASES.get(v, "")


def year_cell(value):
    """Keep 'Rec' as text, year numbers as numbers (matches the existing list)."""
    s = str(value or "").strip()
    if not s:
        return None
    if s.lower() in {"r", "rec", "reception", "0"}:
        return "Rec"
    s = re.sub(r"(?i)^(yr|year)\s*", "", s)
    return int(s) if s.isdigit() else s


# ---------------------------------------------------------------------------
# Student record
# ---------------------------------------------------------------------------
@dataclass
class NewStudent:
    name: str
    edid: str
    program: str                 # JP / PY / SY
    dob: date | None = None
    school: str = ""
    year: str = ""
    gender: str = ""             # "M" / "F" / ""
    start_date: date | None = None
    exit_date: date | None = None
    review: str = ""             # free text, e.g. "22 Oct @ 1pm"
    prev: str | int | None = None
    iesp: str | int | None = None
    flags: dict | None = None    # {"asd": True, ...}
    focus_teacher: str = ""
    bsc: str = ""
    ext_end_date: date | None = None   # approved extension: new end date

    def problems(self) -> list[str]:
        out = []
        if not self.name.strip():
            out.append("no student name")
        if self.program not in ("JP", "PY", "SY"):
            out.append("program not JP / PY / SY")
        return out


# ---------------------------------------------------------------------------
# Workbook structure
# ---------------------------------------------------------------------------
@dataclass
class Section:
    program: str
    header_row: int
    first_row: int
    last_row: int  # last row belonging to this section (before next header)


def _cell_text(ws, row, col="A") -> str:
    v = ws[f"{col}{row}"].value
    return str(v).strip().upper() if v is not None else ""


def find_layout(ws) -> tuple[list[Section], int]:
    """Return the three sections and the TOTAL row number."""
    headers = []
    total_row = None
    for r in range(1, ws.max_row + 1):
        t = _cell_text(ws, r)
        if t in SECTION_HEADERS:
            headers.append((r, SECTION_HEADERS[t]))
        elif t.startswith("TOTAL"):
            total_row = r
            break
    if total_row is None or len(headers) != 3:
        raise ValueError(
            "Couldn't find the JUNIOR PRIMARY / PRIMARY / SENIORS headings and the "
            "TOTAL row in column A. Is this the CLC Allocation List?"
        )
    sections = []
    for i, (r, prog) in enumerate(headers):
        nxt = headers[i + 1][0] if i + 1 < len(headers) else total_row
        sections.append(Section(prog, r, r + 1, nxt - 1))
    return sections, total_row


def norm_name(value) -> str:
    """'PAYTON  HARVEY ' -> 'payton harvey' (the list mixes upper and title case)."""
    return re.sub(r"[^a-z' -]", "", re.sub(r"\s+", " ", str(value or "").strip().lower()))


def existing_names(ws) -> dict[str, str]:
    """normalised student name -> program for every student already on the list."""
    sections, _ = find_layout(ws)
    found = {}
    for s in sections:
        for r in range(s.first_row, s.last_row + 1):
            n = norm_name(ws[f"{COL['name']}{r}"].value)
            if n:
                found[n] = s.program
    return found


def existing_edids(ws) -> dict[str, str]:
    """EDID -> program for every student already on the list."""
    sections, _ = find_layout(ws)
    found = {}
    for s in sections:
        for r in range(s.first_row, s.last_row + 1):
            e = norm_edid(ws[f"{COL['edid']}{r}"].value)
            if e:
                found[e] = s.program
    return found


def _row_is_empty(ws, r) -> bool:
    return not str(ws[f"{COL['name']}{r}"].value or "").strip() and not norm_edid(
        ws[f"{COL['edid']}{r}"].value
    )


# ---------------------------------------------------------------------------
# Row insertion that keeps merges, heights and conditional formatting intact
# (openpyxl's insert_rows doesn't move any of these on its own)
# ---------------------------------------------------------------------------
def _shift_range(ref: str, at: int, n: int, extend_end_at: int | None) -> str:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    if min_row >= at:
        min_row += n
        max_row += n
    elif max_row >= at:
        max_row += n
    elif extend_end_at is not None and max_row == extend_end_at:
        max_row += n  # range ended on the section's last row: grow it
    a = f"{get_column_letter(min_col)}{min_row}"
    b = f"{get_column_letter(max_col)}{max_row}"
    return a if a == b else f"{a}:{b}"


_REF = re.compile(r"(?<![A-Za-z$])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(])")


def _shift_formula(f: str, at: int, n: int) -> str:
    """Move row numbers >= at down by n in a conditional-format formula."""
    if f.startswith('"'):
        return f

    def rep(m):
        row = int(m.group(4))
        if row >= at:
            row += n
        return f"{m.group(1)}{m.group(2)}{m.group(3)}{row}"

    return _REF.sub(rep, f)


def insert_rows_keep_format(ws, at: int, n: int, template_row: int):
    """Insert n rows before `at`, styled like template_row."""
    # merges
    merges = [str(m) for m in ws.merged_cells.ranges]
    for m in merges:
        ws.unmerge_cells(m)
    # row heights
    heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height}
    # conditional formatting
    old_cf = ws.conditional_formatting
    cf_rules = []
    for cf in old_cf:
        refs = str(cf.sqref).split()
        cf_rules.append((refs, cf.rules))

    ws.insert_rows(at, n)
    tmpl = template_row + n if template_row >= at else template_row

    for m in merges:
        ws.merge_cells(_shift_range(m, at, n, None))

    for r in sorted(heights, reverse=True):
        if r >= at:
            ws.row_dimensions[r + n].height = heights[r]
    tmpl_h = heights.get(template_row)
    for i in range(n):
        ws.row_dimensions[at + i].height = tmpl_h

    new_cf = ConditionalFormattingList()
    for refs, rules in cf_rules:
        shifted = " ".join(_shift_range(x, at, n, at - 1) for x in refs)
        # rule formulas are relative to the first range's top-left cell
        moved = range_boundaries(refs[0])[1] >= at
        for rule in rules:
            if moved and rule.formula:
                rule.formula = [_shift_formula(f, at, n) for f in rule.formula]
            new_cf.add(shifted, rule)
    ws.conditional_formatting = new_cf

    max_col = openpyxl.utils.column_index_from_string(LAST_COL)
    for i in range(n):
        for c in range(1, max_col + 1):
            src = ws.cell(tmpl, c)
            dst = ws.cell(at + i, c)
            if src.has_style:
                dst._style = copy.copy(src._style)


# ---------------------------------------------------------------------------
# Writing a student
# ---------------------------------------------------------------------------
def _write_student(ws, r: int, s: NewStudent):
    def put(key, value):
        ws[f"{COL[key]}{r}"] = value if value not in ("", None) else None

    put("name", s.name.strip())
    put("edid", norm_edid(s.edid))
    put("dob", s.dob)
    put("school", s.school.strip())
    put("year", year_cell(s.year))
    g = s.gender.strip().upper()[:1]
    put("male", "y" if g == "M" else None)
    put("female", "y" if g == "F" else None)
    st, sw = term_and_week(s.start_date)
    put("start_term", st)
    put("start_week", sw)
    put("start_date", s.start_date)
    put("review", s.review.strip())
    et, ew = term_and_week(s.exit_date)
    put("exit_term", et)
    put("exit_week", ew)
    put("exit_date", s.exit_date)
    if s.ext_end_date:
        xt, xw = term_and_week(s.ext_end_date)
        put("ext_term_week", f"{xt} {xw.lower()}".strip() if xt else None)
        put("ext_end_date", s.ext_end_date)
        cell = ws[f"{COL['ext_end_date']}{r}"]
        if cell.number_format == "General":
            cell.number_format = "d-mmm"
    else:
        put("ext_term_week", None)
        put("ext_end_date", None)
    put("prev", s.prev)
    put("iesp", s.iesp)
    for f in FLAG_FIELDS:
        put(f, "y" if (s.flags or {}).get(f) else None)
    put("focus_teacher", s.focus_teacher.strip())
    put("bsc", s.bsc.strip())
    for key in ("dob", "start_date", "exit_date"):
        cell = ws[f"{COL[key]}{r}"]
        if cell.number_format == "General":  # keep each column's own date style
            cell.number_format = "d-mmm" if key == "exit_date" else "d/mm/yyyy"


def set_header_week(ws, term: str, week: int) -> bool:
    """Update the printed page header, e.g. 'TERM 4  WEEK 2' -> 'TERM 4  WEEK 3'."""
    changed = False
    for hdr in (ws.oddHeader, ws.evenHeader, ws.firstHeader):
        for part in (hdr.left, hdr.center, hdr.right):
            if part.text and re.search(r"(?i)TERM\s*\d+\s*WEEK\s*\d+", part.text):
                part.text = re.sub(
                    r"(?i)TERM\s*\d+(\s*)WEEK\s*\d+",
                    lambda m: f"TERM {term.lstrip('Tt')}{m.group(1)}WEEK {week}",
                    part.text,
                )
                changed = True
    return changed


def _renumber_and_total(ws):
    sections, total_row = find_layout(ws)
    total = 0
    for s in sections:
        n = 0
        for r in range(s.first_row, s.last_row + 1):
            if _row_is_empty(ws, r):
                ws[f"{COL['num']}{r}"] = None
            else:
                n += 1
                ws[f"{COL['num']}{r}"] = n
        total += n
    ws[f"E{total_row}"] = total
    return total


@dataclass
class AddResult:
    added: list[tuple[str, str]]      # (name, program)
    skipped: list[tuple[str, str]]    # (name, reason)
    total: int
    workbook_bytes: bytes


def add_students(
    xlsx_bytes: bytes,
    students: list[NewStudent],
    header_term: str | None = None,
    header_week: int | None = None,
) -> AddResult:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]

    on_list = existing_edids(ws)
    names_on_list = existing_names(ws)
    added, skipped = [], []
    seen, seen_names = set(), set()

    for s in students:
        label = s.name.strip() or "(no name)"
        issues = s.problems()
        if issues:
            skipped.append((label, ", ".join(issues)))
            continue
        e = norm_edid(s.edid)
        nm = norm_name(s.name)
        if (e and (e in on_list or e in seen)) or nm in names_on_list or nm in seen_names:
            skipped.append((label, "already on the list"))
            continue

        sections, _ = find_layout(ws)
        sec = next(x for x in sections if x.program == s.program)
        target = next(
            (r for r in range(sec.first_row, sec.last_row + 1) if _row_is_empty(ws, r)),
            None,
        )
        if target is None:
            # section full: add a row at the bottom of it, styled like the last row
            target = sec.last_row + 1
            insert_rows_keep_format(ws, target, 1, template_row=sec.last_row)
        _write_student(ws, target, s)
        h = ws.row_dimensions[target].height
        if not h or h < 28:
            ws.row_dimensions[target].height = 30  # spare rows are short
        if e:
            seen.add(e)
        seen_names.add(nm)
        added.append((label, s.program))

    total = _renumber_and_total(ws)
    if header_term and header_week:
        set_header_week(ws, header_term, header_week)
    out = io.BytesIO()
    wb.save(out)
    return AddResult(added, skipped, total, out.getvalue())


def next_week_filename(current_name: str) -> str:
    """CLC_Allocation_List_Week_2.xlsx -> CLC_Allocation_List_Week_3.xlsx"""
    m = re.search(r"(?i)(week[_\s-]*)(\d+)", current_name)
    if m:
        return current_name[: m.start(2)] + str(int(m.group(2)) + 1) + current_name[m.end(2):]
    stem = current_name.rsplit(".", 1)[0]
    return f"{stem}_updated.xlsx"


# ---------------------------------------------------------------------------
# Whole-list build (live list -> Head Office spreadsheet)
# ---------------------------------------------------------------------------
def clear_students(ws):
    """Blank every student row (values only) — keeps headings, styles, rules."""
    sections, _ = find_layout(ws)
    last = openpyxl.utils.column_index_from_string(LAST_COL)
    for sec in sections:
        for r in range(sec.first_row, sec.last_row + 1):
            for c in range(1, last + 1):
                cell = ws.cell(r, c)
                if not isinstance(cell, openpyxl.cell.cell.MergedCell):
                    cell.value = None


def build_list(
    template_bytes: bytes,
    students: list[NewStudent],
    term: str,
    week: int,
) -> AddResult:
    """Fresh list from the template: clear it, then add every active student,
    each program block ordered by start date."""
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    clear_students(ws)
    buf = io.BytesIO()
    wb.save(buf)
    ordered = sorted(
        students,
        key=lambda s: ({"JP": 0, "PY": 1, "SY": 2}.get(s.program, 9),
                       s.start_date or date.max, s.name.lower()),
    )
    return add_students(buf.getvalue(), ordered, header_term=term, header_week=week)


def _tw_key(term, week) -> str:
    t = re.sub(r"\D", "", str(term or ""))
    w = re.sub(r"\D", "", str(week or ""))
    return f"{t}-{w}" if t and w else ""


def fix_date(d: date | None, term, week) -> tuple[date | None, bool]:
    """If a date disagrees with the Term/Week typed beside it but its day/month
    swapped agrees, return the swapped date. Returns (date, was_fixed)."""
    want = _tw_key(term, week)
    if not d or not want:
        return d, False
    if _tw_key(*term_and_week(d)) == want:
        return d, False
    try:
        swapped = date(d.year, d.day, d.month)
    except ValueError:
        return d, False
    if _tw_key(*term_and_week(swapped)) == want:
        return swapped, True
    return d, False


def read_list(xlsx_bytes: bytes) -> list[dict]:
    """Read every student row of an existing allocation list (for the one-off import)."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    sections, _ = find_layout(ws)
    rows = []
    for sec in sections:
        for r in range(sec.first_row, sec.last_row + 1):
            if _row_is_empty(ws, r):
                continue
            v = {k: ws[f"{c}{r}"].value for k, c in COL.items()}
            g = "M" if is_yes(v["male"]) else "F" if is_yes(v["female"]) else ""
            start, f1 = fix_date(to_date(v["start_date"]), v["start_term"], v["start_week"])
            exit_, f2 = fix_date(to_date(v["exit_date"]), v["exit_term"], v["exit_week"])
            ext_tw = re.match(r"(?i)\s*T\s*(\d)\s*w(?:ee)?k\s*(\d+)", str(v["ext_term_week"] or ""))
            ext, f3 = fix_date(to_date(v["ext_end_date"]), *(ext_tw.groups() if ext_tw else ("", "")))
            if not exit_ and v["exit_term"] and v["exit_week"]:
                exit_ = None  # unreadable (e.g. a typo) — flagged in the import report
            rows.append({
                "program": sec.program,
                "student_name": str(v["name"] or "").strip(),
                "edid": norm_edid(v["edid"]),
                "dob": to_date(v["dob"]),
                "school": str(v["school"] or "").strip(),
                "year_level": str(v["year"] or "").strip(),
                "gender": g,
                "start_date": start,
                "exit_date": exit_,
                "ext_end_date": ext,
                "_fixed": [n for n, f in (("start", f1), ("exit", f2), ("extension", f3)) if f],
                "_exit_tw": f"{v['exit_term'] or ''} {v['exit_week'] or ''}".strip(),
                "review": str(v["review"] or "").strip(),
                "prev": str(v["prev"]).strip() if v["prev"] not in (None, "") else "",
                "iesp": str(v["iesp"]).strip() if v["iesp"] not in (None, "") else "",
                "flags": [f for f in FLAG_FIELDS if is_yes(v[f])],
                "focus_teacher": str(v["focus_teacher"] or "").strip(),
                "bsc": str(v["bsc"] or "").strip(),
                "_raw_exit": v["exit_date"],
                "_raw_ext": v["ext_end_date"],
                "_raw_start": v["start_date"],
            })
    return rows


def make_template(xlsx_bytes: bytes) -> bytes:
    """Blank copy of a list (no student data) to keep in the repo as the template."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    clear_students(ws)
    _renumber_and_total(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Built-in blank template (the CLC Allocation List layout, no student data).
# Used when data/allocation_list/allocation_list_template.xlsx isn't in the repo.
# ---------------------------------------------------------------------------
_TEMPLATE_B64 = (
    "UEsDBBQAAAAIAOxgT11Gx01IlQAAAM0AAAAQAAAAZG9jUHJvcHMvYXBwLnhtbE3PTQvCMAwG4L9SdreZih6kDkQ9ip68zy51hbYp"
    "bYT67+0EP255ecgboi6JIia2mEXxLuRtMzLHDUDWI/o+y8qhiqHke64x3YGMsRoPpB8eA8OibdeAhTEMOMzit7Dp1C5GZ3XPlkJ3"
    "sjpRJsPiWDQ6sScfq9wcChDneiU+ixNLOZcrBf+LU8sVU57mym/8ZAW/B7oXUEsDBBQAAAAIAOxgT13VS41oLgEAAIYCAAARAAAA"
    "ZG9jUHJvcHMvY29yZS54bWzFktFLwzAQxv+V0icFuyStOghdYZv4tIHgQPEtJLctmKQhOen235vWrWPou4/57rtfvuOulp7LNsBL"
    "aD0E1BCzgzUuculn+R7Rc0Ki3IMVcZIcLhW3bbAC0zPsiBfyU+yAlJQ+EgsolEBBemDhR2Le1EpyGUBgG054JUe8/wpmgClJwIAF"
    "h5GwCSN5M1dWOx0x9J01uVB6IkKw8UcANWIH9U/2UCH5yXmIenR1XTfpqsGXBmHkfb16HWYutIsonITUFTXHo4dZfv75rVo+bZ7z"
    "pqRsWjBaVGxDK15RXt5/9Fmv8l0C21bprf7nxOVjn5g9bCjj1fQ68TlgU6fbMCLi+iQsjs1CGEjr0O4umzsHiJDdLNtOOKWNEdkK"
    "RHDa7bJlWmKA25r8Rgza9ck131BLAwQUAAAACADsYE9di4JuWOwFAACOGgAAEwAAAHhsL3RoZW1lL3RoZW1lMS54bWztWVuLGzcU"
    "fi/0P4h5d+Y+tpd4gz22kza7SchuUvIoj2WPYs3IjOTdNSFQksdCoTQtfSn0rQ+lbSCBvqS/ZtuUNoX8hWo0vmhsTS7NBlIaG+zR"
    "0XeOPp1zdKSZOX/hJCHgCGUM07Rl2OcsA6A0okOcjlvGjcN+rWEAxmE6hISmqGXMETMu7H74wXm4w2OUICD0U7YDW0bM+XTHNFkk"
    "xJCdo1OUir4RzRLIRTMbm8MMHgu7CTEdywrMBOLUAClMhNmroxGOEDjMTRq7S+M9In5SznJBRLKDSI6oakjscGLnf2zOQpKBI0ha"
    "hhhnSI8P0Qk3AIGMi46WYcmPYe6eN1dKhFfoKnp9+VnoLRSGE0fqZePBStHzfC9or+w7hf1tXK/eC3rByp4EwCgSM7W3sH6n2en6"
    "C6wCKi41trv1rmuX8Ip9dwvf9vNvCe+u8d4Wvt8P1z5UQMWlr/FJ3Qm9Et5f44MtfN1qd716CS9BMcHpZAtt+YEbLme7gowouaSF"
    "N32vX3cW8DXKVLKr0E95Va4l8DbN+gIggws5TgGfT9EIRgIXQoIHGQZ7eByLxJvClDIhthyrb7niN/968kp6BO4gqGgXoohtiXI+"
    "gEUZnvKW8bGwaiiQ509+fP7kEXj+5OHpvcen9345vX//9N7PGsVLMB2ris++/+Lvbz8Ffz367tmDr/R4puJ//+mz3379Ug/kKvDp"
    "1w//ePzw6Tef//nDAw28ncGBCj/ECWLgCjoG12ki5qYZAA2y19M4jCEuacBYIDXAHo9LwCtzSHS4Dio772YmioQOeHF2u8T1IM5m"
    "HGuAl+OkBNynlHRopp3O5XwsdTqzdKwfPJupuOsQHunGDjdC25tNRbZjnckwRiWa14iINhyjFHGQ99EJQhq1WxiX/LqPo4wyOuLg"
    "FgYdiLUuOcQDrle6hBMRl7mOoAh1yTf7N0GHEp35LjoqI8WCgERnEpGSGy/CGYeJljFMiIrcgzzWkTyYZ1HJ4YyLSI8RoaA3RIzp"
    "dK5m8xLdy6K46MO+T+ZJGZlxPNEh9yClKrJLJ2EMk6mWM05jFfsRm4gUheAa5VoStLxC8raIA0wrw30TI/56y/qGqKv6BMl7Zplu"
    "SSBaXo9zMoIoXewBpWqe4PSlpX2jqPvvi7q+qLczrF1am6W8CvcfLOBdOEuvIbFm3tfv9/X7/1i/q9by2VftdaE21dO6NJNUHt1H"
    "mJADPidoj8kSz8T0hn0hlA2ptLpTmMbicjFcCTfOoLwGGeWfYB4fxHAqhrHlCGO2MD1mYEqZ2CSMSttyk5kl+3RYSG17eXMqFCBf"
    "y8Ums5SLLYkX0qC+vgtbmZetMVMJ+NLoq5NQBiuTcDUk6u6rkbCts2LR1LBo2C9iYSpREesPwPy5hu8VjES+QYKGeZwK/WV0zzzS"
    "Vc4sT9vRTK/pnVmkSySUdCuTUNIwhkO0KT7jWDeb+lA7Whr1xtuItbldG0haboFjseZcX5iJ4LRljMTxUFwmU2GP5XUTknHaMiK+"
    "cPS/qSzTjPEuZHEBk13F/BPMUQYITkSuq2Eg6Zqb7dStd5dc03r3PGduBhmNRijiFZJ1U/QVRrS9bwjOG3QmSB/Ew2MwILPsOhSO"
    "8ut27sAhZnzlzSHOlORee3GjXC2WYumh2XqJQjKN4WJHUYt5AZfXKzrKPCTTzVmZOhcOxv2z2HVfrrRRNCs2kHplFXt7m7zCytWz"
    "8rW1rtmwXrxLvPmGoFBr6Km5empVe8cZHgiU4YIKvzmV0XzD3WAza03lXClbW28n6OC2yPyuOK7OCGfFY4ATcY8QLp8rF5VASpfV"
    "5YSDWYZbxh3Lb3uh44c1q+H3ap7rWbWG33Zrbd937Z5vW92Oc1c4hceJ7Rdj98X9DJkvXr5I+dYLmGR5zD4X0cSk8hxsSmX5AsZ2"
    "ql/AACw8cydw+k232QlqTbfdr3ndTqPWDINOrRuE9W6/G/qNZv+uAY4k2Gu7oRf0GrXADsOaF1g5/UazVvccp+3V242e17678LWY"
    "+fJ/6V7Ja/cfUEsDBBQAAAAIAOxgT13V0pYU+gAAACkCAAATAAAAZG9jUHJvcHMvY3VzdG9tLnhtbLWSUWuDMBDHv4rk3eVM1Tai"
    "Qk0s9GEwWNm7aLRCY0KSupax777Itu5h7KWwl4M/d9z9fnD5k1FaGDcKG1zkabIFOjqnM4xtexSysQ++PflOr4xsnI9mwKrvx1Zw"
    "1Z6lmBwmACluz9YpGerbOlTmX+EaTI0UBWJqcn7+cNVi36Ggl27sCvTGE8Z5AklIasrCCKIqpCu6DmEDQCrCdnRbv6NAL8PEb51d"
    "dtKv1plP4Gx29zJ3ql307cuC5IHhAv48AE8Jo0lc7TipgcbRdinEIyaEpetNnOMbQ5njb8tfvo+iG5tnYWZ/eC+bQRyawd7nvfpX"
    "77988M9zlB9QSwMEFAAAAAgA7GBPXc+O9XzhEQAAIn0AABgAAAB4bC93b3Jrc2hlZXRzL3NoZWV0MS54bWzFnW1T6koSx79Kihdb"
    "e6u2hEnIk9dj1RBQkAeRB0HfbHE1nkNdBBfwPNxPvwlMR8Huf/Tqum9E+NEz3TP970zCQI5+LJZ/rr7F8dr6eT+br74Uvq3XD4fF"
    "4urmW3w/WR0sHuJ5Qu4Wy/vJOnm6/FpcPSzjye3G6H5WtEslr3g/mc4Lx0eb17rL46PF43o2ncfdpbV6vL+fLH9V4tnix5eCKtAL"
    "venXb+v0heLx0cPka9yP18OH7jJ5VsxauZ3ex/PVdDG3lvHdl4JWh7rjhKnF5i2X0/jH6tn/1urb4sfpcnrbSrpOIikVrPXkj348"
    "i2/W8e2m8+/J274U0v5ak1+Jk8k7Fg+t+G4dxbNZ0kNYsP5aLO77N5NZ/KXguc+edtIRSN6j7PKzV7tZU5ebplOTdEj/WCz+TF9p"
    "3KZ+JE5uvEhDmSQP32PTX912khH5zza89EkWf2r8/H8K9GQzEcnA/jFZxdFiNprerr99KQQF6za+mzzO1r3Fj3psBjdx5uZxtV7c"
    "Z6+kHdwsZqvNX+vH1tg5CFy37AV+9n7TajJi99P59nHy88mcDFVwULZdP1C2ZGkbS3vf0jmwA1e5nmjpGEtn39I+cP2SI3dZNobl"
    "fcNyXpiusXT3LN0DVS7JnnrGzNszK+eF6BtDf8/QyRvVwBgGLwx9VQodMcDQ2IV7dr407yWa+NLLWcgJTmVJ8yJrHMmEskXtp4uT"
    "M+eKskXtp4uXN5aK8kXtJ0yAp11Ruqj9fPFzDClh1H7GiPNAmaJepgqecUWp4uwPaYiddDL9vRCgkkxoJJ0X0sudBYcG09kfzMQ2"
    "r1LQeDr742m7QPKr9a+0yGetkDicfXXYoOLsNVImwZQ3giluS+2mcFcn68nx0XLxw1pu3p+WYydRUdbotrJv2H7RtlZpOUucSm31"
    "9rkqJUmR0PS4W7wxrLJlzksSiaQqkppITkRyKpK6SBoiORNJUyQtkbRF0hHJuUi6IrkQSU8kfZEMRDIUyaVIRiIZi+RKJNci0VpG"
    "cpZqOU21nKdaTlQtZ6qWU1XLuarlZNVytmo5XbWcr1pOWM1mbDEpL1mNsbc1JvBeX2PsrMbY2xqTHDXS9qfzdDXfXy8TPk0K2vq4"
    "PxhWa53BUXGd9Jq+9FSCjKlbZorQlqVVkm21Vm1UmSarOWbV8wpjVdtapXWdDyGqn5+3GMMTYxgIhlc9xug0p7c2Y1PPsTlhbBo0"
    "M2VxZnRvYFX1oMZYn4HJaQLWMp76Qqe92mWjNpJ6bef5XBs3uFTqAI/OAesCdrFlvuhKpyqF0TPNeoJpNxkFxqxP3kg96haX9IMc"
    "Mz3oNxizoTGTZqoada3G/Pti9j1mrC9zOu03IsZqlGPVqPW7jNk4r7MRNy5XeVYPcXzzzZrMb63WZP71MTlBZ1q5zhvdPte31nlm"
    "1TprV8kbofl6c4GAs41ybFudU6s6vbvjbKs52dCP56vF8pfVXS5u4hXXgqmiSiqGfS4htCmhSqyG8e30ZpJeD+GsTS0tK8E6akXW"
    "yXk07FuDmo7qNa4ea1Ncy9IRrMJ7bqrgbuXYOa462+OqCg68l0dRasYxFU8atrNhp3Hes7q9Rlv3rrjDqGnBdZnDKGBVwGqAnQB2"
    "ClgdsAZgZ4A1AWsB1gasA9g5YF3DxGSs/VxvL1Iyk3hBDXvMqQAxSZu5KdKnFpjD3ACwIWCXgI0AGwN2Bdg1YFojWEEwQrCKYA3B"
    "EwRPEcxkEsiFpbwtLHZwwFy3pYbK24U6d+6/RWUGRVvkMGlY3aLNVe0X9cJ0xpULg7grADKqGxS8RA3jIncRwFhxa1byvsRUii3z"
    "mM7aBjFD1TGjyJWJLfIVUyW2KGQavDCIcaNnYraZywEmZiaugYyGZjhKIaNsY8Z4P5LRGLR4JZtdAzOtZTtdASwCrApYDTCTx8pl"
    "5kBTJnNBmFR2dmdhR8+uucinkJ5dc3hh9OzKenZlPbtAz66sZ1fWs2mxxGRw3ZUF7cqCdmVBk/ucoF1Z0K4saFcWtCsL2pUF7cqC"
    "dmVBu7KgZTR0xdS9lNFIRmMZXcnoWkZaA1YBLKK04tRVRbAGWjV57HDMJDIvZzdXzh5zzf6lnj1Zz56sZ0/Wswf07Ml69mQ9e0DP"
    "nqxnT9azJ+uZ3Of07Ml69mQ9e7KePVnPnqxnT9azJ+vZk/Uso6En61lGIxmNPVkoV7LZNTDTGsGK3KiOAKuiRmvA0GQyr2gPKNrL"
    "VbRvFtweVLQvK3qLbCY9Il9WtA8U7cuK9mVF+0DRvqxoX1a0Lyua3OcU7cuKNmaKMev4sqR9WdK+LGlflrQvS9qXJS2joS9LWkYj"
    "GY1ldCWjaxlpDVgFsAiwKmA1wEwa83L2gZz9XDkHrzpAB7KcA1nOgSznAMg5kOUcyHI2LXI+1gNZzoEs50CWM7nPyTmQ5Uw+cnIO"
    "ZDkHspwDWc6BLOdAlnMgy1lGw0A+Tl0Gsp5lNKah4j5ZJ8b4fw2Y1ghWEIxkT3UVsBpgJpl5UQdA1EGuqMNXiTqURR3Kog5lUYdA"
    "1KEs6lAWdQhEHcqiDmVRh7KoyX1O1KEsavKRE3UoizqURR3Kog5lUYeyqENZ1DIahvIxWkYjGY1ppDhNh7L6rgHTGsEKgpHsqa4C"
    "VgPM5DKv6RBoOszVdJpf0v5kaiZ9jyRpw1hNG8aKmtpkVU2NcrImxu53k1mdGCdscpTd9VaSpZ0FwWnbQFbchpWZRjvEOHUbxsrb"
    "MFbfxDiBU/Ccwil4TuKADYlxIgdsBNgYsCvArgHTGsEKghGCVQRrCFJ681InS1brZInErl4hdgXEroDYFRC7QmJXQOwKiF1mdWKs"
    "2BUQuwJipyBYsSsgdgXEroDYFRC7AmJXQOwKiF0BsctsSIwVu8xGgI0BuwLsGjCtEawgGCFYRbCGIKW3IHaFxK7yxW6/Quw2ELsN"
    "xG4DsdtI7DYQuw3ELrM6MVbsxlHGlzNi3K6XLAh2U/sWeszMtA0rM850iDF254b5zGh3DQsZuwtytMQos0cRsmq3gdplNiTGql1m"
    "I8DGgF0Bdg2Y1ghWEIwQrCJYQ5DyW1C7jdRu56v9VVvhFO2gknbigU1wZOsyLkYIVhGsIXiC4CmCdQQbCJ4h2ESwhWAbwQ6C5wh2"
    "sxlVwozCPXFZ09ymuAxKu+LAdrjMltsPh+AQwUsERwiOEbxC8BpBrSGtQBpBWoW0BukJpKeQPikHbI9TtD/OB7XGbOdhjqAVw1zm"
    "ABMRY8pkldrk1g81MmSOdifE2LWFzOrEmJFqGGYz2jkjxp5IoI1yBvoMaxtWZjrskKPsiYRpk/3iHO0OY4b7IoPMkPaoR8abPjHm"
    "9GQA2BCwS2LMympEjKmQY8M8ZkSviLGLC8PYy4ME+a/Uyd7oCHSpq8iS8tvjFW8gM//6FEhKoxTXlKu7bLcUvGprnXJBMXBBMXBB"
    "MaD9aWwxcEExcEExkFmdOlSMNw0D+WoANtllUbDVwAXVwAXVwAXVwAXVgPZ6sdWAIFsNXFANXFANZDYE7JIYWw1cUA1okydjeGUg"
    "Xw5cVA5cVA5kd3QEutRVZEkZLpQDF5UDWVSaklwoB25uOfBeVQ48UA48UA48UA5oextbDjxQDjxQDmgPIZOD9QwySd8wkK8HYJNe"
    "FgZbDzxQDzxQDzxQDzxQD2inGFsPCLL1wEDFdNnPIOPrAMEhgpcZZCZrlEFG+ONsJtmqYCCXdddZs+zuvaxddvsezQpfGDxUGIAl"
    "5bpQGDxUGGR5acp2oTB4uYXB7PBLL2KIZQHs7zOMLws+KAu0R44tCz4oCz4oCzKrE2NPGXxQFMA+vywItij4oCj4oCj4oCj4oCjQ"
    "Fke2KBBki4IPFgk+WCTIbAjYJTF2keCDRYIPThl8sEbw0RrBR2sE2RsdgS51FVlSfgulwEelQJaURimuKVdRKQheda2SdsOJX5iu"
    "pV/K7LPXKo0tN4MRglUEawieIHiKYB3BBoJnCDYRbCHYRrCD4DmCXYJl6evz+FolNc2kZC+D4o+kiMnSRz4PEBwieIngCMExglcI"
    "XiOoNaRQMxqKRkPVaCgbDXWjoXB0XciJ3WpD2xZRrTE76tjzkRAsPEKw8KCNf+zCIwQLjxAsPGRWJ8YuPGRHz0CbzSwIxrBloM99"
    "/d8wj114hGDhEYKFR4gWHgTZhUcIFh4hWHjIbAjYJTF24RGChUcIFh4hWHiEaOERooWH7I2OQJe6iiwpv4WFR4gWHnKmaiHFd39q"
    "q0TnGQc+uAJhm5103Jf3iTGzFxHjPue0aTsgk2Y1AwNu6UCMifmUGi1xP/hBhsy8N4hx+5yIcecaWRRMoy0y5CRvWMhJnuwYdk6M"
    "mYpuFj0TxUUGmTB6GeQ0T10yMzwAbAjYJTEmM0Y0NNzRnhi3+YEYt32ZGCdrSm/FHbV0JaOcbQQ61VUQiaYcD1nVE2RVD2SlhSzf"
    "Vb3Z2Vj20GVH22zhMxspXq4SB+cD3bLMD+pZutU6j/Sgxv2MVIXaYrdDIFhFsGag/5S1349LR8XvXMy05cN9Oo3K+VFBNxuKhm0+"
    "5uVm+QzBJoItAxmB6LZhjCB1h2W74ZrPnJV9ELp/I94yihfAJoItA/l4yyBeju3G674vXhfFC2ATwZaBfLwuiJdju/F674vXQ/EC"
    "2ESwZSAfrwfi5dhuvP774vVRvAA2EWwZyMfrg3g5thtv8L54AxQvgE0EWwby8QYgXo7txhu+L94QxQtgE8GWgXy8IYiXY7u/xVd6"
    "V7xOCcSLYBPBloFsvIax8bJsN171vngVihfAJoItA/l4zdqCjVflxmsz8b5cYzUcGwUGYBPBloF8YGaDNBsYx3YDc14ZmFn3sHHJ"
    "rAlYixgH2w5YL5Eh2BLscAsmLq4yiEtmTcBaxPi4wLqIDFFc3MKIiwutgBBsIthywArIMD4RObYbmJe/q7LteKALju12QYuO8kEZ"
    "jF3b8UEvHNvtJXhlLwHohWO7vYSv7CUEvXBs20vx2a0i7uPl1819glbWzeJxnrbvFJ69vL0pUlcFh+nHFMUXpJeQzdVqDjmHm023"
    "L5G2Dys283rDPmxyr3edwwu2naSLc6EHdZiej74kbfuwy/aRtHXBt5XEeM6HmEbImiR+be+z9DTAydQs5rfT9PeIJ7Pt4XQ9nX+l"
    "OzOdlg/ryrZOVfroJ4/hYXp9ILG66z3OYmv96yH+UrhJWmqsCtbDcrpYTte/0vlK5njxEC8n60Uy8/F/HiezgnX78y69K1SyuDk+"
    "Sm+r9TibHBd+FY6K9CTJh0276T+cWznepptk91ybL9bRYr6eTOerymwy/3PHy/Q7e8Yn55lLrVrnn4Neo/3PpL3ffvvH1/XvJc7F"
    "1/ezuWEJBa/+pz2le7w+JST1vpD+7vwe9pI07CVp2LNL1ih5HL1MR+x4+clx/3P9fpOf6aXEz5jI0Pl/zONm/t40Hn5+voUf7udb"
    "fcwTn136aBf75cPLpED3kwxLt5gk/4SH6cVfa7x5wc4v1Y5cqEv/ozrdt0uHl685jKSbrD7buV6qVdsalQ9H6YNKH703lphyTh58"
    "uKK23r7Ry2fK57wcfbyXbz0+u3gcRx9VA/2ckfiwdYD6nIA2H8t/TkRPSyi2Kn9YR+qzhk591tCpdw3d3xXg4eiNC5H0LBetlz5s"
    "PJycqvlxHX2W3J+tBz5tht+6ZMmTFVyyvKGfPFV9VD95ovr4Jdg4Wc1sPozPXc544KT4HcuZpIP5+vwhBSvrW9LXX+nwzKJ4vo6X"
    "29s/m1tNtyfLr8mwWbP4bp10eeCXAs9T5SRRbTtw0jOj5fZ6Tno9cPGweUs5KDmqHHolzw7D9PPyPxbr9eKeZ9/iyW28TNnz1203"
    "PdO+WyzWAnu6E/bjg5UEkHi+uUXSl0Iywberm8lDnIxj8nfZn/4Vb64GrrY3qk53RtxN14MFXYmi5+ZGpZubkW69Otn0f3y0uL2t"
    "b144/sfk/uH31ubv6eZvtPlb0MvpZPavymJ2W9i8oAIrWvxIPJnOZhOrFU+W83T+0wFexpZdsr1//yyVSpN/99ePt8mLVmO+vWt4"
    "EoIhg1qvbZUta1SrNS37qPjkROoPuVZ89n/8PZ7TW4rPn6T/ZwbPn9xNl6t1ZrLzbPMkM9p9tjs6xey26Mf/BVBLAwQUAAAACADs"
    "YE9dgnS++XYJAACHyQAADQAAAHhsL3N0eWxlcy54bWztXW2P4jYQ/isR6te7xM57y6JeqZAqtaeqd1Irtf0QlgBR80JDuGP762sn"
    "LCQsBifEzgTY1S4hxDPPjMfjmbFJhuvsJfQ/LX0/U7ZRGK+fBsssW32rquvnpR956/fJyo/JJ/MkjbyMvE0X6nqV+t5sTRtFoYo1"
    "zVIjL4gHo2G8iSZRtlaek02cPQ3w/pRSvPw0exogyxgoBblxMvOfBn9+887Q3L9nf72Loui77wfqyUZmtdFMjSL1hfwUDdQd59Fw"
    "nsQlAO6gOENIepGvfPHCp8HYC4NpGtBmcy8KwpfiNKYnnpMwSZWMiE54IHpm/V/xMSreUa3s6ERBnKQ584LDMZ8PaeCFdbngSwSV"
    "j16aJl9P0p3uiIjk0ZBuQ6XjRkqvp6N0MX0aTCYu/lEzNW6Bfg/ixYz8rZVChqWXrskgOjAqc0Yt9c45LXKSb2bGDS0CSbDmczxa"
    "0FdtC+TlGQjqL17+ew6GKA5npGqVp1w+O3cxmWjkR1xXiaF6RlWacEMTIMmFzmgLuTAn8MrBlT8GHbEsdz2jabY2ljBMZEkjdtAz"
    "GY+lM96ZTBfCvtUyvujH8xca/gdhuA//EU1A6JnRcOVlmZ/GE/Imb5SffPORsjv+/LIisi9S7wVhc8DdYJ2EwYyyXIyr3mk8dvMR"
    "Pt19EMQzf+uTvMbKBVNLFK/kRTlNJoJ57WzDHShZQPWsvddd17VN0zGRiw3yJ0dY1x2PZQlrH4Q1ibCu7rgWdh2kGY4cANorgHfa"
    "e81wXcfVkWNiZDm6IQeBpF79YNLf1rGXtEeN1bF0WzN03bIcJKsDa7LJX4hPmybpzE/3Xs3WB6/nRsPQn2ekfRoslvQ1S1aUSZJl"
    "SUQOZoG3SGIvd3mvLdgtlbwYRHyrPws20eDVKx9DJVc2YMFJPL9WHppsmdetnlkDDZOBZtuI/thkwDEgNqFSD/c1nGj7+xZblQyb"
    "MSAZFv0m6lGLy/nVc0TwhEw6nTAsbFmWrdm66dyKYVwvuWzb4HRrTIT87pCzBa+xSXTkgpD3Q5Fv3QEgVXJjbsF+mw+ms3BODnXJ"
    "YQgXhhvB3G3AV9/dgpoKryFTd0y3FFvW7m+Icy4oI7h9CQXNUreULcFUUavRM0wRu/JrciK3rjMmsFLew+BvLYp/21J8HN9tUUeQ"
    "ItrPDQV1zcV6hbiq1xU6El+Ku2SHiNih6xi2odmGiS3crds5iaZB+nGSTtP+vCPRJQeqJ7E+VPcwuhbjYp4p6p7UVC9MASX51VVV"
    "3rD7xi2krVTjobtrk9Feid7SzPWwmq6TexCqqxtl35LZtJBV90JO+XsLRK5LcRVcgGCUuClCxlJg/Tqc+MpKZ7XBR62olocTVue7"
    "UTFvuwbUVzU9jKnjQKnnbgTqnAJqePRUSf2ILvqx0A9BzofRd1bH7jRLAFoKkd7zPa1Pg9NTT4bIbdadwVnDjWcawOvJEHYbgZXp"
    "leAs2UxD/5ignv+yzPlim7fyX2xyQtaLbe4nI1YhC/mm7gt2Z0Ift1JKXFa4Ymvx7mBN2vph+InS+2NeuQnkdl66l6NG7+QY7w+D"
    "MNwdFmSKNyq7EbYPrdxyK6uYVMogCkglNMjSG+FRVsGXJPthQ7QQ5+//3SSZ/2vqz4Nt/n47vwj8DHVvtQpfPoTBIo78HCY/w9HQ"
    "e22nLJM0+I9wo3fIeSYn/OIWOds5GxQSKrIBUeQzoFoQWRdKHUNUqAkRlAURlA0RlAsRlAMRFILpwlk+HLl3CEv54qdZ8FwDKC4B"
    "tSpATblAlczfZr8lmZcFlLSrXQHd6C/0zrWufE291Wdydnd3wn7I0Ryo1bXCL7oS1mwg2cqbTu/AYCJWbAQMJ2u2hQATGf3ACU+d"
    "l70rj251WLplBzs2LKAsfQKbrHrS7SyYdMKCNaY4ZinZoBvP+hCMgMdHgUh0eIBioflrm0ARAKBlE+1LxFeqG5oVE4Ue+mMNKvJa"
    "PhduB9QTQ78NMUpGhStiSJ5aWpRDr/jJrsVoHC9jCPFyeSrCrAWd3gAFNheBVijAKkm9fJmpXQwhw+NDCizIY+aisr1snamBCRpC"
    "8gQQZzvKhTDGeHBKrvg3nhKgKRSz9l/okAs+rA0vuuT6ROPY9gZQG5JLLC3pWnJM1pKuZWej7aCGHE6wUUOep9l2LXuPQTu6Bpd5"
    "lFCzKqEgpvBy3sECCiH9LONErG2V4EINLtQYsvEi1h5IeLq2eFJoEGOOCymWPOsxoPLkIhAcRNlmS7mIUyliw6z+lpXM9BKwJ+aS"
    "adgV1MAqVszqugFhRZqjQGFCzpaYoIFZARSc7Si3J3VLE8JqC89WChNEjMADFFixkhkrmuCCRY0nVIQwsLgCMEvypNC0rGr1spDT"
    "U9S9LD/JtuSWUIPO4JmoQWcUUMKdlnQNebkAs3JOE0JSVK2WlLaeOLDrOiykIDaT80EFUoLiggohWNM4gIIolvHpFARUHp2CiH/5"
    "dAoC6uls3QGX/JT7nlXTtSDXbPCZmzW0e98cVK1mwuo8GzLOys5LViQEzBMyw0wIARuj4xG0ji+PUybQa9w1EtrxsBWKWUD7sh/w"
    "CHVfttbdBOzebAk8gg2uUIhojacPX3DnAgpjwxSXSiHEC1xIgaT4HEghzHB8ZtobpCaEzi+71tICl1GxUqErnc1hCY1lmsPq9u52"
    "pSCqshfIgFD45AEKbVmd9dVy2cFd/e8yM78U3/m34pvflqBX0A1p9tIYlViFNoclNAdpDMsQGgM3h9XpBF2+FYrTB1hC04PGnajD"
    "NPluYz9WH3Yc+zFhderkmaURW3JFpw5QeDUGVrWpihTylhI27N58LbIK25J9K+CWtA15dZS5LAhiiwyfWUC2ZlbZB9VAyLf6rB0H"
    "XSLJ262QZ4YbttCiWOM+04QqVWyX0VhJJHm3HYtg1dNsKRmqunsgU+lhUZVHRe3PKrFHnx32MUkjLzywUaabIMyCeC/ncYPx0n/+"
    "RxmTs4cF6nIjXDz27MB9NJxtDw+I0qkkM4qUKuzoQWa49CAzzXVdw0HIMXXdcTWncFRFI6rf0XDlZUTyeJK/mS7G9ShVWqsFRTUH"
    "xkJXRsC40ttkCdeFJ0g+hGoo1ENvD6EeUCWbCmxXDRsdKEO7NwB3MUndq5YfKjiBQDt6QrFhmS62NMdAdi1sl+iwkdH/JA/IvGno"
    "V1MSQnPmz71NmH3ef/g0OBz/kj82F++v+pVmS7urDsc/0wf0IitPVyjiNYWTP3j309L3s9H/UEsDBBQAAAAIAOxgT11CSfAe0AAA"
    "AKYCAAALAAAAX3JlbHMvLnJlbHO1kstuAjEMRX9llD2YR8WiYlixYYcQP+AmnodmEkeOEcPfk7Ipg9pSIXXp5Obo5MrrA/WoLYfU"
    "tDEVg+9DKk2jGt8Bkm3IY5pypJBvKhaPmkepIaLtsCZYzGYrkHuG2azvmcXxEukvRK6q1tKW7clT0G/ADwlTHFFq0tIMPZxZug/m"
    "bpqhpti50sjOzQ286PLz78CTokNFsCw0iZJfi7aUvnQc230+TrfESGjxstDzcmhQCo7c70oY48ho+Y9G9pSU/ZOKbpmR0tunEozW"
    "cnMFUEsDBBQAAAAIAOxgT11bIEnZXgEAAGQCAAAPAAAAeGwvd29ya2Jvb2sueG1sjZHLTsMwEEV/JfIekkY82qrpBgRUQlBRRNdu"
    "PGlG+BHZE0L79YwTpVRiw8q+d+yZ6+NF5/znzrnP5NtoGwpREzXzNA1lDUaGS9eA5UrlvJHE0u/T0HiQKtQAZHSaZ9lNaiRasVyM"
    "vdY+XS7i5gOhC79+lMkXBtyhRjoUot9rEIlBiwaPoAqRiSTUrntyHo/OktSb0jutCzEZCh/gCcs/9ibmeZe70DvfW7TKdYW4yKez"
    "nFseTsYkyq5XW1RUFyKfZVcn7wlwXxM3ub6NB0nu3iShK8RNxrJCH6gf1QeVJeEX8NRBteQeUBP4e0nw6F3boN3HPIwjPePRsxvX"
    "Afzc/we9qyos4d6VrQFLA3sPOga0ocYmiMRKA4XYUKv4RLKyw3WuR0w8bqUGZMQRzz7Az5ELfqWGsGNCBRVaUC/cNLDP2Mu1T+LS"
    "95lNsnzKVFqt79h7tc9OqvHB468vfwBQSwMEFAAAAAgA7GBPXSQem6KtAAAA+AEAABoAAAB4bC9fcmVscy93b3JrYm9vay54bWwu"
    "cmVsc7WRPQ6DMAyFrxLlADVQqUMFTF1YKy4QBfMjEhLFrgq3L4UBkDp0YbKeLX/vyU6faBR3bqC28yRGawbKZMvs7wCkW7SKLs7j"
    "ME9qF6ziWYYGvNK9ahCSKLpB2DNknu6Zopw8/kN0dd1pfDj9sjjwDzC8XeipRWQpShUa5EzCaLY2wVLiy0yWoqgyGYoqlnBaIOLJ"
    "IG1pVn2wT06053kXN/dFrs3jCa7fDHB4dP4BUEsDBBQAAAAIAOxgT12hWMQ+IwEAAEwEAAATAAAAW0NvbnRlbnRfVHlwZXNdLnht"
    "bK2UTW6DMBCFr4K8jcBpF11UIZum2zaLXsA1Q7DiP3mGlNy+AySRWqUoEd1g4XnzPps3YvVxjIBZ56zHUjRE8VlK1A04hUWI4LlS"
    "h+QU8Wvayaj0Xu1APi6XT1IHT+App95DrFcbqFVrKXvteBtN8KVIYFFkL6OwZ5VCxWiNVsR1efDVL0p+IhTcOWiwMREXLBDyKqGv"
    "/A049b0fICVTQbZVid6UY5XsrEQ6WsBi2uLKGUNdGw1V0K3jlgJjAlVhA0DOFqPpYppM/IVhfD7M5g82U0BWblOIyIkluB93jqTv"
    "ziMbQSIzfcULka1n3w/6tCuo7mXrFim42fjR5kY4Z/sV0n4YBpTDMj/gnwN28b/hHJ8h7P97vvu1cMr4M18OP5H1N1BLAQIUAxQA"
    "AAAIAOxgT11Gx01IlQAAAM0AAAAQAAAAAAAAAAAAAACAAQAAAABkb2NQcm9wcy9hcHAueG1sUEsBAhQDFAAAAAgA7GBPXdVLjWgu"
    "AQAAhgIAABEAAAAAAAAAAAAAAIABwwAAAGRvY1Byb3BzL2NvcmUueG1sUEsBAhQDFAAAAAgA7GBPXYuCbljsBQAAjhoAABMAAAAA"
    "AAAAAAAAAIABIAIAAHhsL3RoZW1lL3RoZW1lMS54bWxQSwECFAMUAAAACADsYE9d1dKWFPoAAAApAgAAEwAAAAAAAAAAAAAAgAE9"
    "CAAAZG9jUHJvcHMvY3VzdG9tLnhtbFBLAQIUAxQAAAAIAOxgT13PjvV84REAACJ9AAAYAAAAAAAAAAAAAACAgWgJAAB4bC93b3Jr"
    "c2hlZXRzL3NoZWV0MS54bWxQSwECFAMUAAAACADsYE9dgnS++XYJAACHyQAADQAAAAAAAAAAAAAAgAF/GwAAeGwvc3R5bGVzLnht"
    "bFBLAQIUAxQAAAAIAOxgT11CSfAe0AAAAKYCAAALAAAAAAAAAAAAAACAASAlAABfcmVscy8ucmVsc1BLAQIUAxQAAAAIAOxgT11b"
    "IEnZXgEAAGQCAAAPAAAAAAAAAAAAAACAARkmAAB4bC93b3JrYm9vay54bWxQSwECFAMUAAAACADsYE9dJB6boq0AAAD4AQAAGgAA"
    "AAAAAAAAAAAAgAGkJwAAeGwvX3JlbHMvd29ya2Jvb2sueG1sLnJlbHNQSwECFAMUAAAACADsYE9doVjEPiMBAABMBAAAEwAAAAAA"
    "AAAAAAAAgAGJKAAAW0NvbnRlbnRfVHlwZXNdLnhtbFBLBQYAAAAACgAKAH8CAADdKQAAAAA="
)


def builtin_template() -> bytes:
    import base64
    return base64.b64decode(_TEMPLATE_B64)
