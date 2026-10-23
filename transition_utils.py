"""
transition_utils.py — shared helpers for CLC-LBU Transition Timetables
=====================================================================

Used by:  pages/71_Transition_Timetables.py
Can also be imported by the Daily Bulletin / Calendar:

    from transition_utils import get_students_client, bulletin_lines_for
    lines = bulletin_lines_for(get_students_client(), date.today())
    # -> ["AW · PY — Home school 8:45–12:00", ...]

Privacy model (matches Communication Log):
  * Supabase stores INITIALS only — never full name, EDID or DOB.
  * Full name / preferred name / EDID / DOB are typed (or read from an
    uploaded Word timetable) at export time, held in the browser session
    only, and written into the downloaded Word file. They are never saved.
"""

from __future__ import annotations

import io
import os
import re
from copy import deepcopy
from datetime import date, datetime, timedelta
from pathlib import Path

# ──────────────────────────────────────────────────────────────────────────
# SA DfE term dates (education.sa.gov.au/teaching/term-dates)
# Add the next year here each year.
# ──────────────────────────────────────────────────────────────────────────
TERM_DATES: dict[int, dict[int, tuple[date, date]]] = {
    2026: {
        1: (date(2026, 1, 27), date(2026, 4, 10)),
        2: (date(2026, 4, 27), date(2026, 7, 3)),
        3: (date(2026, 7, 20), date(2026, 9, 25)),
        4: (date(2026, 10, 12), date(2026, 12, 11)),
    },
    2027: {
        1: (date(2027, 1, 27), date(2027, 4, 9)),
        2: (date(2027, 4, 26), date(2027, 7, 2)),
        3: (date(2027, 7, 19), date(2027, 9, 24)),
        4: (date(2027, 10, 11), date(2027, 12, 10)),
    },
    2028: {
        1: (date(2028, 1, 31), date(2028, 4, 13)),
        2: (date(2028, 5, 1), date(2028, 7, 7)),
        3: (date(2028, 7, 24), date(2028, 9, 29)),
        4: (date(2028, 10, 16), date(2028, 12, 15)),
    },
}

DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

# Default session times (match the current Word timetables)
CLC_START, CLC_END = "9:00", "2:45"
HOME_AM_START, HOME_AM_END = "8:45", "12:00"

LOC_CLC, LOC_HOME, LOC_OTHER = "CLC", "HOME", "OTHER"

TABLE_TIMETABLES = "transition_timetables"
TABLE_HISTORY = "transition_timetable_history"

LOGO_PATH = Path(__file__).parent / "data" / "transition_timetables" / "clc_logo.jpeg"


# ──────────────────────────────────────────────────────────────────────────
# Term / week maths
# ──────────────────────────────────────────────────────────────────────────
def week_monday(year: int, term: int, week: int) -> date:
    """Monday of a given term week (Week 1 = the week the term starts in)."""
    start = TERM_DATES[year][term][0]
    first_monday = start - timedelta(days=start.weekday())
    return first_monday + timedelta(weeks=week - 1)


def weeks_in_term(year: int, term: int) -> int:
    start, end = TERM_DATES[year][term]
    first_monday = start - timedelta(days=start.weekday())
    return (end - first_monday).days // 7 + 1


def term_week_for(d: date) -> tuple[int, int, int] | None:
    """(year, term, week) for a date inside a term, else None."""
    terms = TERM_DATES.get(d.year, {})
    for term, (start, end) in terms.items():
        first_monday = start - timedelta(days=start.weekday())
        if first_monday <= d <= end:
            return d.year, term, (d - first_monday).days // 7 + 1
    return None


def next_term_week(year: int, term: int, week: int) -> tuple[int, int, int]:
    if week < weeks_in_term(year, term):
        return year, term, week + 1
    if term < 4:
        return year, term + 1, 1
    if (year + 1) in TERM_DATES:
        return year + 1, 1, 1
    return year, term, week  # no data for next year yet


def current_or_next_term_week(today: date | None = None) -> tuple[int, int, int]:
    today = today or date.today()
    tw = term_week_for(today)
    if tw:
        return tw
    # In holidays → first week of next term
    for y in sorted(TERM_DATES):
        for t in sorted(TERM_DATES[y]):
            if TERM_DATES[y][t][0] >= today:
                return y, t, 1
    y = max(TERM_DATES)
    return y, 4, weeks_in_term(y, 4)


# ──────────────────────────────────────────────────────────────────────────
# Data model
#   week = {"year", "term", "week", "days": [day x5]}
#   day  = {"date": "YYYY-MM-DD", "loc": CLC|HOME|OTHER, "start", "end",
#           "full_day": bool, "label": "LAST DAY"/"", "other": "place/notes"}
# ──────────────────────────────────────────────────────────────────────────
def make_day(d: date, loc: str = LOC_CLC, start: str = "", end: str = "",
             full_day: bool = False, label: str = "", other: str = "") -> dict:
    if loc == LOC_CLC and not start and not end and not full_day:
        start, end = CLC_START, CLC_END
    return {"date": d.isoformat(), "loc": loc, "start": start, "end": end,
            "full_day": full_day, "label": label, "other": other}


def make_week(year: int, term: int, week: int, pattern: list[dict] | None = None) -> dict:
    mon = week_monday(year, term, week)
    days = []
    for i in range(5):
        d = mon + timedelta(days=i)
        if pattern and i < len(pattern):
            nd = deepcopy(pattern[i])
            nd["date"] = d.isoformat()
            nd["label"] = ""  # labels like LAST DAY don't carry over
            days.append(nd)
        else:
            days.append(make_day(d))
    return {"year": year, "term": term, "week": week, "days": days}


def sort_weeks(weeks: list[dict]) -> list[dict]:
    return sorted(weeks, key=lambda w: (w["year"], w["term"], w["week"]))


def fmt_time_range(day: dict) -> str:
    if day.get("full_day"):
        return "Full Day"
    s, e = day.get("start", ""), day.get("end", "")
    if s and e:
        return f"{s} – {e}"
    return s or e or ""


def short_day_text(day: dict, home_school: str = "Home school") -> str:
    """Compact text for screens / Bulletin (no names)."""
    loc = day.get("loc")
    place = {LOC_CLC: "CLC", LOC_HOME: home_school or "Home school"}.get(
        loc, day.get("other") or "Other")
    bits = [place, fmt_time_range(day)]
    if day.get("label"):
        bits.append(day["label"].upper())
    return " ".join(b for b in bits if b)


def doc_day_text(day: dict, name: str, home_school_short: str) -> list[str]:
    """Sentence used in the Word export, e.g. 'Abi will attend Bellevue. 8:45 – 12:00'."""
    loc = day.get("loc")
    if loc == LOC_CLC:
        place = "CLC"
    elif loc == LOC_HOME:
        place = home_school_short or "home school"
    else:
        place = day.get("other") or ""
    who = name or "Student"
    first = f"{who} will attend {place}." if place else f"{who}: {day.get('other') or 'not attending'}."
    tail = day["label"].upper() if day.get("label") else fmt_time_range(day)
    return [first, tail] if tail else [first]


def school_short_name(school: str) -> str:
    """'Bellevue Heights Primary School' -> 'Bellevue Heights'."""
    s = re.sub(r"\b(Primary|Junior Primary|High|Area|R-12|R-7|Secondary|College)\b.*$", "",
               school or "", flags=re.I).strip()
    s = re.sub(r"\bSchool\b", "", s, flags=re.I).strip()
    return s or (school or "").strip()


def week_summary(week: dict) -> dict:
    """Counts used for the step-down chart."""
    out = {"At CLC": 0, "Home school – part day": 0, "Home school – full day": 0, "Other": 0}
    for d in week["days"]:
        if d["loc"] == LOC_CLC:
            out["At CLC"] += 1
        elif d["loc"] == LOC_HOME:
            out["Home school – full day" if d.get("full_day") else "Home school – part day"] += 1
        else:
            out["Other"] += 1
    return out


def last_clc_day(weeks: list[dict]) -> date | None:
    days = [date.fromisoformat(d["date"]) for w in weeks for d in w["days"] if d["loc"] == LOC_CLC]
    return max(days) if days else None


def day_for(weeks: list[dict], d: date) -> dict | None:
    iso = d.isoformat()
    for w in weeks:
        for day in w["days"]:
            if day["date"] == iso:
                return day
    return None


def initials_from_name(full_name: str) -> str:
    parts = [p for p in re.split(r"[\s.\-]+", full_name or "") if p]
    return "".join(p[0].upper() for p in parts[:3])


def recommencement_text(d: date | None) -> str:
    if not d:
        return ""
    tw = term_week_for(d)
    if tw:
        return f"{d.strftime('%A')} Term {tw[1]} Week {tw[2]} - {d.strftime('%d/%m/%Y')}"
    return d.strftime("%A %d/%m/%Y")


# ──────────────────────────────────────────────────────────────────────────
# Read an existing Word transition timetable
# ──────────────────────────────────────────────────────────────────────────
_TIME_RE = re.compile(r"(\d{1,2}[:.]\d{2})\s*(?:am|pm)?\s*[–\-—to]+\s*(\d{1,2}[:.]\d{2})", re.I)
_DATE_DM_RE = re.compile(r"(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?")


def _cell_texts(row) -> list[str]:
    """Unique cell texts in order (Word repeats merged cells)."""
    out, prev_tc = [], None
    for c in row.cells:
        if c._tc is prev_tc:
            continue
        prev_tc = c._tc
        out.append(re.sub(r"\s+", " ", c.text.replace("\n", " ")).strip())
    return out


def _after(label: str, text: str) -> str:
    m = re.search(re.escape(label) + r"\s*:?\s*(.*)", text, flags=re.I)
    return m.group(1).strip() if m else ""


def _parse_day_text(text: str, d: date, home_tokens: list[str]) -> dict:
    t = text or ""
    low = t.lower()
    label = "LAST DAY" if "last day" in low else ""
    full = "full day" in low
    m = _TIME_RE.search(t)
    start, end = (m.group(1).replace(".", ":"), m.group(2).replace(".", ":")) if m else ("", "")

    if re.search(r"\bclc\b|cowandilla", low):
        loc = LOC_CLC
    elif any(tok and tok in low for tok in home_tokens) or "home school" in low:
        loc = LOC_HOME
    elif not t.strip():
        loc = LOC_CLC  # blank → default, staff will check
    else:
        loc = LOC_OTHER

    other = ""
    if loc == LOC_OTHER:
        mm = re.search(r"attend\s+(.+?)(?:\.|$)", t, flags=re.I)
        other = (mm.group(1) if mm else t).strip()

    if loc == LOC_CLC and not (start or full):
        start, end = CLC_START, CLC_END
        if label:  # "LAST DAY" with no times → normal CLC day
            pass
    return {"date": d.isoformat(), "loc": loc, "start": start, "end": end,
            "full_day": full, "label": label, "other": other}


def parse_timetable_docx(file_bytes: bytes) -> dict:
    """
    Returns {"header": {...}, "sensitive": {...}, "weeks": [...], "warnings": [...]}
    'sensitive' (full name, EDID, DOB) is for the export form only — never save it.
    """
    import docx  # python-docx

    doc = docx.Document(io.BytesIO(file_bytes))
    header = {"home_school": "", "home_school_short": "", "home_school_contact": "", "bsc": "",
              "year_level": "", "recommencement_date": None}
    sensitive = {"full_name": "", "preferred_name": "", "edid": "", "dob": ""}
    warnings: list[str] = []
    doc_year = None

    # ── pass 1: header fields anywhere in the tables
    for table in doc.tables:
        for row in table.rows:
            for txt in _cell_texts(row):
                low = txt.lower()
                if low.startswith("student") and "name" in low:
                    sensitive["full_name"] = sensitive["full_name"] or _after("Name", txt)
                elif low.startswith("school:"):
                    header["home_school"] = header["home_school"] or _after("School", txt)
                elif "contact person" in low:
                    header["home_school_contact"] = header["home_school_contact"] or _after("Person", txt)
                elif low.startswith("bsc"):
                    header["bsc"] = header["bsc"] or _after("BSC", txt)
                elif low.startswith("year level"):
                    header["year_level"] = header["year_level"] or _after("Year Level", txt)
                elif low.startswith("edid"):
                    sensitive["edid"] = sensitive["edid"] or _after("EDID", txt)
                elif low.startswith("dob"):
                    sensitive["dob"] = sensitive["dob"] or _after("DOB", txt)
                elif low.startswith("date:") and not doc_year:
                    m = _DATE_DM_RE.search(txt)
                    if m and m.group(3):
                        y = int(m.group(3))
                        doc_year = y + 2000 if y < 100 else y
                elif "recommencement" in low:
                    pass
            # Recommencement value sits in the cell after its label
            texts = _cell_texts(row)
            if texts and "recommencement" in texts[0].lower():
                joined = " ".join(texts[1:]) or texts[0]
                m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", joined)
                if m:
                    y = int(m.group(3)); y = y + 2000 if y < 100 else y
                    try:
                        header["recommencement_date"] = date(y, int(m.group(2)), int(m.group(1))).isoformat()
                    except ValueError:
                        pass

    doc_year = doc_year or date.today().year
    full = sensitive["full_name"]
    if full:
        sensitive["preferred_name"] = full.split()[0]

    # Short name staff actually used in the day cells, e.g. "attend Bellevue."
    seen = {}
    for table in doc.tables:
        for row in table.rows:
            for txt in _cell_texts(row):
                mm = re.search(r"will attend\s+([A-Za-z' ]+?)\s*(?:\.|\d|full|$)", txt, flags=re.I)
                if mm and mm.group(1).strip().lower() not in ("clc", "cowandilla"):
                    nm = mm.group(1).strip()
                    seen[nm] = seen.get(nm, 0) + 1
    if seen:
        header["home_school_short"] = max(seen, key=seen.get)
    elif header["home_school"]:
        header["home_school_short"] = school_short_name(header["home_school"])

    home_tokens = [header["home_school_short"].lower()] if header["home_school_short"] else []
    if header["home_school"]:
        home_tokens += [w.lower() for w in school_short_name(header["home_school"]).split() if len(w) > 3]
        home_tokens.append(school_short_name(header["home_school"]).lower())

    # ── pass 2: week blocks (Term row → Week/Date row → activity row)
    weeks: list[dict] = []
    for table in doc.tables:
        rows = [_cell_texts(r) for r in table.rows]
        i = 0
        while i < len(rows):
            first = rows[i][0] if rows[i] else ""
            mterm = re.search(r"term\s*:?\s*(\d)", first, flags=re.I)
            if mterm and i + 2 < len(rows):
                term = int(mterm.group(1))
                week_row, act_row = rows[i + 1], rows[i + 2]
                mweek = re.search(r"week\s*:?\s*(\d{1,2})", week_row[0] if week_row else "", flags=re.I)
                if not mweek:
                    i += 1
                    continue
                wk = int(mweek.group(1))
                # Monday dd/mm from the date row (used to pick the right year)
                mon_dm = None
                if len(week_row) > 1:
                    m = _DATE_DM_RE.search(week_row[1])
                    if m:
                        mon_dm = (int(m.group(1)), int(m.group(2)))
                year = None
                for y in (doc_year, doc_year + 1, doc_year - 1):
                    if y in TERM_DATES and term in TERM_DATES[y]:
                        mon = week_monday(y, term, wk)
                        if mon_dm is None or (mon.day, mon.month) == mon_dm:
                            year = y
                            break
                if year is None:
                    warnings.append(f"Term {term} Week {wk}: dates didn't match the SA term calendar — check this week.")
                    year = doc_year if doc_year in TERM_DATES else max(TERM_DATES)
                mon = week_monday(year, term, wk)
                acts = act_row[1:6] if len(act_row) >= 6 else (act_row + [""] * 6)[1:6]
                days = [_parse_day_text(acts[k], mon + timedelta(days=k), home_tokens) for k in range(5)]
                for k, dd in enumerate(days):
                    if dd["loc"] == LOC_OTHER:
                        warnings.append(f"T{term} W{wk} {DAY_NAMES[k]}: couldn't tell CLC vs home school — set to 'Other', please check.")
                weeks.append({"year": year, "term": term, "week": wk, "days": days})
                i += 3
            else:
                i += 1

    if not weeks:
        warnings.append("No week blocks found — is this the standard CLC transition timetable layout?")
    return {"header": header, "sensitive": sensitive, "weeks": sort_weeks(weeks), "warnings": warnings}


# ──────────────────────────────────────────────────────────────────────────
# Word export (matches the current CLC transition timetable layout)
# ──────────────────────────────────────────────────────────────────────────
GREEN = "C5E0B3"
YELLOW = "FFFF00"
FONT = "Open Sans"


def _shade(cell, fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _write(cell, text: str, bold=False, size=10, align="center", lines: list[str] | None = None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    from docx.enum.table import WD_ALIGN_VERTICAL
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    p.paragraph_format.space_after = Pt(0)
    for li, line in enumerate(lines if lines is not None else [text]):
        if li:
            p.add_run().add_break()
        r = p.add_run(line)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.name = FONT


def _label_value(cell, label: str, value: str, size=10):
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: "); r.bold = True; r.font.size = Pt(size); r.font.name = FONT
    r = p.add_run(value or ""); r.font.size = Pt(size); r.font.name = FONT


def _keep_row(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)
    for c in row.cells:
        for p in c.paragraphs:
            p.paragraph_format.keep_with_next = True


def build_timetable_docx(header: dict, weeks: list[dict], sensitive: dict) -> bytes:
    """
    header: home_school, home_school_contact, bsc, year_level, recommencement_date (iso)
    sensitive: full_name, preferred_name, edid, dob  (session-only, never stored)
    """
    import docx
    from docx.shared import Pt, Cm, Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Cm(1.27))
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)

    # Title + logo
    top = doc.add_table(rows=1, cols=2)
    top.alignment = WD_TABLE_ALIGNMENT.CENTER
    top.columns[0].width = Cm(14.5); top.columns[1].width = Cm(3.5)
    c0, c1 = top.rows[0].cells
    c0.width, c1.width = Cm(14.5), Cm(3.5)
    _write(c0, "COWANDILLA LEARNING CENTRE", bold=True, size=16)
    if LOGO_PATH.exists():
        c1.text = ""
        p = c1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run().add_picture(str(LOGO_PATH), height=Cm(2.6))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # Header block (5 rows like the original)
    rec = date.fromisoformat(header["recommencement_date"]) if header.get("recommencement_date") else None
    h = doc.add_table(rows=5, cols=2)
    h.style = "Table Grid"
    h.alignment = WD_TABLE_ALIGNMENT.CENTER
    _label_value(h.cell(0, 0), "Student's Name", sensitive.get("full_name", ""))
    _label_value(h.cell(0, 1), "Date", date.today().strftime("%d/%m/%Y"))
    m = h.cell(1, 0).merge(h.cell(1, 1))
    _label_value(m, "School", header.get("home_school", ""))
    _label_value(h.cell(2, 0), "Home School Contact Person", header.get("home_school_contact", ""))
    _label_value(h.cell(2, 1), "BSC", header.get("bsc", ""))
    _label_value(h.cell(3, 0), "Year Level", header.get("year_level", ""))
    yl = h.cell(3, 1)
    _label_value(yl, "EDID", sensitive.get("edid", ""))
    p = yl.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    r = p.add_run("DOB: "); r.bold = True; r.font.size = Pt(10); r.font.name = FONT
    r = p.add_run(sensitive.get("dob", "")); r.font.size = Pt(10); r.font.name = FONT
    m = h.cell(4, 0).merge(h.cell(4, 1))
    _label_value(m, "Recommencement Date at Home School", recommencement_text(rec))
    for row in h.rows:
        row.cells[0].width = Cm(9.5)
        if len(row.cells) > 1:
            row.cells[1].width = Cm(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # Week blocks — one table, 3 rows per week
    name = sensitive.get("preferred_name") or (sensitive.get("full_name", "").split() or [""])[0]
    short = header.get("home_school_short") or school_short_name(header.get("home_school", ""))
    wks = sort_weeks(weeks)
    t = doc.add_table(rows=0, cols=6)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(3.0)] + [Cm(3.2)] * 5
    for w in wks:
        mon = week_monday(w["year"], w["term"], w["week"])
        r1 = t.add_row(); r2 = t.add_row(); r3 = t.add_row()
        _write(r1.cells[0], f"Term: {w['term']}", bold=True, align="left")
        _write(r2.cells[0], f"Week: {w['week']}", bold=True, align="left")
        _write(r3.cells[0], "")
        for k in range(5):
            day = w["days"][k]
            d = mon + timedelta(days=k)
            _write(r1.cells[k + 1], DAY_NAMES[k], bold=True)
            _shade(r1.cells[k + 1], GREEN)
            home = day["loc"] != LOC_CLC
            _write(r2.cells[k + 1], "", bold=home, lines=["Date:", d.strftime("%d/%m")])
            # make 'Date:' bold always
            r2.cells[k + 1].paragraphs[0].runs[0].bold = True
            _write(r3.cells[k + 1], "", bold=home, lines=doc_day_text(day, name, short))
            if home:
                _shade(r2.cells[k + 1], YELLOW)
                _shade(r3.cells[k + 1], YELLOW)
        for row in (r1, r2, r3):
            for ci, c in enumerate(row.cells):
                c.width = widths[ci]
        _keep_row(r1); _keep_row(r2)
        r3.height = Cm(1.6)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_filename(sensitive: dict, initials: str) -> str:
    full = (sensitive.get("full_name") or "").split()
    if len(full) >= 2:
        stem = f"{full[0][0].upper()}.{full[-1].upper()}"
    else:
        stem = initials or "Student"
    return f"{stem} - Transition Timetable.docx"


# ──────────────────────────────────────────────────────────────────────────
# Supabase (students project — same one as the Student Register / Archive)
# ──────────────────────────────────────────────────────────────────────────
def _secret(*names: str) -> str:
    for n in names:
        v = os.environ.get(n)
        if v:
            return v
        try:
            import streamlit as st
            if n in st.secrets:
                return str(st.secrets[n])
        except Exception:
            pass
    return ""


def get_students_client():
    """Supabase client for the students/behaviour project. Returns None if not configured."""
    from supabase import create_client
    url = _secret("SUPABASE_URL_STUDENTS", "SUPABASE_URL")
    key = _secret("SUPABASE_KEY_STUDENTS", "SUPABASE_SERVICE_KEY_STUDENTS",
                  "SUPABASE_SERVICE_ROLE_KEY_STUDENTS", "SUPABASE_ANON_KEY_STUDENTS",
                  "SUPABASE_KEY", "SUPABASE_SERVICE_KEY", "SUPABASE_ANON_KEY")
    if not url or not key:
        return None
    return create_client(url, key)


def load_timetables(client, include_complete: bool = True) -> list[dict]:
    q = client.table(TABLE_TIMETABLES).select("*").order("updated_at", desc=True)
    if not include_complete:
        q = q.eq("status", "active")
    return q.execute().data or []


def save_timetable(client, record: dict, saved_by: str, change_note: str = "") -> dict:
    """Insert or update; every save also writes a history snapshot."""
    now = datetime.utcnow().isoformat()
    payload = {k: record.get(k) for k in (
        "student_initials", "register_id", "program", "home_school", "home_school_short", "home_school_contact",
        "bsc", "year_level", "recommencement_date", "status", "weeks", "notes")}
    payload["weeks"] = sort_weeks(payload.get("weeks") or [])
    payload["updated_at"] = now
    payload["updated_by"] = saved_by
    if record.get("id"):
        payload["version"] = int(record.get("version") or 1) + 1
        res = client.table(TABLE_TIMETABLES).update(payload).eq("id", record["id"]).execute()
    else:
        payload["version"] = 1
        payload["created_by"] = saved_by
        payload.setdefault("status", "active")
        res = client.table(TABLE_TIMETABLES).insert(payload).execute()
    row = (res.data or [payload])[0]
    try:
        client.table(TABLE_HISTORY).insert({
            "timetable_id": row.get("id"),
            "version": row.get("version"),
            "snapshot": {k: row.get(k) for k in payload},
            "saved_by": saved_by,
            "change_note": change_note,
        }).execute()
    except Exception:
        pass  # history is best-effort
    return row


def load_history(client, timetable_id: str) -> list[dict]:
    return (client.table(TABLE_HISTORY).select("*").eq("timetable_id", timetable_id)
            .order("version", desc=True).execute().data or [])


def set_status(client, timetable_id: str, status: str):
    client.table(TABLE_TIMETABLES).update(
        {"status": status, "updated_at": datetime.utcnow().isoformat()}).eq("id", timetable_id).execute()


def delete_timetable(client, timetable_id: str):
    client.table(TABLE_HISTORY).delete().eq("timetable_id", timetable_id).execute()
    client.table(TABLE_TIMETABLES).delete().eq("id", timetable_id).execute()


# ──────────────────────────────────────────────────────────────────────────
# For the Daily Bulletin / Calendar
# ──────────────────────────────────────────────────────────────────────────
def where_is_everyone(timetables: list[dict], d: date) -> list[dict]:
    """[{initials, program, day, home_school}] for active timetables on date d."""
    out = []
    for tt in timetables:
        if tt.get("status") != "active":
            continue
        day = day_for(tt.get("weeks") or [], d)
        if day:
            out.append({"initials": tt.get("student_initials", ""), "program": tt.get("program", ""),
                        "home_school": tt.get("home_school_short") or school_short_name(tt.get("home_school", "")),
                        "day": day})
    return sorted(out, key=lambda x: (x["program"], x["initials"]))


def bulletin_lines_for(client, d: date | None = None) -> list[str]:
    """Initials-only lines for the Bulletin, e.g. 'AW · PY — Home school 8:45–12:00'."""
    d = d or date.today()
    try:
        tts = load_timetables(client, include_complete=False)
    except Exception:
        return []
    lines = []
    for e in where_is_everyone(tts, d):
        day = e["day"]
        if day["loc"] == LOC_CLC and not day.get("label"):
            continue  # normal CLC day — nothing to flag
        place = "Home school" if day["loc"] == LOC_HOME else (day.get("other") or "Other")
        if day["loc"] == LOC_CLC:
            place = "CLC"
        bits = [place, fmt_time_range(day).replace(" – ", "–")]
        if day.get("label"):
            bits.append(f"({day['label'].upper()})")
        lines.append(f"{e['initials']} · {e['program']} — " + " ".join(b for b in bits if b))
    return lines


# ──────────────────────────────────────────────────────────────────────────
# Document Archive (uses the portal's archive_utils.py)
# ──────────────────────────────────────────────────────────────────────────
ARCHIVE_SOURCE_TOOL = "Transition Timetables"
ARCHIVE_DOC_TYPE = "Transition Timetable"


def archive_timetable(record: dict, saved_by: str, change_note: str = "") -> dict | None:
    """File an INITIALS-ONLY Word copy in the student's Document Archive folder.
    Replaces the student's previous timetable there (older copy goes to the
    archive's restorable Deleted list). Never raises; returns the archive
    record or None."""
    try:
        import archive_utils as au
    except ImportError:
        return None
    try:
        ini = (record.get("student_initials") or "").strip().upper()
        weeks = record.get("weeks") or []
        if not ini or not weeks:
            return None
        sens = {"full_name": ini, "preferred_name": ini, "edid": "", "dob": ""}
        file_bytes = build_timetable_docx(record, weeks, sens)
        tw = term_week_for(date.today())
        term, year = (f"Term {tw[1]}", tw[0]) if tw else (None, None)
        tags = [t for t in (record.get("program"), record.get("home_school_short")) if t]
        return au.archive_document(
            source_tool=ARCHIVE_SOURCE_TOOL,
            doc_type=ARCHIVE_DOC_TYPE,
            title=f"Transition Timetable — {ini} — v{record.get('version') or 1}",
            file_bytes=file_bytes,
            filename=f"Transition_Timetable_{ini}.docx",
            student_initials=ini,
            term=term,
            year=year,
            tags=tags,
            created_by=saved_by or "unknown",
            notes=(change_note or None),
            supersede=True,
        )
    except Exception:
        return None
